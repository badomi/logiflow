"""로컬 LLM 추출 → 룰 검증 → 보완 문항 / 견적서 생성 흐름 테스트.

실제 LLM 대신 FakeLlm을 쓴다: 메일 입력을 보고 "로컬 LLM이 이렇게 답했다"고 가정한 JSON을 돌려준다.
LLM 품질(정밀도, NFR-01)은 여기서 재지 않는다 — 실제 모델로 샘플 측정이 따로 필요하다.
"""

import base64
from datetime import date

import pytest
from openpyxl import load_workbook

from app import pipeline, validation
from app.config import settings
from app.db import SessionLocal
from app.extraction import LlmExtractor, join_broken_list, normalize, strip_quoted
from app.llm import LlmError, ensure_local
from app.models import Case
from app.seed import seed_containers, seed_sample_rates
from app.validation import RuleValidator

from .conftest import sample

TODAY = date(2026, 9, 29)

# 01_lcl_request.eml 만 읽었을 때 LLM이 낼 법한 값 (8-1 LCL: 선적항 미정, 조건 2개, 부피 불일치)
FIRST = {
    "commodity": "스프링노트(일반 문구류)", "qty": 10, "qtyUnit": "박스", "packing": None,
    "boxL": 310, "boxW": 450, "boxH": 270, "totalCbm": 0.19,
    "pol": None, "polName": None, "pod": None, "podName": "호주 시드니", "pickupLocation": "인천",
    "cargoReadyDate": None, "incoterms": "FOB 및 CIF", "containerType": "LCL",
    "grossWeightKg": 500, "invoiceValue": None, "ccy": None, "paymentTerm": None,
    "customerName": None, "contactName": None, "contactEmail": None,
    "multipleItems": False, "evidence": {"pod": "도착항: 호주 시드니"},
}
# 회신(03_lcl_reply.eml)까지 읽었을 때 + 부족분을 모두 채운 경우
AFTER_REPLY = {
    **FIRST, "pol": "KRINC", "polName": "인천", "cargoReadyDate": "2026-10-20", "incoterms": "CIF",
    "invoiceValue": 3000, "ccy": "USD", "paymentTerm": "T/T 30% Advance", "totalCbm": 0.38,
}


class FakeLlm:
    def __init__(self, first=FIRST, after_reply=AFTER_REPLY):
        self.first, self.after_reply, self.calls = first, after_reply, []

    def complete_json(self, system, user, schema):
        self.calls.append(user)
        return dict(self.after_reply if "인보이스 금액" in user else self.first)


class BrokenLlm:
    def complete_json(self, system, user, schema):
        raise LlmError("LLM 서버(http://localhost:11434)에 연결하지 못했습니다")


@pytest.fixture
def connect(monkeypatch):
    """가짜 LLM으로 파이프라인 연결 + 마스터·테스트 요율 등록. 끝나면 연결 해제."""
    monkeypatch.setattr(validation, "today_kst", lambda: TODAY)
    with SessionLocal() as s:
        seed_containers(s)
        seed_sample_rates(s)
        s.commit()

    def _connect(llm):
        pipeline.register(extractor=LlmExtractor(client=llm), validator=RuleValidator())
        return llm

    yield _connect
    pipeline.register(extractor=pipeline._NotConnected(), validator=pipeline._NotConnected())


def body(name: str) -> dict:
    return {"emlBase64": base64.b64encode(sample(name)).decode(), "actor": "quote@leona.example.com"}


def create(client, name="01_lcl_request.eml") -> dict:
    case_id = client.post("/api/cases", json=body(name)).json()["case"]["caseId"]
    return client.get(f"/api/cases/{case_id}").json()


# ------------------------------------------------------------------ 전체 흐름 (MUST-SHIP ①②③④)


def test_missing_fields_become_supplement_questions(client, connect):
    connect(FakeLlm())
    case = create(client)
    assert case["status"] == "정보부족"
    questions = case["extraction"]["validation"]["questions"]
    joined = "\n".join(questions)
    assert "선적항" in joined and "화물 준비일" in joined and "결제 조건" in joined
    assert "FOB / CIF" in joined  # 조건 2개 → 하나로 정해 달라
    assert "0.38CBM" in joined  # 박스 규격 기준 0.38 vs 기재 0.19
    assert all("(" in q for q in questions)  # 국·영문 병기 (FR-510)

    # 문항이 그대로 A트랙 보완 요청 초안이 된다 (FR-304·305)
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/supplement", json={"questions": questions}).json()
    assert draft["subject"].startswith(f"[{case['caseId']}]")
    assert draft["htmlBody"].count("<li>") == len(questions)


def test_reply_merge_fills_gaps_and_builds_quote(client, connect):
    llm = connect(FakeLlm())
    case_id = create(client)["caseId"]
    client.post(f"/api/cases/{case_id}/replies", json=body("03_lcl_reply.eml"))
    case = client.get(f"/api/cases/{case_id}").json()

    assert "인보이스 금액" in llm.calls[-1] and "최초 요청" in llm.calls[-1]  # 원문+회신을 함께 다시 읽음
    assert case["status"] == "계산완료"
    result = case["extraction"]["validation"]
    assert result["questions"] == []
    quote = result["quote"]
    # 0.38CBM·0.5톤 → 최소 1 R/T. USD 500 / KRW 5,500+6,500+203+5,000+90,000 (통관·보험은 AT COST)
    assert quote["totals"] == {"USD": 500.0, "KRW": 107203.0}

    # 견적서가 A트랙 송부 초안에 자동 첨부된다 (FR-505)
    draft = client.post(f"/api/cases/{case_id}/drafts/quote", json={}).json()
    assert [a["filename"] for a in draft["attachments"]] == [f"{case_id}_견적서.xlsx"]


def test_quote_xlsx_follows_template(client, connect):
    connect(FakeLlm(first=AFTER_REPLY))
    case_id = create(client)["caseId"]
    path = settings.storage_dir / case_id / "quotes" / f"{case_id}_견적서.xlsx"
    ws = load_workbook(path).active
    assert case_id in ws["B12"].value  # FR-102 견적서에 케이스 ID
    assert ws["C17"].value == ": INCHEON - SYDNEY"
    assert ws["C20"].value == ": CIF"
    assert ws["B23"].value == "OCEAN FREIGHT" and ws["C23"].value == 500.0 and ws["D23"].value == 1.0
    assert ws["C35"].value == "USD500.00 + KRW107,203"
    assert len(ws._images) == 1  # 양식 로고 유지


def test_requote_moves_old_file_to_history(client, connect):
    connect(FakeLlm(first=AFTER_REPLY, after_reply=AFTER_REPLY))
    case_id = create(client)["caseId"]
    client.post(f"/api/cases/{case_id}/replies", json=body("03_lcl_reply.eml"))
    quotes = settings.storage_dir / case_id
    assert [p.name for p in (quotes / "quotes").iterdir()] == [f"{case_id}_견적서.xlsx"]
    assert [p.name for p in (quotes / "quote-history").iterdir()] == [f"{case_id}_견적서_v1.xlsx"]


# ------------------------------------------------------------------ 보류·실패


def test_no_rate_means_no_quote(client, connect):
    connect(FakeLlm(first={**AFTER_REPLY, "podName": "뉴욕", "pod": None}))
    case = create(client)
    assert case["status"] == "보류"
    assert case["extraction"]["validation"]["hold"][0]["code"] == "NO_RATE"  # FR-501
    assert not (settings.storage_dir / case["caseId"] / "quotes").exists()


def test_multi_item_and_dangerous_goods_are_held(client, connect):
    connect(FakeLlm(first={**AFTER_REPLY, "multipleItems": True, "commodity": "화장품(향수 포함)"}))
    case = create(client, "02_fob_multi_item.eml")
    codes = {h["code"] for h in case["extraction"]["validation"]["hold"]}
    assert case["status"] == "보류" and codes == {"MULTIPLE_ITEMS", "DANGEROUS_GOODS"}


def test_llm_down_keeps_case_and_marks_failed(client, connect):
    connect(BrokenLlm())
    case = create(client)
    assert case["status"] == "실패"  # NFR-03: 케이스·원문은 남고 다시 실행 가능
    assert "연결하지 못했습니다" in case["extraction"]["error"]
    assert case["mails"][0]["hasRaw"]


def test_past_ready_date_and_over_weight(client, connect):
    connect(FakeLlm(first={**AFTER_REPLY, "cargoReadyDate": "2026-09-01", "grossWeightKg": 30000,
                           "containerType": "40FT HC", "pol": "KRPUS"}))
    case = create(client)
    codes = {i["code"] for i in case["extraction"]["validation"]["issues"]}
    assert {"PAST_DATE", "OVER_WEIGHT"} <= codes


# ------------------------------------------------------------------ 단위 규칙


def test_normalize_codes_and_ports():
    values, notes = normalize({**FIRST, "qtyUnit": "박스", "containerType": "40' HC", "ccy": "달러",
                               "cargoReadyDate": "2026년 10월 20일", "polName": "부산항"})
    assert values["packing"] == "Carton" and values["qtyUnit"] == "CTNS"
    assert values["containerType"] == "40HQ" and values["ccy"] == "USD"
    assert values["cargoReadyDate"] == "2026-10-20"
    assert values["pol"] == "KRPUS" and values["pod"] == "AUSYD"
    assert values["incoterms"] is None and notes[0]["code"] == "MULTIPLE_INCOTERMS"


def test_unmapped_port_gives_candidates():
    values, notes = normalize({**FIRST, "podName": "SIDNEY"})
    assert values["pod"] is None
    assert notes[0]["code"] == "PORT_UNMAPPED" and "AUSYD" in notes[0]["candidates"]


def test_strip_quoted_and_broken_list():
    outlook = "회신드립니다.\r\n  1.\r\n픽업지: 남양주\r\n________________\r\n보낸 사람: A\r\n이전 메일"
    text = join_broken_list(strip_quoted(outlook))
    assert "1. 픽업지: 남양주" in text and "이전 메일" not in text
    gmail = "준비일: 2026-10-20\n\n2026년 9월 29일 (화) 오전 11:46, 박재석 <a@b.c>님이 작성:\n> 이전"
    assert strip_quoted(gmail) == "준비일: 2026-10-20"


def test_llm_must_be_local():
    ensure_local("http://localhost:11434")
    ensure_local("http://192.168.0.10:11434")  # 사내망 허용
    with pytest.raises(LlmError, match="NFR-04"):
        ensure_local("http://8.8.8.8:11434")  # 외부 주소 거부
