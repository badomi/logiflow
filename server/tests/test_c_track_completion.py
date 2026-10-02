from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import validation
from app.config import settings
from app.extraction import parsers
from app.models import QuoteInput
from app.sample_rate_quote import SOURCE, calculate
from app.seed import seed_containers
from .test_sample_rate_quote import request
from .test_outlook_sample_quote import BODY, import_client_rates
from .test_pipeline_flow import connect, create, text_mail


def workbook(tmp_path, sheet, coordinate, value):
    wb = load_workbook(SOURCE)
    wb[sheet][coordinate] = value
    target = tmp_path / "rates.xlsx"
    wb.save(target)
    return target


def test_selected_carrier_schedule():
    result = calculate(SOURCE, request())
    assert result["schedule"] == dict(etd="2026-10-07", eta="2026-10-09", transitTime=2,
                                      freeTimeDem=5, freeTimeDet=5)


def test_ready_after_departure_is_held():
    with pytest.raises(ValueError, match="스케줄"):
        calculate(SOURCE, request(cargoReadyDate="2026-10-08"))


def test_eta_is_derived_when_blank(tmp_path):
    path = workbook(tmp_path, "OceanFreight", "N8", None)
    assert calculate(path, request())["schedule"]["eta"] == "2026-10-09"


@pytest.mark.parametrize("sheet,cell,value,match", [
    ("OceanFreight", "N8", "2026-10-10", "ETA"),
    ("OceanFreight", "O8", -1, "free_time_dem"),
    ("OceanFreight", "P8", 1.5, "free_time_det"),
    ("OceanFreight", "M4", None, "컬럼"),
    ("OceanFreight", "I8", -430, "amount"),
    ("Surcharges", "K14", -0.002, "rate_pct"),
])
def test_invalid_master_blocks_calculation(tmp_path, sheet, cell, value, match):
    with pytest.raises(ValueError, match=match):
        calculate(workbook(tmp_path, sheet, cell, value), request())


def test_excel_native_dates_supported(tmp_path):
    assert calculate(workbook(tmp_path, "OceanFreight", "L8", datetime(2026, 10, 7)), request())["totals"] == {"USD":870}


@pytest.mark.parametrize("field,value", [("grossWeightKg",0), ("grossWeightKg",-500), ("invoiceValue",0),
    ("totalCbm",-1), ("boxL",0), ("boxW",-5), ("boxH",float('inf'))])
def test_invalid_numeric_inputs_ask_for_correction(session, field, value):
    qi = QuoteInput(**{field:value})
    issues = validation.check(session, qi, today=date(2026,10,2))
    assert any(i["field"] == field and i["code"] == "INVALID" and i["severity"] == "Block" for i in issues)


@pytest.mark.parametrize("parser,text", [(parsers.parse_weight_kg,"-500 KG"), (parsers.parse_cbm,"-1 CBM"),
    (parsers.parse_dims_mm,"500 x -400 x 500 mm"), (parsers.parse_money,"USD -3000"),
    (parsers.parse_qty,"-10 CTNS"), (parsers.parse_qty,"1.5 CTNS")])
def test_negative_values_not_silently_changed_to_positive(parser, text):
    assert parser(text) is None


def test_overweight_per_container_and_volume_mismatch(session):
    seed_containers(session)
    session.flush()
    qi = QuoteInput(containerType="20FT GP",grossWeightKg=30000,qty=10,boxL=500,boxW=400,boxH=500,totalCbm=2)
    codes = {i["code"] for i in validation.check(session,qi,today=date(2026,10,2),container_count=1)}
    assert {"OVER_WEIGHT","CBM_MISMATCH"} <= codes


def test_schedule_stored_and_printed_in_app_workbook(client, connect, monkeypatch):
    import_client_rates(client)
    monkeypatch.setattr(validation,"today_kst",lambda:date(2026,10,2))
    case = create(client,text_mail(BODY))
    assert case["status"] == "계산완료"
    quote = case["extraction"]["validation"]["quote"]
    assert quote["validUntil"] == "2026-10-15"
    assert quote["schedule"]["eta"] == "2026-10-09"
    wb = load_workbook(settings.storage_dir/case["caseId"]/"quotes"/f"{case['caseId']}_견적서.xlsx")
    ws = wb.active
    assert ws["B20"].value == "** INCOTERMS"
    assert ws["B35"].value == "TOTAL (USD)"
    assert "2026-10-07" in ws["B42"].value and "2026-10-09" in ws["B42"].value
    assert any("DET 5 DAYS / DEM 5 DAYS" in str(ws.cell(r,2).value) for r in range(39,48))
    assert "T/T 100% advance" in ws["B41"].value
    assert "Calculation Details" in wb.sheetnames


def test_previous_quote_is_blocked_after_invalid_correction(client, connect, monkeypatch):
    import_client_rates(client)
    case = create(client,text_mail(BODY))
    case_id = case["caseId"]
    response = client.post(f"/api/cases/{case_id}/fields/grossWeightKg",json={"value":-1})
    assert response.status_code == 200
    assert client.get(f"/api/cases/{case_id}").json()["status"] == "정보부족"
    assert client.post(f"/api/cases/{case_id}/drafts/quote",json={}).status_code == 400


def test_failed_pdf_regeneration_does_not_attach_previous_pdf(client, connect, monkeypatch):
    from app import quotation
    import_client_rates(client)
    monkeypatch.setattr(settings,"quote_pdf",True)
    monkeypatch.setattr(quotation,"to_pdf",lambda data:(b"%PDF-1.4 mock",None))
    case = create(client,text_mail(BODY))
    case_id = case["caseId"]
    monkeypatch.setattr(quotation,"to_pdf",lambda data:(None,"test failure"))
    client.post(f"/api/cases/{case_id}/extract",json={})
    assert not (settings.storage_dir/case_id/"quotes"/f"{case_id}_견적서.pdf").exists()
    assert (settings.storage_dir/case_id/"quote-history"/f"{case_id}_견적서_v1.pdf").exists()
    assert client.post(f"/api/cases/{case_id}/drafts/quote",json={}).status_code == 400
