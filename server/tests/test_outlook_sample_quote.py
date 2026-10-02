"""Outlook과 같은 API 경로에서 제공 요율표 → 견적 → 첨부 목록까지 검증."""
import io
import os

import pytest
from pypdf import PdfReader

from openpyxl import load_workbook

from app import quotation
from app.config import settings, SERVER_DIR
from app.sample_rate_quote import SOURCE
from .test_pipeline_flow import connect, create, text_mail


def import_client_rates(client):
    from app.db import SessionLocal
    from app.models import Rate
    from sqlalchemy import delete
    # 이 통합 테스트는 제공 요율표만 사용한다 (connect의 예시 요율과 혼합 금지).
    with SessionLocal() as session:
        session.execute(delete(Rate))
        session.commit()
    response = client.post("/api/rates/import", files={"file": (SOURCE.name, SOURCE.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text


BODY = """안녕하세요. 아래 조건으로 견적을 요청드립니다.
품명: 플라스틱 수납함
포장 형태: Carton
수량: 10 CTNS
박스 1개 규격: 500 x 400 x 500 MM
총부피: 1 CBM
총중량: 500 KG
선적항: CNSHA
도착항: KRPUS
컨테이너: 20FT GP 1대
거래조건: CIF
화물 준비일: 2026-10-05
인보이스 금액: USD 15,000
결제 조건: T/T 100% advance
감사합니다.
김테스트
테스트상사
"""


def test_api_builds_workbook_quote_and_two_attachments(client, connect, monkeypatch):
    import_client_rates(client)
    monkeypatch.setattr(settings, "quote_pdf", True)
    monkeypatch.setattr(quotation, "to_pdf", lambda data: (b"%PDF-1.4\nmock", None))
    case = create(client, text_mail(BODY))
    assert case["status"] == "계산완료", case
    quote = case["extraction"]["validation"]["quote"]
    assert quote["totals"] == {"USD": 870.0}
    assert quote["carrier"] == "SITC"
    assert quote["lines"][0]["source"].startswith("OceanFreight#")
    assert quote["file"] and quote["pdf"]
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/quote", json={}).json()
    assert {a["filename"].split(".")[-1] for a in draft["attachments"]} == {"pdf", "xlsx"}
    for attachment in draft["attachments"]:
        response = client.get(attachment["url"])
        assert response.status_code == 200 and response.content
        if attachment["filename"].endswith("xlsx"):
            ws = load_workbook(io.BytesIO(response.content)).active
            assert case["caseId"] in ws["B12"].value
            assert ws["C35"].value == "USD870.00"
            assert "Calculation Details" in load_workbook(io.BytesIO(response.content)).sheetnames


def test_missing_lane_is_held_instead_of_using_other_rates(client, connect, monkeypatch):
    import_client_rates(client)
    case = create(client, text_mail(BODY.replace("CNSHA", "KRPUS").replace("도착항: KRPUS", "도착항: SGSIN")))
    assert case["status"] == "보류"
    assert case["extraction"]["validation"]["quote"] is None
    assert "요율" in case["extraction"]["validation"]["hold"][0]["message"]


@pytest.mark.skipif(os.environ.get("LEONA_EXCEL_TEST") != "1", reason="Windows Excel 실변환은 명시 실행")
def test_real_excel_pdf_through_case_api(client, connect, monkeypatch):
    import_client_rates(client)
    monkeypatch.setattr(settings, "quote_pdf", True)
    monkeypatch.setattr(settings, "pdf_backend", "excel")
    case = create(client, text_mail(BODY))
    quote = case["extraction"]["validation"]["quote"]
    assert case["status"] == "계산완료" and quote["pdf"], case
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/quote", json={}).json()
    assert len(draft["attachments"]) == 2
    preview = SERVER_DIR.parent / "output" / "quote-preview"
    preview.mkdir(parents=True, exist_ok=True)
    for attachment in draft["attachments"]:
        response = client.get(attachment["url"])
        assert response.status_code == 200
        if attachment["filename"].endswith("pdf"):
            doc = PdfReader(io.BytesIO(response.content))
            assert len(doc.pages) == 1
            assert "870.00" in doc.pages[0].extract_text()
            (preview / "api_quote.pdf").write_bytes(response.content)
        else:
            (preview / "api_quote.xlsx").write_bytes(response.content)
