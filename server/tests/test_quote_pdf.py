"""WBS 4.9 견적서 PDF — 운영체제별 LibreOffice 탐색·변환 실패·내장 엔진 대체 검증."""

from io import BytesIO
from pathlib import Path
import subprocess

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject, NumberObject

from app import quotation, quote_pdf
from app.config import settings
from app.quote_pdf import render_pdf
from app.quote_template import save_filled_template


def filled_template() -> bytes:
    """render_xlsx와 같은 칸에 값을 채운 양식 (DB 없이)."""
    wb = load_workbook(settings.quote_template)
    ws = wb.active
    ws["B9"] = "수 신 : 한빛무역 / 김민수 님"
    ws["B12"] = "Subject : 해상 수출 운임 제안서  [LQ-2026-1001-001]"
    ws["C17"], ws["C18"], ws["C20"] = ": BUSAN - SINGAPORE", ": BY CONTAINER SHIPMENT (40HQ X 1)", ": CIF"
    ws["F21"] = "2026-10-01"
    for r in range(23, 35):
        for c in "CDEF":
            ws[f"{c}{r}"] = None
    ws["B23"], ws["C23"], ws["D23"], ws["E23"], ws["F23"] = "OCEAN FREIGHT", 750.0, 1, 750.0, "PER CNTR"
    ws["C23"].number_format = ws["E23"].number_format = '"US$"#,##0.00'
    ws["B24"], ws["C24"], ws["D24"], ws["E24"], ws["F24"] = "THC", 410000, 1, 410000, "PER CNTR"
    ws["C24"].number_format = ws["E24"].number_format = '"₩"#,##0'
    ws["E33"], ws["F33"] = "AT COST", "INV.V x 1/1,000"
    ws["C35"] = "USD750.00 + KRW410,000"
    ws["B39"] = "* 실제 출항일 환율 적용"
    return save_filled_template(wb, settings.quote_template)


def test_render_pdf_from_filled_template():
    pdf = render_pdf(filled_template())
    assert pdf.startswith(b"%PDF") and len(pdf) > 5_000


def test_to_pdf_falls_back_without_libreoffice(monkeypatch):
    monkeypatch.setattr(quotation, "find_soffice", lambda: None)
    pdf, note = quotation.to_pdf(filled_template())
    assert pdf is not None and pdf.startswith(b"%PDF")
    assert "내장 엔진" in note


def test_to_pdf_reports_failure_if_builtin_fails(monkeypatch):
    monkeypatch.setattr(quotation, "find_soffice", lambda: None)
    pdf, note = quotation.to_pdf(b"not an xlsx")
    assert pdf is None and "PDF·XLSX 송부 보류" in note


@pytest.fixture
def no_installed_soffice(monkeypatch):
    """실제 PC 설치 상태와 개인 .env에 영향을 받지 않는 탐색 테스트."""
    monkeypatch.setattr(settings, "soffice_path", None)
    monkeypatch.setattr(quotation.shutil, "which", lambda name: None)
    monkeypatch.setattr(Path, "is_file", lambda path: False)


def test_configured_soffice_path_takes_priority(monkeypatch, tmp_path):
    executable = tmp_path / "한글 설치 경로" / "soffice.exe"
    executable.parent.mkdir()
    executable.write_bytes(b"test executable")
    monkeypatch.setattr(settings, "soffice_path", str(executable))
    monkeypatch.setattr(quotation.shutil, "which", lambda name: "different executable")
    assert quotation.find_soffice() == str(executable)


@pytest.mark.parametrize("directory_instead", [False, True])
def test_invalid_configured_path_does_not_select_another_installation(monkeypatch, tmp_path, directory_instead):
    configured = tmp_path if directory_instead else tmp_path / "missing-soffice"
    monkeypatch.setattr(settings, "soffice_path", str(configured))
    monkeypatch.setattr(quotation.shutil, "which", lambda name: "different executable")
    assert quotation.find_soffice() is None


@pytest.mark.parametrize("platform", ["win32", "darwin", "linux"])
@pytest.mark.parametrize("command", ["soffice", "libreoffice"])
def test_path_lookup_still_works_on_each_platform(no_installed_soffice, monkeypatch, platform, command):
    monkeypatch.setattr(quotation.sys, "platform", platform)
    monkeypatch.setattr(quotation.shutil, "which", lambda name: "path executable" if name == command else None)
    assert quotation.find_soffice() == "path executable"


@pytest.mark.parametrize("available", [0, 1])
def test_windows_default_installations_remain_supported(no_installed_soffice, monkeypatch, available):
    monkeypatch.setattr(quotation.sys, "platform", "win32")
    expected = Path(quotation._WINDOWS_SOFFICE[available])
    monkeypatch.setattr(Path, "is_file", lambda path: path == expected)
    assert quotation.find_soffice() == str(expected)


def test_windows_keeps_64_bit_installation_priority(no_installed_soffice, monkeypatch):
    monkeypatch.setattr(quotation.sys, "platform", "win32")
    monkeypatch.setattr(Path, "is_file", lambda path: path in map(Path, quotation._WINDOWS_SOFFICE))
    assert quotation.find_soffice() == str(Path(quotation._WINDOWS_SOFFICE[0]))


@pytest.mark.parametrize("user_installation", [False, True])
def test_macos_application_bundle_locations(no_installed_soffice, monkeypatch, user_installation):
    monkeypatch.setattr(quotation.sys, "platform", "darwin")
    expected = Path(quotation._MACOS_SOFFICE[int(user_installation)]).expanduser()
    monkeypatch.setattr(Path, "is_file", lambda path: path == expected)
    assert quotation.find_soffice() == str(expected)


def test_macos_prefers_shared_applications_installation(no_installed_soffice, monkeypatch):
    monkeypatch.setattr(quotation.sys, "platform", "darwin")
    installations = [Path(p).expanduser() for p in quotation._MACOS_SOFFICE]
    monkeypatch.setattr(Path, "is_file", lambda path: path in installations)
    assert quotation.find_soffice() == str(installations[0])


@pytest.mark.parametrize("platform", ["win32", "darwin"])
def test_missing_installation_returns_none(no_installed_soffice, monkeypatch, platform):
    monkeypatch.setattr(quotation.sys, "platform", platform)
    assert quotation.find_soffice() is None


def sample_pdf(kind="valid"):
    writer = PdfWriter()
    if kind != "no_pages":
        page = writer.add_blank_page(width=595, height=842)
        stream = DecodedStreamObject()
        stream.set_data(b"q 1 0 0 1 0 0 cm Q")
        page[NameObject("/Contents")] = writer._add_object(stream)
        if kind == "broken_contents":
            page[NameObject("/Contents")] = NumberObject(123)
    if kind == "encrypted":
        writer.encrypt("test-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def fake_conversion(monkeypatch, output, returncode=0):
    monkeypatch.setattr(quotation, "find_soffice", lambda: "test soffice")

    def run(command, **kwargs):
        if output is not None:
            folder = Path(command[command.index("--outdir") + 1])
            (folder / "quote.pdf").write_bytes(output)
        return subprocess.CompletedProcess(command, returncode, b"private conversion output", b"private error output")

    monkeypatch.setattr(quotation.subprocess, "run", run)


def test_successful_conversion_returns_a_readable_pdf(monkeypatch):
    expected = sample_pdf()
    fake_conversion(monkeypatch, expected)
    pdf, note = quotation._soffice_pdf(b"test xlsx")
    assert pdf == expected and note is None


def test_nonzero_exit_rejects_even_an_existing_pdf(monkeypatch):
    fake_conversion(monkeypatch, sample_pdf(), returncode=7)
    pdf, note = quotation._soffice_pdf(b"test xlsx")
    assert pdf is None and "종료 코드 7" in note
    assert "private" not in note


@pytest.mark.parametrize("kind", ["missing", "empty", "not_pdf", "truncated", "no_pages", "encrypted", "broken_contents"])
def test_success_exit_rejects_missing_or_unreadable_pdf(monkeypatch, kind):
    outputs = {"missing": None, "empty": b"", "not_pdf": b"not a PDF", "truncated": b"%PDF-1.7\nbroken"}
    output = outputs[kind] if kind in outputs else sample_pdf(kind)
    fake_conversion(monkeypatch, output)
    pdf, note = quotation._soffice_pdf(b"test xlsx")
    assert pdf is None
    assert ("변환 결과 없음" if kind == "missing" else "PDF 오류") in note


@pytest.mark.parametrize("failure", ["timeout", "permission", "not_found"])
def test_conversion_process_errors_are_reported(monkeypatch, failure):
    monkeypatch.setattr(quotation, "find_soffice", lambda: "test soffice")
    monkeypatch.setattr(settings, "pdf_timeout_s", 2.5)

    def run(command, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired(command, kwargs["timeout"])
        raise (PermissionError if failure == "permission" else FileNotFoundError)("private path")

    monkeypatch.setattr(quotation.subprocess, "run", run)
    pdf, note = quotation._soffice_pdf(b"test xlsx")
    assert pdf is None and ("시간 초과 (2.5초)" if failure == "timeout" else "실행·파일 처리 실패") in note
    assert "private" not in note


def test_pdf_read_error_is_reported(monkeypatch):
    fake_conversion(monkeypatch, sample_pdf())
    original = Path.read_bytes

    def read_bytes(path):
        if path.name == "quote.pdf":
            raise PermissionError("private path")
        return original(path)

    monkeypatch.setattr(Path, "read_bytes", read_bytes)
    pdf, note = quotation._soffice_pdf(b"test xlsx")
    assert pdf is None and "실행·파일 처리 실패" in note
    assert "private" not in note


def test_corrupt_libreoffice_result_uses_valid_builtin_pdf(monkeypatch):
    fake_conversion(monkeypatch, b"%PDF-1.7\nbroken")
    pdf, note = quotation.to_pdf(filled_template())
    assert len(PdfReader(BytesIO(pdf)).pages) == 1
    assert "PDF 오류" in note and "내장 엔진" in note


def test_corrupt_builtin_result_is_not_sent(monkeypatch):
    fake_conversion(monkeypatch, None)
    monkeypatch.setattr(quote_pdf, "render_pdf", lambda xlsx: b"%PDF-1.7\nbroken")
    pdf, note = quotation.to_pdf(b"test xlsx")
    assert pdf is None and "PDF·XLSX 송부 보류" in note


def test_installed_libreoffice_preserves_quote_content(tmp_path):
    """실제 변환만 통과시킨다. 내장 엔진으로 대체되면 성공으로 취급하지 않는다."""
    if quotation.find_soffice() is None:
        pytest.skip("LibreOffice 미설치 — 실제 변환 검증 제외")
    xlsx = filled_template()
    pdf, note = quotation.to_pdf(xlsx)
    assert pdf is not None and note is None, note
    document = PdfReader(BytesIO(pdf))
    assert len(document.pages) == 1
    text = document.pages[0].extract_text()
    for expected in ("LQ-2026-1001-001", "한빛무역", "USD750.00 + KRW410,000", "OCEAN FREIGHT", "EXPORT SERVICE CHARGE"):
        assert expected in text
    # 실제 결과를 눈으로 비교할 수 있게 pytest 임시 경로에만 보관한다.
    (tmp_path / "quote.xlsx").write_bytes(xlsx)
    (tmp_path / "quote.pdf").write_bytes(pdf)
