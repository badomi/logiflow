"""FR-507: 임시 DB에서 필수 견적 정보 누락·보완·재생성을 검증한다."""

from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app import quotation
from app.config import settings
from app.models import LaneRule, Quote, Rate
from .test_pipeline_flow import CLIENT_CASE, FCL_SG, client_rates, connect, create, quote_folder, text_mail


BODY = CLIENT_CASE.format(carrier="선사: HMM\n").replace("2026-10-05", "2026-10-20")
EXPORT_FOB = FCL_SG.replace("조건: CIF", "조건: FOB")


def export_freight_rows(session):
    return session.scalars(select(Rate).where(
        Rate.charge_code == "OCEAN_FREIGHT", Rate.pol == "KRPUS",
        Rate.pod == "SGSIN", Rate.container_type == "40HQ",
    )).all()


@pytest.mark.parametrize("field,value,expected", [
    ("pol", "CNSHA", "etd"),
    ("pod", "AUSYD", "etd"),
    ("container_type", "20FT GP", "etd"),
    ("valid_from", date(2026, 11, 6), "etd"),
    ("valid_until", date(2026, 11, 4), "etd"),
    ("etd", None, "etd"),
    ("etd", "2026-11-04", "etd"),
    ("eta", "2026-11-05", "eta"),
])
def test_export_fob_requires_matching_valid_schedule(client, connect, session, field, value, expected):
    rows = export_freight_rows(session)
    assert rows and all(row.incoterms == "CIF" for row in rows)
    for row in rows:
        setattr(row, field, value)
    session.commit()
    case = create(client, text_mail(EXPORT_FOB))
    result = case["extraction"]["validation"]
    assert case["status"] == "보류" and result["quote"] is None
    assert result["questions"] == []
    assert result["hold"][0]["code"] == "QUOTE_INFO_INCOMPLETE"
    assert expected in result["hold"][0]["fields"]
    assert quote_folder(case["caseId"]) == []
    assert session.scalar(select(func.count()).select_from(Quote)) == 0


@pytest.mark.parametrize("wanted,selected,etd,eta,det,dem", [
    ("HMM", "HMM", "2026-11-07", "2026-11-10", 8, 9),
    ("", "SITC", "2026-11-06", "2026-11-08", 5, 6),
    ("COSCO", "SITC", "2026-11-06", "2026-11-08", 5, 6),
])
def test_export_fob_schedule_obeys_carrier_selection_without_billing_freight(
    client, connect, session, wanted, selected, etd, eta, det, dem,
):
    for row in export_freight_rows(session):
        row.carrier, row.unit_price = "SITC", 750
        row.transit_days, row.free_time_det, row.free_time_dem = 2, 5, 6
    session.add(Rate(
        charge_code="OCEAN_FREIGHT", charge_label="Ocean Freight", basis="PER_CNTR",
        pol="KRPUS", pod="SGSIN", container_type="40HQ", incoterms="CIF",
        currency="USD", unit_price=900, carrier="HMM", etd="2026-11-07", eta=None,
        transit_days=3, free_time_det=8, free_time_dem=9,
    ))
    session.commit()
    body = EXPORT_FOB.replace("감사합니다.", f"선사: {wanted}\n감사합니다.") if wanted else EXPORT_FOB
    case = create(client, text_mail(body))
    assert case["status"] == "계산완료", case["extraction"]["validation"]
    quote = case["extraction"]["validation"]["quote"]
    assert (quote["carrier"], quote["etd"], quote["eta"]) == (selected, etd, eta)
    assert quote["totals"] == {"KRW": 498220.0}
    ws = load_workbook(settings.storage_dir / quote["file"]).active
    remarks = "\n".join(str(ws[f"B{row}"].value or "") for row in range(39, 49))
    assert f"DET {det} days / DEM {dem} days" in remarks
    assert ws["E23"].value is None
    assert len(quote_folder(case["caseId"])) == 2


def test_export_fob_ambiguous_carrier_currency_holds_for_review(client, connect, session):
    session.add(Rate(
        charge_code="OCEAN_FREIGHT", charge_label="Ocean Freight", basis="PER_CNTR",
        pol="KRPUS", pod="SGSIN", container_type="40HQ", incoterms="CIF",
        currency="KRW", unit_price=900000, carrier="HMM", etd="2026-11-07", eta="2026-11-09",
        transit_days=2, free_time_det=7, free_time_dem=7,
    ))
    session.commit()
    case = create(client, text_mail(EXPORT_FOB))
    result = case["extraction"]["validation"]
    assert case["status"] == "보류" and result["quote"] is None
    assert result["hold"][0]["code"] == "QUOTE_INFO_INCOMPLETE"
    assert result["hold"][0]["fields"] == ["carrier"]
    assert quote_folder(case["caseId"]) == []


def freight_rows(session):
    return session.scalars(select(Rate).where(
        Rate.charge_code == "OCEAN_FREIGHT", Rate.carrier == "HMM",
        Rate.pol == "CNSHA", Rate.pod == "KRPUS", Rate.container_type == "20FT GP",
    )).all()


@pytest.mark.parametrize("field,value,expected", [
    ("transit_days", None, "transitTime"),
    ("free_time_det", None, "freeTimeDet"),
    ("free_time_dem", None, "freeTimeDem"),
    ("etd", None, "etd"),
    ("etd", "invalid", "etd"),
    ("eta", "invalid", "eta"),
    ("transit_days", -1, "transitTime"),
    ("free_time_det", -1, "freeTimeDet"),
    ("free_time_dem", -1, "freeTimeDem"),
    ("etd", "2026-10-19", "etd"),
    ("eta", "2026-10-19", "eta"),
])
def test_missing_or_invalid_master_info_holds_before_file_generation(
    client, client_rates, session, monkeypatch, field, value, expected,
):
    rows = freight_rows(session)
    assert rows
    for row in rows:
        setattr(row, field, value)
    session.commit()

    def unexpected_render(*args, **kwargs):
        pytest.fail("필수 정보 확인 전에 견적서 생성이 호출됨")

    monkeypatch.setattr(quotation, "render_xlsx", unexpected_render)
    case = create(client, text_mail(BODY))
    result = case["extraction"]["validation"]
    assert case["status"] == "보류" and result["quote"] is None
    assert result["questions"] == []  # 요율 관리 정보는 화주 보완 문항으로 보내지 않는다.
    assert result["hold"][0]["code"] == "QUOTE_INFO_INCOMPLETE"
    assert result["hold"][0]["severity"] == "Block"
    assert expected in result["hold"][0]["fields"]
    assert "FR-507" in result["hold"][0]["message"]
    assert result["severityCount"]["Block"] == 1
    assert session.scalar(select(func.count()).select_from(Quote)) == 0
    assert quote_folder(case["caseId"]) == []
    assert client.post(f"/api/cases/{case['caseId']}/drafts/quote", json={}).status_code == 400


@pytest.mark.parametrize("days", [None, 0, -1, 1.5])
def test_invalid_validity_days_are_not_used(days):
    qi = SimpleNamespace(transitTime=2, freeTimeDet=7, freeTimeDem=7, validityDays=days)
    freight = SimpleNamespace(etd="2026-10-20", eta="2026-10-22")
    with pytest.raises(quotation.QuoteInformationError) as error:
        quotation.required_schedule(qi, freight, date(2026, 10, 20), date(2026, 10, 15), [])
    assert error.value.fields == ["validityDays"]


def test_lane_rule_fallback_and_derived_eta_make_complete_quote(client, client_rates, session):
    for row in freight_rows(session):
        row.transit_days = row.free_time_det = row.free_time_dem = row.eta = None
    session.add(LaneRule(pol="CNSHA", pod="KRPUS", container_type="20FT GP",
                         transit_days=2, free_time_det=0, free_time_dem=9, validity_days=7))
    session.commit()
    case = create(client, text_mail(BODY))
    result = case["extraction"]["validation"]
    assert case["status"] == "계산완료" and result["hold"] == []
    quote = result["quote"]
    assert (quote["etd"], quote["eta"], quote["validUntil"]) == ("2026-10-20", "2026-10-22", "2026-10-06")
    assert quote["totals"] == {"USD": 963.0}
    ws = load_workbook(settings.storage_dir / quote["file"]).active
    remarks = "\n".join(ws[f"B{row}"].value or "" for row in range(39, 48))
    assert "ETD / ETA : 2026-10-20 / 2026-10-22" in remarks
    assert "DET 0 days / DEM 9 days" in remarks and "to be confirmed" not in remarks
    assert (settings.storage_dir / quote["pdf"]).read_bytes().startswith(b"%PDF")


def test_missing_info_retires_previous_files_and_recovery_makes_both_formats(client, client_rates, session):
    case = create(client, text_mail(BODY))
    case_id = case["caseId"]
    first = case["extraction"]["validation"]["quote"]
    assert first["version"] == 1 and len(quote_folder(case_id)) == 2
    rows = freight_rows(session)
    originals = [(row, row.free_time_det) for row in rows]
    for row in rows:
        row.free_time_det = None
    session.commit()
    client.post(f"/api/cases/{case_id}/extract", json={})
    held = client.get(f"/api/cases/{case_id}").json()
    result = held["extraction"]["validation"]
    assert held["status"] == "보류" and result["quote"] is None
    assert result["hold"][0]["code"] == "QUOTE_INFO_INCOMPLETE"
    assert len(result["quoteRetired"]) == 2 and quote_folder(case_id) == []
    assert len(quote_folder(case_id, "quote-history")) == 2
    assert client.post(f"/api/cases/{case_id}/drafts/quote", json={}).status_code == 400
    assert session.scalar(select(func.count()).select_from(Quote)) == 1  # 실패한 버전은 없음

    for row, value in originals:
        row.free_time_det = value
    session.commit()
    client.post(f"/api/cases/{case_id}/extract", json={})
    recovered = client.get(f"/api/cases/{case_id}").json()
    quote = recovered["extraction"]["validation"]["quote"]
    assert recovered["status"] == "계산완료" and quote["version"] == 2
    assert quote["totals"] == first["totals"]
    assert len(quote_folder(case_id)) == 2
    draft = client.post(f"/api/cases/{case_id}/drafts/quote", json={}).json()
    assert sorted(a["filename"] for a in draft["attachments"]) == [f"{case_id}_견적서.pdf", f"{case_id}_견적서.xlsx"]
