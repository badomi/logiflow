"""견적마다 1~14일 선택, 재발행 및 기존 버전 보존."""
import io
import json
import os
from datetime import date

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app import validation
from app.config import settings, SERVER_DIR
from app.db import SessionLocal
from app.models import Case, CaseEvent
from .test_outlook_sample_quote import BODY, import_client_rates
from .test_pipeline_flow import connect, create, text_mail


@pytest.fixture
def case(client, connect, monkeypatch):
    monkeypatch.setattr(validation, "today_kst", lambda: date(2026, 10, 2))
    import_client_rates(client)
    return create(client, text_mail(BODY))


@pytest.mark.parametrize("days,until", [(1,"2026-10-03"),(3,"2026-10-05"),(7,"2026-10-09"),(14,"2026-10-15")])
def test_selected_days_persist_and_version_is_preserved(client, case, days, until):
    cid = case["caseId"]
    response = client.post(f"/api/cases/{cid}/quote-validity", json={"validityDays": days, "actor": "tester"})
    assert response.status_code == 200
    detail = client.get(f"/api/cases/{cid}").json()
    quote = detail["extraction"]["validation"]["quote"]
    assert detail["validityDays"] == quote["validityDays"] == days
    assert quote["validUntil"] == until and quote["version"] == 2
    assert quote["validityLimitedByRate"] == (days == 14)
    events = [e for e in detail["events"] if e["eventType"] == "QUOTE_VALIDITY_CHANGED"]
    assert events[-1]["actor"] == "tester" and events[-1]["detail"] == {"before": 14, "after": days}
    assert (settings.storage_dir/cid/"quote-history"/f"{cid}_견적서_v1.xlsx").exists()
    wb = load_workbook(settings.storage_dir/quote["file"])
    remarks = " ".join(str(wb.active.cell(r,2).value) for r in range(39,48))
    assert until in remarks and f"selected: {days} days" in remarks
    client.post(f"/api/cases/{cid}/extract", json={})
    assert client.get(f"/api/cases/{cid}").json()["validityDays"] == days


@pytest.mark.parametrize("value", [0,15,-1,1.5,"7",True,None])
def test_invalid_days_do_not_change_case(client, case, value):
    cid = case["caseId"]
    before = client.get(f"/api/cases/{cid}").json()
    response = client.post(f"/api/cases/{cid}/quote-validity", json={"validityDays": value})
    assert response.status_code == 422
    after = client.get(f"/api/cases/{cid}").json()
    assert after["events"] == before["events"] and after["validityDays"] == before["validityDays"]


def test_days_are_independent_per_case(client, case):
    other = create(client, text_mail(BODY))
    client.post(f"/api/cases/{case['caseId']}/quote-validity", json={"validityDays":3})
    assert client.get(f"/api/cases/{other['caseId']}").json()["validityDays"] == 14


def test_running_pipeline_prevents_change(client, case):
    cid = case["caseId"]
    with SessionLocal() as session:
        record = session.scalar(select(Case).where(Case.case_id == cid))
        session.add(CaseEvent(case=record, event_type="PIPELINE_STARTED", detail=json.dumps({"trigger":"MANUAL_RERUN"})))
        session.commit()
    response = client.post(f"/api/cases/{cid}/quote-validity", json={"validityDays":7})
    assert response.status_code == 409
    assert client.get(f"/api/cases/{cid}").json()["validityDays"] == 14


@pytest.mark.skipif(os.environ.get("LEONA_EXCEL_TEST") != "1", reason="Excel 실변환 명시 실행")
def test_selected_date_matches_pdf_and_excel(client, case, monkeypatch):
    monkeypatch.setattr(settings,"quote_pdf",True)
    monkeypatch.setattr(settings,"pdf_backend","excel")
    cid=case["caseId"]
    assert client.post(f"/api/cases/{cid}/quote-validity",json={"validityDays":7}).status_code == 200
    detail=client.get(f"/api/cases/{cid}").json()
    q=detail["extraction"]["validation"]["quote"]
    assert q["validUntil"] == "2026-10-09" and q["pdf"]
    pdf=PdfReader(settings.storage_dir/q["pdf"])
    assert "VALID UNTIL: 2026-10-09" in pdf.pages[0].extract_text()
    assert "selected: 7 days" in pdf.pages[0].extract_text()
    draft=client.post(f"/api/cases/{cid}/drafts/quote",json={}).json()
    assert len(draft["attachments"]) == 2
    wb=load_workbook(settings.storage_dir/q["file"])
    assert any("2026-10-09" in str(wb.active.cell(r,2).value) for r in range(39,48))
    preview = SERVER_DIR.parent / "output" / "quote-preview"
    preview.mkdir(parents=True, exist_ok=True)
    (preview / "validity-7-days.pdf").write_bytes((settings.storage_dir/q["pdf"]).read_bytes())
