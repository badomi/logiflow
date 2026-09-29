"""요율 엑셀 올리기·내려받기 (NFR-06)."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from app import rates
from app.db import SessionLocal
from app.models import LaneRule, Rate
from app.seed import seed_containers, seed_sample_rates


@pytest.fixture
def db():
    with SessionLocal() as s:
        seed_containers(s)
        s.commit()
        yield s


def workbook(rows, lanes=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = rates.RATE_SHEET
    ws.append(rates.RATE_COLUMNS)
    for r in rows:
        ws.append(r)
    if lanes is not None:
        ls = wb.create_sheet(rates.LANE_SHEET)
        ls.append(rates.LANE_COLUMNS)
        for r in lanes:
            ls.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


OK = ["OCEAN_FREIGHT", "OCEAN FREIGHT", "부산", "SINGAPORE", "40' HC", "CIF", "CNTR", "750", "USD", "", "2026-10-01",
      "2026-10-31", "2026-10 D선사"]


def test_template_has_guide_and_example(db):
    wb = load_workbook(io.BytesIO(rates.template()))
    assert wb.sheetnames == [rates.RATE_SHEET, rates.LANE_SHEET, rates.GUIDE_SHEET]
    assert [c.value for c in wb[rates.RATE_SHEET][1]] == rates.RATE_COLUMNS


def test_import_normalizes_names_and_replaces_same_source(db):
    result = rates.import_workbook(db, workbook([OK, ["THC", "", "KRPUS", "", "40HQ", "", "컨테이너", "410,000", "", "",
                                                      "", "", "2026-10 D선사"]]))
    assert result["rates"] == 2 and result["replacedSources"] == ["2026-10 D선사"]
    row = db.query(Rate).filter_by(charge_code="OCEAN_FREIGHT").one()
    assert (row.pol, row.pod, row.container_type, row.basis, row.unit_price) == ("KRPUS", "SGSIN", "40HQ", "PER_CNTR", 750)
    assert db.query(Rate).filter_by(charge_code="THC").one().currency == "KRW"  # 통화 비우면 KRW

    rates.import_workbook(db, workbook([[*OK[:7], "800", *OK[8:]]]))  # 같은 출처 다시 → 통째로 교체
    assert [r.unit_price for r in db.query(Rate).all()] == [800]


def test_bad_rows_change_nothing_and_say_where(db):
    seed_sample_rates(db)
    db.commit()
    before = db.query(Rate).count()
    bad = [["OCEAN_FREIGHT", "", "SIDNEY", "", "53FT", "FCA", "개당", "", "US", "", "10/01", "", "x"]]
    with pytest.raises(rates.RateImportError) as err:
        rates.import_workbook(db, workbook([OK, *bad]))
    text = "\n".join(err.value.errors)
    for expected in ("요율 3행", "AUSYD", "컨테이너 마스터", "인코텀즈", "기준", "통화", "날짜"):
        assert expected in text, expected
    assert db.query(Rate).count() == before  # 하나라도 틀리면 반영 안 함


def test_lane_rules_sheet_replaces_rules(db):
    rates.import_workbook(db, workbook([OK], lanes=[["KRPUS", "SGSIN", "40HQ", "USD", 5, 7, 7, 10]]))
    rule = db.query(LaneRule).one()
    assert (rule.transit_days, rule.validity_days) == (5, 10)


def test_export_round_trip(db):
    seed_sample_rates(db)
    db.commit()
    count = db.query(Rate).count()
    data = rates.export(db)
    assert rates.import_workbook(db, data)["rates"] == count and db.query(Rate).count() == count


def test_api_endpoints(client, db):
    assert client.get("/api/rates/template").headers["content-type"].startswith(rates.XLSX_MEDIA if hasattr(rates, "XLSX_MEDIA") else "application/vnd")
    res = client.post("/api/rates/import", files={"file": ("r.xlsx", workbook([OK]), "application/octet-stream")})
    assert res.status_code == 200 and res.json()["rates"] == 1
    bad = client.post("/api/rates/import", files={"file": ("r.xlsx", b"not excel", "application/octet-stream")})
    assert bad.status_code == 400 and bad.json()["detail"]["errors"]
    assert client.get("/api/rates/summary").json()[0]["source"] == "2026-10 D선사"


def test_health_detail_tells_what_to_fix(client, db, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "extraction_mode", "off")
    info = client.get("/api/health/detail").json()
    joined = "\n".join(info["warnings"])
    assert not info["ok"] and "EXTRACTION_MODE=off" in joined and "요율이 0개" in joined

    monkeypatch.setattr(settings, "extraction_mode", "hybrid")
    monkeypatch.setattr(settings, "llm_url", "http://127.0.0.1:9")  # 아무것도 없는 포트
    info = client.get("/api/health/detail").json()
    assert info["llm"]["reachable"] is False and any("Ollama" in w for w in info["warnings"])
