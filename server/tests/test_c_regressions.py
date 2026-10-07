"""C트랙 회귀 검증: 원본 DB·업체 양식 대신 conftest의 임시 저장소만 사용한다."""

import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app import quotation, rates, services, validation
from app.models import Rate
from app.seed import seed_containers
from .test_pipeline_flow import FULL_LCL, TODAY, connect, create, text_mail


CLIENT_FILE = Path(__file__).resolve().parents[2] / "샘플_요율표.xlsx"


def save(wb):
    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()


def rate_values(session):
    return [{col.name: getattr(row, col.name) for col in Rate.__table__.columns if col.name != "id"}
            for row in session.scalars(select(Rate).order_by(Rate.id))]


def test_client_export_preserves_every_rate_field_and_quote(session, client, connect):
    rates.import_workbook(session, CLIENT_FILE.read_bytes())
    before = rate_values(session)
    body = ("품목: 부품\n수량: 10박스\n전체 부피: 5CBM\n총중량: 500kg\n선적항: 상하이\n도착항: 부산\n"
            "화물 준비일: 2026-10-20\n조건: CIF\n컨테이너: 20GP 1대\n인보이스 금액: USD 15000\n결제조건: T/T Advance")
    first = create(client, text_mail(body))["extraction"]["validation"]["quote"]
    exported = rates.export(session)
    result = rates.import_workbook(session, exported)
    session.expire_all()
    assert result["rates"] == len(before)
    assert rate_values(session) == before
    second = create(client, text_mail(body))["extraction"]["validation"]["quote"]
    for key in ("totals", "carrier", "etd", "eta", "lines", "rateSources", "validUntil"):
        assert first[key] == second[key], key


def test_export_preserves_shipment_basis_and_source_priority(session):
    seed_containers(session)
    for source, amount in (("Z old", 30), ("A new", 40)):
        session.add(Rate(charge_code="DOC", charge_label="DOC", basis="PER_SHIPMENT", unit_price=amount,
                         currency="USD", source=source, sort_no=25))
    session.commit()
    before = rate_values(session)
    assert quotation.one_table_per_charge(list(session.scalars(select(Rate))))[0].source == "A new"
    rates.import_workbook(session, rates.export(session))
    session.expire_all()
    assert rate_values(session) == before
    assert quotation.one_table_per_charge(list(session.scalars(select(Rate))))[0].source == "A new"


@pytest.mark.parametrize("value", [-1, "NaN", "Infinity", "-Infinity"])
@pytest.mark.parametrize("client_format", [True, False])
def test_invalid_rate_amount_is_rejected_without_replacing_data(session, value, client_format):
    seed_containers(session)
    rates.import_workbook(session, CLIENT_FILE.read_bytes())
    before = rate_values(session)
    wb = load_workbook(CLIENT_FILE) if client_format else load_workbook(io.BytesIO(rates.export(session)))
    if client_format:
        wb["OceanFreight"]["I5"] = value
    else:
        ws = wb["요율"]
        column = [c.value for c in ws[1]].index("단가") + 1
        ws.cell(2, column, value)
    with pytest.raises(rates.RateImportError, match="틀린 곳"):
        rates.import_workbook(session, save(wb))
    assert rate_values(session) == before


@pytest.mark.parametrize("cell,value", [("M5", -2), ("M5", 1.5), ("O5", "NaN"), ("N5", "2026-10-01")])
def test_invalid_schedule_is_reported_as_import_error(session, cell, value):
    seed_containers(session)
    session.commit()
    wb = load_workbook(CLIENT_FILE)
    wb["OceanFreight"][cell] = value
    with pytest.raises(rates.RateImportError):
        rates.import_workbook(session, save(wb))
    assert rate_values(session) == []


def test_old_management_format_is_still_accepted(session):
    seed_containers(session)
    wb = Workbook()
    ws = wb.active
    ws.title = "요율"
    ws.append(rates.RATE_COLUMNS)
    ws.append(["THC", "THC", "KRPUS", "", "40HQ", "", "CNTR", 0, "USD", "", "", "", "legacy"])
    assert rates.import_workbook(session, save(wb))["rates"] == 1


@pytest.mark.parametrize("field", ["boxL", "totalCbm", "grossWeightKg", "invoiceValue"])
@pytest.mark.parametrize("value", [float("inf"), float("nan")])
def test_nonfinite_cargo_is_invalid(client, connect, session, field, value):
    case = services.get_case(session, create(client, text_mail(FULL_LCL))["caseId"])
    setattr(case.quote_input, field, value)
    issues = validation.check(session, case.quote_input, today=TODAY)
    expected = "volume" if field in ("boxL", "totalCbm") else field
    assert any(i["field"] == expected and i["code"] == "INVALID" for i in issues)


def test_confirming_one_dimension_does_not_confirm_another(client, connect, session):
    case = services.get_case(session, create(client, text_mail(FULL_LCL))["caseId"])
    result = validation.RuleValidator().validate(session, case, {
        "fields": {"boxW": {"status": "review", "evidence": "폭은 재확인이 필요합니다"}},
        "lockedByManual": ["boxL"],
    })
    assert result["status"] == "정보부족" and result["quote"] is None
    assert any(i["field"] == "volume" for i in result["issues"])


def test_unreadable_llm_value_blocks_stale_value_without_quoting_it(client, connect, session):
    case = services.get_case(session, create(client, text_mail(FULL_LCL))["caseId"])
    result = validation.RuleValidator().validate(session, case, {
        "fields": {"grossWeightKg": {"status": "review", "reason": "UNPARSED", "method": "llm",
                                     "evidence": "LLM이 잘못 짚었을 수 있는 줄"}},
    })
    assert result["status"] == "정보부족" and result["quote"] is None
    assert result["questions"] == [validation.rules()["questions"]["grossWeightKg"]]


def test_same_charge_at_origin_and_destination_is_not_dropped():
    origin = Rate(id=1, charge_code="THC", source="origin table", direction="ORIGIN")
    destination = Rate(id=2, charge_code="THC", source="destination table", direction="DEST")
    newer_origin = Rate(id=3, charge_code="THC", source="new origin table", direction="ORIGIN")
    selected = quotation.one_table_per_charge([origin, destination, newer_origin])
    assert [r.id for r in selected] == [2, 3]


@pytest.mark.parametrize("sheet", ["OceanFreight", "Surcharges", "요율"])
def test_duplicate_upload_reports_both_rows_and_preserves_rates(session, client, sheet):
    seed_containers(session)
    rates.import_workbook(session, CLIENT_FILE.read_bytes())
    before = rate_values(session)
    wb = load_workbook(io.BytesIO(rates.export(session))) if sheet == "요율" else load_workbook(CLIENT_FILE)
    ws = wb[sheet]
    original_row = 2 if sheet == "요율" else 5
    duplicate_row = ws.max_row + 1
    ws.append([cell.value for cell in ws[original_row]])
    # 관리용 원본 행·정렬순서가 달라도 같은 비용이면 중복이다.
    if sheet == "요율":
        headers = [cell.value for cell in ws[1]]
        ws.cell(duplicate_row, headers.index("원본 행") + 1, "copied row")
        ws.cell(duplicate_row, headers.index("정렬순서") + 1, 999)
    response = client.post("/api/rates/import", files={"file": ("duplicate.xlsx", save(wb))})
    assert response.status_code == 400
    message = "\n".join(response.json()["detail"]["errors"])
    assert f"{sheet} {original_row}행" in message and f"{sheet} {duplicate_row}행" in message
    assert "동일한 요율" in message
    session.expire_all()
    assert rate_values(session) == before


def test_same_charge_with_distinct_direction_or_period_is_allowed(session):
    seed_containers(session)
    wb = Workbook()
    ws = wb.active
    ws.title = "요율"
    headers = [*rates.RATE_COLUMNS, *rates.RATE_EXTRA_COLUMNS]
    ws.append(headers)
    for direction in ("ORIGIN", "DEST"):
        for month in (10, 11):
            values = {"구분코드": "THC", "견적서 표기": "THC", "선적항(POL)": "CNSHA",
                      "도착항(POD)": "KRPUS", "컨테이너": "20GP", "기준": "CNTR", "단가": 120,
                      "통화": "USD", "출처": "direction and period", "비용 방향": direction,
                      "적용 시작": f"2026-{month}-01", "적용 종료": f"2026-{month}-28"}
            ws.append([values.get(header) for header in headers])
    assert rates.import_workbook(session, save(wb))["rates"] == 4


@pytest.mark.parametrize("wanted,second_carrier", [(None, "B"), ("missing", "B"), ("A", "A")])
def test_mixed_freight_currencies_require_review(wanted, second_carrier):
    rows = [Rate(id=1, charge_code="OCEAN_FREIGHT", carrier="A", unit_price=500, currency="USD"),
            Rate(id=2, charge_code="OCEAN_FREIGHT", carrier=second_carrier, unit_price=1000, currency="KRW")]
    with pytest.raises(quotation.RateSelectionError, match="통화가 서로 다릅니다"):
        quotation.choose_carrier(rows, wanted)


def test_requested_carrier_can_be_selected_without_comparing_other_currencies():
    rows = [Rate(id=1, charge_code="OCEAN_FREIGHT", carrier="A", unit_price=500, currency="USD"),
            Rate(id=2, charge_code="OCEAN_FREIGHT", carrier="B", unit_price=1000, currency="KRW")]
    chosen, carrier = quotation.choose_carrier(rows, "A")
    assert chosen == [rows[0]] and carrier == "A"


def test_same_currency_freight_still_uses_lowest_price_with_other_currency_surcharge():
    rows = [Rate(id=1, charge_code="OCEAN_FREIGHT", carrier="A", unit_price=500, currency="USD"),
            Rate(id=2, charge_code="OCEAN_FREIGHT", carrier="B", unit_price=400, currency="USD"),
            Rate(id=3, charge_code="THC", unit_price=10000, currency="KRW")]
    chosen, carrier = quotation.choose_carrier(rows, None)
    assert chosen == rows[1:] and carrier == "B"


def test_mixed_currency_recalculation_holds_case_and_retires_old_quote(session, client, connect):
    rates.import_workbook(session, CLIENT_FILE.read_bytes())
    body = ("품목: 부품\n수량: 10박스\n전체 부피: 5CBM\n총중량: 500kg\n선적항: 상하이\n도착항: 부산\n"
            "화물 준비일: 2026-10-20\n조건: CIF\n컨테이너: 20GP 1대\n인보이스 금액: USD 15000\n결제조건: T/T Advance")
    created = create(client, text_mail(body))
    first = created["extraction"]["validation"]["quote"]
    assert first is not None
    from app.config import settings
    original_file = settings.storage_dir / first["file"]
    assert original_file.exists()
    case = services.get_case(session, created["caseId"])
    selected, reason = quotation.rates_with_reason(session, case.quote_input,
                                                  quotation.rate_date(case.quote_input, TODAY))
    assert reason is None
    freight = next(row for row in selected if row.charge_code == "OCEAN_FREIGHT")
    alternative = {col.name: getattr(freight, col.name) for col in Rate.__table__.columns if col.name != "id"}
    alternative.update(carrier="OTHER", currency="KRW", unit_price=1000)
    session.add(Rate(**alternative))
    session.commit()

    result = validation.RuleValidator().validate(session, case, created["extraction"])
    assert result["status"] == "보류" and result["quote"] is None
    assert any("통화가 서로 다릅니다" in item["message"] for item in result["hold"])
    assert not original_file.exists() and result["quoteRetired"]
