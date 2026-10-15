"""C트랙 견적 출력 회귀: 원본 고정 양식, 합계, 적용 비용, 필수 산출물."""

import io
from copy import copy
from datetime import datetime, timezone
from types import SimpleNamespace as NS
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import quotation, quote_pdf
from app.config import settings
from app.models import QuoteItem, QuoteTotal
from .test_pipeline_flow import CLIENT_CASE, FCL_SG, FULL_LCL, client_rates, connect, create, quote_folder, text_mail


@pytest.fixture(autouse=True)
def english_default(monkeypatch):
    monkeypatch.setattr(settings, "quote_language", "en")


def item(code, rate, basis="PER_CNTR", qty=1, currency="USD", label=None):
    return QuoteItem(charge_code=code, charge_label=label or code, basis=basis, rate=rate, qty=qty,
                     amount=rate * qty if rate is not None and qty is not None else None,
                     currency=currency, remark="internal memo", source_ref="Surcharges#1")


def example():
    qi = NS(customerName="TEST TRADING", contactName="Test Buyer", contactEmail="buyer@example.com",
            containerType="20GP", incoterms="CIF", pol="CNSHA", pod="KRPUS", qty=10, qtyUnit="CTNS",
            boxL=1000, boxW=1000, boxH=500, totalCbm=5, grossWeightKg=500,
            exchangeRateAsOf="2026-10-15", quoteCurrency="USD", transitTime=2, freeTimeDet=7, freeTimeDem=7)
    case = NS(case_id="LQ-2026-1015-002", quote_input=qi)
    items = [item("OCEAN_FREIGHT", 520), item("OTHC", 120), item("DTHC", 130),
             item("DOC", 30, "PER_BL"), item("BL", 30, "PER_BL"),
             item("HANDLING", 50, "PER_BL", label="Handling Charge"), item("INS", None, "PERCENT", None), item("BAF", 50)]
    items[6].amount = 33
    items[6].remark = "INV.V x 110% x 0.2% (MIN USD 30)"
    quote = NS(items=items, totals=[QuoteTotal(currency="USD", total=963)], version_no=1,
               created_at=datetime(2026, 10, 15, tzinfo=timezone.utc), etd="2026-10-20", eta="2026-10-22",
               carrier="HMM", valid_until="2026-10-29")
    return case, quote


def workbook(case, quote):
    data = quotation.render_xlsx(case, quote, 1)
    return load_workbook(io.BytesIO(data)).active, data


def test_original_layout_drawing_and_cost_labels_are_preserved():
    case, quote = example()
    ws, data = workbook(case, quote)
    source = load_workbook(settings.quote_template).active
    assert [ws[f"B{r}"].value for r in range(22, 36)] == [source[f"B{r}"].value for r in range(22, 36)]
    assert set(map(str, ws.merged_cells.ranges)) == set(map(str, source.merged_cells.ranges))
    assert str(ws.print_area) == str(source.print_area)
    for col in "ABCDEF":
        assert ws.column_dimensions[col].width == source.column_dimensions[col].width
    for row in range(1, 50):
        assert ws.row_dimensions[row].height == source.row_dimensions[row].height
        assert ws.row_dimensions[row].hidden == source.row_dimensions[row].hidden
    for row in range(22, 36):
        for col in "BCDEF":
            assert copy(ws[f"{col}{row}"].border) == copy(source[f"{col}{row}"].border)
            assert copy(ws[f"{col}{row}"].fill) == copy(source[f"{col}{row}"].fill)
    with ZipFile(settings.quote_template) as original, ZipFile(io.BytesIO(data)) as filled:
        for name in original.namelist():
            if name.startswith(("xl/drawings/", "xl/media/", "xl/worksheets/_rels/", "xl/printerSettings/")):
                assert original.read(name) == filled.read(name), name
        assert "xl/calcChain.xml" not in filled.namelist()
    assert ws["E21"].value == source["E21"].value and ws["F21"].value == "2026-10-15"
    assert ws["D22"].value == "20FT"


@pytest.mark.parametrize("configured_language", ["en", "ko"])
def test_quote_language_preserves_styles_and_matches_pdf(monkeypatch, configured_language):
    monkeypatch.setattr(settings, "quote_language", configured_language)
    case, quote = example()
    case.quote_input.customerName, case.quote_input.contactName = "테스트무역", "김테스트"
    ws, data = workbook(case, quote)
    original = load_workbook(settings.quote_template).active
    pdf = PdfReader(io.BytesIO(quote_pdf.render_pdf(data)))
    assert len(pdf.pages) == 1
    text = pdf.pages[0].extract_text()
    english = {"B12": "Subject : Ocean Freight Quotation",
               "B14": "* We wish your company continued success.",
               "B15": "* Please find our quotation below for your reference.",
               "B48": "* Thank you.", "F49": "(Company seal omitted)"}
    for ref in ("B12", "B14", "B15", "B48", "F49"):
        expected = english[ref] if configured_language == "en" else original[ref].value
        assert ws[ref].value == expected
        assert copy(ws[ref].font) == copy(original[ref].font)
        assert expected in text
    if configured_language == "en":
        assert ws["B9"].value == "To : 테스트무역 / Attn : 김테스트"  # 고객 고유명사는 번역하지 않는다
        assert ws["B10"].value == "From : LEONA SEA & AIR CO., LTD. / Quotation Team"
        assert "included in TOTAL" in text and "수 신 :" not in text
    else:
        assert ws["B9"].value == "수 신 : 테스트무역 / 김테스트"
        assert ws["B10"].value == "발 신 : 레오나 해운항공㈜ / 견적 담당 드림"
        assert "TOTAL 포함" in text and "Thank you" not in text
    assert ws["B9"].value in text and ws["B10"].value in text
    assert case.case_id in text


@pytest.mark.parametrize("configured_language", [None, "", "invalid", "EN"])
def test_missing_or_invalid_quote_language_defaults_to_english(monkeypatch, configured_language):
    monkeypatch.setattr(settings, "quote_language", configured_language)
    ws, _ = workbook(*example())
    assert ws["B12"].value == "Subject : Ocean Freight Quotation"
    assert ws["B48"].value == "* Thank you."


@pytest.mark.parametrize("lang", ["en", "ko"])
@pytest.mark.parametrize("body,currencies,totals", [
    (FCL_SG.replace("조건: CIF", "조건: FOB"), "KRW", {"KRW": 498220.0}),
    (FCL_SG, "USD / KRW", {"USD": 750.0, "KRW": 498220.0}),
    (CLIENT_CASE.format(carrier="선사: HMM\n").replace("2026-10-05", "2026-10-20"),
     "USD", {"USD": 963.0}),
], ids=["export-fob-krw", "export-cif-mixed", "import-cif-usd"])
def test_currency_notice_matches_actual_totals_in_xlsx_and_pdf(
    client, client_rates, monkeypatch, lang, body, currencies, totals,
):
    monkeypatch.setattr(settings, "quote_language", lang)
    case = create(client, text_mail(body))
    result = case["extraction"]["validation"]
    assert case["status"] == "계산완료", result
    quote = result["quote"]
    assert quote["totals"] == totals  # 통화 안내를 바꾸어도 금액·비용 분기는 그대로다.
    ws = load_workbook(settings.storage_dir / quote["file"]).active
    expected = quotation.TEXT[lang]["fx"].format(currency=currencies, as_of="2026-09-29")
    assert ws["B39"].value == expected
    pdf = PdfReader(settings.storage_dir / quote["pdf"])
    assert len(pdf.pages) == 1
    assert " ".join(expected.split()) in " ".join(pdf.pages[0].extract_text().split())
    assert len(quote_folder(case["caseId"])) == 2


def test_grouping_and_handling_keep_original_963_total_and_sources():
    case, quote = example()
    before = [(i.charge_code, i.amount, i.source_ref) for i in quote.items]
    ws, data = workbook(case, quote)
    assert ws["C23"].value == ws["E23"].value == 570
    assert ws["C24"].value == ws["E24"].value == 250
    assert ws["C25"].value == ws["E25"].value == 60
    assert ws["E34"].value == 33 and "110%" in ws["F34"].value
    assert "BAF USD50.00" in ws["F23"].value
    assert "Origin THC" in ws["F24"].value and "Destination THC" in ws["F24"].value
    assert "B/L USD30.00" in ws["F25"].value
    remarks = "\n".join(ws[f"B{r}"].value or "" for r in range(39, 48))
    assert "Handling Charge: USD50.00 PER B/L x 1 = USD50.00 (included in TOTAL)" in remarks
    assert sum(ws[f"E{r}"].value or 0 for r in range(23, 35)) + 50 == 963
    assert ws["C35"].value == "USD963.00"
    assert [(i.charge_code, i.amount, i.source_ref) for i in quote.items] == before
    assert all(ws[f"D{r}"].value is None for r in range(26, 34))
    pdf = PdfReader(io.BytesIO(quote_pdf.render_pdf(data)))
    assert len(pdf.pages) == 1
    text = pdf.pages[0].extract_text()
    for word in ("SHORING CHARGE", "PSF", "PFS", "Handling Charge", "USD963.00", "110%", case.case_id):
        assert word in text


def test_handling_not_applied_is_not_printed():
    case, quote = example()
    quote.items = [i for i in quote.items if i.charge_code != "HANDLING"]
    quote.totals[0].total = 913
    ws, _ = workbook(case, quote)
    assert "Handling" not in " ".join(str(c.value or "") for row in ws for c in row)
    assert ws["C35"].value == "USD913.00"


@pytest.mark.parametrize("other", [item("BL", 30000, "PER_BL", currency="KRW"),
                                    item("BL", 30, "PER_CNTR"), item("BL", 30, "PER_BL", qty=2)])
def test_incompatible_currency_basis_or_quantity_is_not_combined(other):
    base = item("DOC", 30, "PER_BL")
    groups, extras = quotation.display_groups([base, other])
    assert groups == {25: [base]} and extras == [other]
    assert other.currency in quotation.extra_charge_text(other, "en")


def test_baf_without_freight_is_a_separate_remark():
    baf = item("BAF", 50)
    groups, extras = quotation.display_groups([baf])
    assert groups == {} and extras == [baf]


def test_zero_free_time_is_zero_and_missing_values_are_explicit():
    case, quote = example()
    case.quote_input.freeTimeDet = case.quote_input.transitTime = 0
    case.quote_input.freeTimeDem = None
    quote.etd = quote.eta = None
    ws, _ = workbook(case, quote)
    remarks = "\n".join(ws[f"B{r}"].value or "" for r in range(39, 48))
    assert "DET 0 days / DEM to be confirmed" in remarks
    assert "T/T : approx. 0 days" in remarks
    assert "ETD / ETA : to be confirmed / to be confirmed" in remarks


def test_overflow_never_silently_drops_a_charge():
    case, quote = example()
    quote.items.extend(item(f"EXTRA_{i}", 10) for i in range(3))
    with pytest.raises(quotation.QuoteGenerationError, match="Remark 공간"):
        workbook(case, quote)


def test_rate_validity_controls_handling_in_real_generation(client, client_rates, session):
    from sqlalchemy import select
    from datetime import date
    from app.models import Rate

    body = CLIENT_CASE.format(carrier="선사: HMM\n").replace("2026-10-05", "2026-10-20")
    for rate in session.scalars(select(Rate).where(Rate.charge_code == "HANDLING")):
        rate.valid_until = date(2026, 10, 19)
    session.commit()
    case = create(client, text_mail(body))
    result = case["extraction"]["validation"]["quote"]
    assert result["totals"] == {"USD": 913.0}
    assert not any("HANDLING" in line["charge"].upper() for line in result["lines"])
    ws = load_workbook(settings.storage_dir / result["file"]).active
    assert not any("Handling" in (ws[f"B{row}"].value or "") for row in range(39, 48))


def test_zero_transit_derives_same_day_eta(client, client_rates, session):
    from sqlalchemy import select
    from app.models import Rate

    for rate in session.scalars(select(Rate).where(Rate.charge_code == "OCEAN_FREIGHT", Rate.carrier == "HMM")):
        rate.transit_days = rate.free_time_det = rate.free_time_dem = 0
        rate.eta = None
    session.commit()
    case = create(client, text_mail(CLIENT_CASE.format(carrier="선사: HMM\n").replace("2026-10-05", "2026-10-20")))
    quote = case["extraction"]["validation"]["quote"]
    assert quote["etd"] == quote["eta"] == "2026-10-20"


def test_pdf_disabled_is_held_without_a_quote_version(client, connect, session, monkeypatch):
    from sqlalchemy import select, func
    from app.models import Quote, QuoteItem, QuoteTotal

    monkeypatch.setattr(settings, "quote_pdf", False)
    case = create(client, text_mail(FULL_LCL))
    assert case["status"] == "보류" and quote_folder(case["caseId"]) == []
    assert "PDF 생성 설정이 꺼져" in case["extraction"]["validation"]["hold"][0]["message"]
    for model in (Quote, QuoteItem, QuoteTotal):
        assert session.scalar(select(func.count()).select_from(model)) == 0


def test_partial_save_is_retired_and_retry_preserves_history(client, connect, session, monkeypatch):
    from app import services
    from app.validation import RuleValidator
    from sqlalchemy import select, func
    from app.models import Quote

    case_id = create(client, text_mail(FULL_LCL))["caseId"]
    folder = settings.storage_dir / case_id
    old = {p.name: p.read_bytes() for p in (folder / "quotes").iterdir()}
    real_save = quotation._save

    def broken(case, quote, ext, data):
        path = real_save(case, quote, ext, data)
        if ext == "pdf":
            raise OSError("test disk failure after write")
        return path

    case = services.get_case(session, case_id)
    monkeypatch.setattr(quotation, "_save", broken)
    result = RuleValidator().validate(session, case, {})
    assert result["status"] == "보류" and quote_folder(case_id) == []
    assert session.scalar(select(func.count()).select_from(Quote)) == 1
    for name, data in old.items():
        path = folder / "quote-history" / name.replace("_견적서.", "_견적서_v1.")
        assert path.read_bytes() == data
    history_before = {p.name: p.read_bytes() for p in (folder / "quote-history").iterdir()}
    monkeypatch.setattr(quotation, "_save", real_save)
    assert RuleValidator().validate(session, case, {})["status"] == "계산완료"
    # 같은 번호의 실패 파일이 보관되어 있어도 다음 보관 시 덮어쓰지 않는다.
    quotation.retire_current(session, case)
    for name, data in history_before.items():
        assert (folder / "quote-history" / name).read_bytes() == data
