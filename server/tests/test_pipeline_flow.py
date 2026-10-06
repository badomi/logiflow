"""케이스 생성 → 추출(B) → 검증(C) → 보완 요청 / 회신 병합 → 견적서 흐름 테스트.

룰 모드는 실제 샘플 메일로 돌린다. 하이브리드 모드는 FakeLlm(가짜 LLM 응답)으로 결합·환각 차단 규칙을 확인한다.
실제 LLM 품질은 eval/run_eval.py로 따로 잰다 (NFR-01).
"""

import base64
from pathlib import Path
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook

from app import pipeline, validation
from app.config import settings
from app.db import SessionLocal
from app.extraction import HybridExtractor
from app.seed import seed_containers, seed_sample_rates
from app.validation import RuleValidator

from .conftest import sample

TODAY = date(2026, 9, 29)
REST_OF_LCL = "요청하신 정보 회신드립니다.\n선적항: 인천항\n조건: CIF\n결제조건: T/T 30% Advance\n전체 부피: 0.38CBM\n감사합니다."


@pytest.fixture
def connect(monkeypatch):
    monkeypatch.setattr(validation, "today_kst", lambda: TODAY)
    monkeypatch.setattr(settings, "quote_pdf", False)  # PDF는 LibreOffice 있는 PC에서만 → 전용 테스트에서 확인
    with SessionLocal() as s:
        seed_containers(s)
        seed_sample_rates(s)
        s.commit()

    def _connect(llm=None, mode="rules"):
        pipeline.register(extractor=HybridExtractor(llm=llm, llm_mode=mode), validator=RuleValidator())

    _connect()
    yield _connect
    pipeline.register(extractor=pipeline._NotConnected(), validator=pipeline._NotConnected())


def eml(name: str) -> dict:
    return {"emlBase64": base64.b64encode(sample(name)).decode(), "actor": "quote@leona.example.com"}


def text_mail(body: str, subject="견적 요청", sender="buyer@abc-trading.example.com") -> dict:
    return {"snapshot": {"subject": subject, "from": {"name": "", "email": sender}, "bodyText": body,
                         "internetMessageId": f"<{uuid.uuid4()}@test>"}, "actor": "quote@leona.example.com"}


def create(client, body: dict) -> dict:
    case_id = client.post("/api/cases", json=body).json()["case"]["caseId"]
    return client.get(f"/api/cases/{case_id}").json()


def reply(client, case_id: str, body: dict) -> dict:
    client.post(f"/api/cases/{case_id}/replies", json=body)
    return client.get(f"/api/cases/{case_id}").json()


FULL_LCL = (
    "품목: 스프링노트\n박스 수량: 10박스\n박스 1개 규격: 310 × 450 × 270mm\n전체 부피: 0.38CBM\n예상 총중량: 500kg\n"
    "선적항: 인천\n도착항: 시드니\n화물 준비일: 2026-10-20\n조건: CIF\n선적방식: LCL\n인보이스 금액: USD 3,000\n"
    "결제조건: T/T 30% Advance\n감사합니다."
)


# ------------------------------------------------------------------ MUST-SHIP ①②③④ (룰 모드, 실제 샘플)


def test_lcl_request_becomes_supplement_questions(client, connect):
    case = create(client, eml("01_lcl_request.eml"))
    assert case["status"] == "정보부족"
    questions = "\n".join(case["extraction"]["validation"]["questions"])
    for expected in ("선적항", "화물 준비일", "FOB / CIF", "0.38CBM", "결제 조건"):
        assert expected in questions
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["pol"]["value"] is None  # 픽업지 '인천'을 선적항으로 쓰지 않는다
    assert fields["commodity"]["evidence"] == "품목: 스프링노트(일반 문구류)"  # FR-202 근거 원문
    assert fields["boxL"]["value"] == 310 and fields["grossWeightKg"]["score"] == 1.0


def test_reply_merge_then_quote(client, connect):
    case_id = create(client, eml("01_lcl_request.eml"))["caseId"]
    case = reply(client, case_id, eml("03_lcl_reply.eml"))  # 준비일·인보이스 금액 보충
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case_id}/fields").json()}
    assert fields["cargoReadyDate"]["value"] == "2026-10-20" and fields["ccy"]["value"] == "USD"
    assert fields["cargoReadyDate"]["origin"] == "메일 2 · 회신 본문"
    assert case["status"] == "정보부족"  # 선적항·조건·결제조건은 아직

    case = reply(client, case_id, text_mail(REST_OF_LCL, subject=f"RE: [{case_id}] 견적 보완"))
    assert case["status"] == "계산완료", case["extraction"]["validation"]
    assert case["extraction"]["validation"]["quote"]["totals"] == {"USD": 500.0, "KRW": 140203.0}
    draft = client.post(f"/api/cases/{case_id}/drafts/quote", json={}).json()
    assert [a["filename"] for a in draft["attachments"]] == [f"{case_id}_견적서.xlsx"]


def test_quote_xlsx_follows_template(client, connect):
    case_id = create(client, text_mail(FULL_LCL))["caseId"]
    ws = load_workbook(settings.storage_dir / case_id / "quotes" / f"{case_id}_견적서.xlsx").active
    assert case_id in ws["B12"].value and ws["C17"].value == ": INCHEON - SYDNEY" and ws["C20"].value == ": CIF"
    assert ws["C35"].value == "USD500.00 + KRW140,203" and len(ws._images) == 1
    # 예시 견적서 형식: 한 페이지 맞춤(F열 주소·DATE·REMARK 포함), R/T 기준, 쓰지 않은 줄 숨김. 문구는 기본 영문 (FR-510)
    assert ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1
    assert ws["F2"].value == "LEONA SEA & AIR CO., LTD." and ws["F21"].value.startswith("DATE : 20")
    assert ws["D22"].value == "R/T" and ws["C19"].value.endswith("(based on 1 R/T)")
    assert ws["B29"].value == "CUSTOMS CLEARANCE FEE" and ws["F29"].value == "INV.V x 1/1,000 (MIN)"
    assert ws.row_dimensions[31].hidden and not ws.row_dimensions[30].hidden
    assert any(ws[f"B{r}"].value.startswith("* VALIDITY : 2026-") for r in range(39, 45) if ws[f"B{r}"].value)


def fixed_text(ws) -> str:
    """견적서에서 시스템이 쓰는 고정 문구 칸 (화주가 적은 값이 들어가는 수신·화물 줄은 뺀다)."""
    refs = ["B10", "B12", "B14", "B15", "B48", "F49", *[f"B{r}" for r in range(38, 48)], *[f"F{r}" for r in range(22, 36)]]
    return "\n".join(str(ws[ref].value) for ref in refs if ws[ref].value)


def test_quote_is_english_by_default(client, connect):
    """FR-510: 견적서는 기본 영문 — 고정 문구·비고에 한글이 없다."""
    case_id = create(client, text_mail(FULL_LCL))["caseId"]
    ws = load_workbook(settings.storage_dir / case_id / "quotes" / f"{case_id}_견적서.xlsx").active
    text = fixed_text(ws)
    assert not [ch for ch in text if "가" <= ch <= "힣"], text
    assert ws["B10"].value == "From : LEONA SEA & AIR CO., LTD. / Quotation Team"
    assert ws["B12"].value == f"Subject : Ocean Freight Quotation  [{case_id}]"
    # 회사명·담당자명이 없는 메일이면 수신 줄에 회신 주소를 적는다 (빈 줄로 나가지 않게)
    assert ws["B9"].value == "To : buyer@abc-trading.example.com" and ws["B48"].value == "* Thank you."
    assert f"* QUOTE NO. : {case_id} (Rev. 1)" in text and "* Exchange rate on the actual sailing date" in text
    assert "* Destination charges are for consignee's account." in text  # CIF 수출 견적 고지


def test_quote_language_can_be_korean(client, connect, monkeypatch):
    """QUOTE_LANGUAGE=ko 이면 발주 측 예시 견적서의 한글 문구를 쓴다."""
    monkeypatch.setattr(settings, "quote_language", "ko")
    case_id = create(client, text_mail(FULL_LCL))["caseId"]
    ws = load_workbook(settings.storage_dir / case_id / "quotes" / f"{case_id}_견적서.xlsx").active
    assert ws["B10"].value == "발 신 : 레오나 해운항공㈜ / 견적 담당 드림" and "해상 수출 운임 제안서" in ws["B12"].value
    assert ws["C19"].value.endswith("(1 R/T 기준)") and ws["F29"].value == "INV.V x 1/1,000 (MIN 기준)"
    assert ws["B48"].value == "* 감사합니다." and ws["B9"].value == "수 신 : buyer@abc-trading.example.com"
    assert "* 도착지 비용 수하인 부담" in [ws[f"B{r}"].value for r in range(39, 48)]
    assert any((ws[f"B{r}"].value or "").startswith("* VALIDITY : 2026년") for r in range(39, 48))


def test_multi_item_sample_is_held_without_mixed_values(client, connect):
    case = create(client, eml("02_fob_multi_item.eml"))
    assert case["status"] == "보류"
    assert {h["code"] for h in case["extraction"]["validation"]["hold"]} >= {"MULTIPLE_ITEMS"}
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["qty"]["value"] is None and fields["containerType"]["value"] is None  # 품목 값 섞임 방지
    assert fields["contactEmail"]["value"] == "shipper@example.com"


def test_no_rate_and_dangerous_goods(client, connect):
    no_rate = create(client, text_mail(FULL_LCL.replace("시드니", "뉴욕")))
    assert no_rate["status"] == "보류" and no_rate["extraction"]["validation"]["hold"][0]["code"] == "NO_RATE"
    dg = create(client, text_mail(FULL_LCL.replace("스프링노트", "리튬 배터리")))
    assert dg["status"] == "보류" and dg["extraction"]["validation"]["hold"][0]["code"] == "DANGEROUS_GOODS"


# ------------------------------------------------------------------ 하이브리드 (가짜 LLM)


class FakeLlm:
    def __init__(self, fields: dict, multi=False):
        self.fields, self.multi, self.calls = fields, multi, 0

    def complete_json(self, system, user, schema):
        self.calls += 1
        return {"fields": self.fields, "multipleItems": self.multi}


def test_llm_fills_prose_mail_rules_miss(client, connect):
    prose = ("안녕하세요. 인천에서 시드니로 LCL 견적 부탁드립니다. 스프링노트 10박스이고 한 박스가 310 × 450 × 270mm, "
             "전부 합쳐 0.38CBM에 500kg 정도입니다. 2026년 10월 20일에 준비되며 CIF 조건, 인보이스 금액은 USD 3,000, "
             "결제는 T/T 30% Advance 입니다.\n감사합니다.")
    llm = FakeLlm({
        "commodity": {"value": "스프링노트", "quote": "스프링노트 10박스이고"},
        "quantity": {"value": "10박스", "quote": "스프링노트 10박스이고"},
        "portOfLoading": {"value": "인천", "quote": "인천에서 시드니로 LCL 견적 부탁드립니다."},
        "portOfDischarge": {"value": "시드니", "quote": "인천에서 시드니로 LCL 견적 부탁드립니다."},
        "grossWeight": {"value": "500kg", "quote": "전부 합쳐 0.38CBM에 500kg 정도입니다."},
        "totalVolume": {"value": "0.38CBM", "quote": "전부 합쳐 0.38CBM에 500kg 정도입니다."},
        "cargoReadyDate": {"value": "2026년 10월 20일", "quote": "2026년 10월 20일에 준비되며 CIF 조건"},
        "incoterms": {"value": "CIF", "quote": "2026년 10월 20일에 준비되며 CIF 조건"},
        "invoiceValue": {"value": "USD 3,000", "quote": "인보이스 금액은 USD 3,000"},
        "paymentTerm": {"value": "T/T 30% Advance", "quote": "결제는 T/T 30% Advance 입니다."},
    })
    connect(llm, mode="hybrid")
    case = create(client, text_mail(prose))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert llm.calls == 1
    assert fields["pol"]["value"] == "KRINC" and fields["pol"]["method"] == "llm"
    assert fields["containerType"]["method"] == "rule+llm" or fields["containerType"]["value"] == "LCL"
    assert case["status"] == "계산완료"


def test_hallucinated_llm_value_is_dropped(client, connect):
    connect(FakeLlm({
        "paymentTerm": {"value": "T/T 100%", "quote": "결제조건: T/T 100%"},  # 메일에 없는 문장
        "commodity": {"value": "스프링노트", "quote": "품목: 스프링노트(일반 문구류)"},
    }), mode="hybrid")
    case = create(client, eml("01_lcl_request.eml"))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["paymentTerm"]["value"] is None  # 근거가 원문에 없어 버림 (FR-202 정밀도 우선)
    assert "commodity" in case["extraction"]["result"]["llm"]["asked"]  # LLM 중심: 모든 항목을 묻는다
    assert fields["commodity"]["method"] == "rule+llm"  # 규칙·LLM이 같은 값 → 점수 가산


def test_rule_llm_conflict_goes_to_review(client, connect):
    connect(FakeLlm({"grossWeight": {"value": "50kg", "quote": "예상 총중량: 최대 500kg"}}), mode="hybrid")
    case = create(client, eml("01_lcl_request.eml"))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    # 인용 문장 안에 '50'이 없으므로 LLM 값은 근거 검증에서 탈락 → 룰 값 유지
    assert fields["grossWeightKg"]["value"] == 500.0


def test_llm_down_falls_back_to_rules(client, connect):
    class Down:
        def complete_json(self, *a):
            raise RuntimeError("LLM 서버(http://localhost:11434)에 연결하지 못했습니다")

    connect(Down(), mode="hybrid")
    case = create(client, eml("01_lcl_request.eml"))
    assert case["status"] == "정보부족"  # 실패로 멈추지 않고 룰 결과로 진행 (NFR-03)
    assert "연결하지 못했습니다" in case["extraction"]["result"]["llm"]["error"]


# ------------------------------------------------------------------ FR-206 수동 보정


def test_manual_correction_survives_reextraction(client, connect):
    case_id = create(client, eml("01_lcl_request.eml"))["caseId"]
    res = client.post(f"/api/cases/{case_id}/fields/pol", json={"value": "KRINC", "actor": "kim@leona"})
    assert res.status_code == 200
    reply(client, case_id, eml("03_lcl_reply.eml"))  # 재추출해도 담당자 값 유지
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case_id}/fields").json()}
    assert fields["pol"]["value"] == "KRINC" and fields["pol"]["method"] == "manual"
    events = [e["eventType"] for e in client.get(f"/api/cases/{case_id}").json()["events"]]
    assert "FIELD_CORRECTED" in events
    assert client.post(f"/api/cases/{case_id}/fields/qty", json={"value": "열개"}).status_code == 400
    assert client.post(f"/api/cases/{case_id}/fields/unknown", json={"value": "x"}).status_code == 400


# ------------------------------------------------------------------ A트랙 임시 부분 연결 (보완 문항·받는 사람)


def test_supplement_draft_uses_validation_questions_when_empty(client, connect):
    case = create(client, eml("01_lcl_request.eml"))
    expected = case["extraction"]["validation"]["questions"]
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/supplement", json={"questions": []}).json()
    assert draft["htmlBody"].count("<li>") == len(expected) > 0  # FR-305: 문항 = 검증 결과
    edited = client.post(f"/api/cases/{case['caseId']}/drafts/supplement", json={"questions": ["직접 쓴 문항"]}).json()
    assert edited["htmlBody"].count("<li>") == 1  # 담당자가 고친 문항이 있으면 그것을 쓴다


def test_supplement_draft_reports_question_mismatch(client, connect):
    """FR-305: 초안 문항이 검증 결과와 같은지 확인해 돌려주고 이력에 남긴다 (담당자 수정은 막지 않는다)."""
    case = create(client, eml("01_lcl_request.eml"))
    expected = case["extraction"]["validation"]["questions"]
    assert case["extraction"]["validation"]["questionCount"] == len(expected) >= 2
    url = f"/api/cases/{case['caseId']}/drafts/supplement"
    same = client.post(url, json={"questions": []}).json()
    assert same["questionCheck"] == {"matches": True, "expected": len(expected), "actual": len(expected),
                                     "added": [], "missing": []}
    edited = client.post(url, json={"questions": [expected[0], "직접 쓴 문항"]}).json()
    check = edited["questionCheck"]
    assert check["matches"] is False and check["added"] == ["직접 쓴 문항"] and check["missing"] == expected[1:]
    events = client.get(f"/api/cases/{case['caseId']}").json()["events"]
    assert events[-1]["detail"]["questionCheck"]["matches"] is False


def test_issues_have_severity(client, connect):
    """FR-301: 문제마다 심각도(Block/Warn/Info)가 붙고, 누락은 Block이다."""
    validation_result = create(client, eml("01_lcl_request.eml"))["extraction"]["validation"]
    issues = validation_result["issues"]
    assert issues and all(i["severity"] in ("Block", "Warn", "Info") for i in issues)
    missing = [i for i in issues if i["code"] == "MISSING"]
    assert missing and all(i["severity"] == "Block" for i in missing)
    # 정의서 2장: Block = 견적 산출 불가. 화주 문항이 있으면 견적을 만들지 않으므로 전부 Block이어야 한다
    assert all(i["severity"] == "Block" for i in issues if i["question"])
    assert sum(validation_result["severityCount"].values()) == len(issues) + len(validation_result["hold"])


def test_draft_recipient_uses_extracted_contact(client, connect):
    body = "품목: 노트\n감사합니다.\n김민수 대리\n㈜한빛무역"
    case = create(client, text_mail(body, sender="ms.kim@hanbit.example.com"))
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/supplement", json={}).json()
    assert draft["to"] == [{"name": "김민수 대리", "email": "ms.kim@hanbit.example.com"}]  # FR-208


def test_quote_draft_message_explains_status(client, connect):
    case = create(client, eml("01_lcl_request.eml"))
    res = client.post(f"/api/cases/{case['caseId']}/drafts/quote", json={})
    assert res.status_code == 400 and "정보부족" in res.json()["detail"]


# ------------------------------------------------------------------ NFR-03 다시 추출


def test_rerun_revives_case_created_while_extraction_was_off(client, connect):
    pipeline.register(extractor=pipeline._NotConnected(), validator=pipeline._NotConnected())  # 추출 꺼짐
    case_id = create(client, eml("01_lcl_request.eml"))["caseId"]
    assert client.get(f"/api/cases/{case_id}").json()["extraction"]["state"] == "not-connected"

    connect()  # 추출 켬 (.env 수정 + 서버 재시작에 해당)
    again = client.post("/api/cases", json=eml("01_lcl_request.eml")).json()
    assert again["duplicate"] and again["case"]["extraction"]["state"] == "not-connected"  # 같은 메일 = 예전 케이스

    res = client.post(f"/api/cases/{case_id}/extract", json={"actor": "kim@leona"})
    assert res.status_code == 200
    case = client.get(f"/api/cases/{case_id}").json()
    assert case["extraction"]["state"] == "done" and case["status"] == "정보부족"
    assert case["extraction"]["trigger"] == "MANUAL_RERUN"


def test_rerun_keeps_manual_fields_and_blocks_double_click(client, connect, session):
    case_id = create(client, eml("01_lcl_request.eml"))["caseId"]
    client.post(f"/api/cases/{case_id}/fields/pol", json={"value": "KRINC"})
    client.post(f"/api/cases/{case_id}/extract", json={})
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case_id}/fields").json()}
    assert fields["pol"]["value"] == "KRINC" and fields["pol"]["method"] == "manual"

    from app import services
    pipeline.start(session, services.get_case(session, case_id), trigger="MANUAL_RERUN", actor=None)  # 실행 중 상태
    assert client.post(f"/api/cases/{case_id}/extract", json={}).status_code == 409
    assert client.post("/api/cases/LQ-2099-0101-999/extract", json={}).status_code == 404



# ------------------------------------------------------------------ LLM에게 빈 칸만 묻기 (CPU 속도, NFR-02)


def test_llm_not_called_when_rules_find_everything(client, connect, monkeypatch):
    monkeypatch.setattr(settings, "llm_scope", "missing")  # 느린 PC용: 빈 칸만 LLM
    llm = FakeLlm({})
    connect(llm, mode="hybrid")
    case = create(client, text_mail(FULL_LCL))
    assert llm.calls == 0 and case["status"] == "계산완료"
    assert case["extraction"]["result"]["llm"]["skipped"] == "규칙으로 모두 찾음"


def test_reply_prose_correction(client, connect, monkeypatch):
    monkeypatch.setattr(settings, "llm_scope", "missing")
    """회신에서 문장으로 고친 값: 룰이 읽으면 룰로, 룰이 못 읽는데 그 항목 이야기가 나오면 LLM에게 다시 묻는다."""
    llm = FakeLlm({})
    connect(llm, mode="hybrid")
    base = FULL_LCL.replace("결제조건: T/T 30% Advance\n", "")
    case_id = create(client, text_mail(base))["caseId"]
    reply(client, case_id, text_mail("총중량은 650kg로 정정합니다.\n결제조건: T/T 30% Advance", subject=f"RE: [{case_id}]"))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case_id}/fields").json()}
    assert fields["grossWeightKg"]["value"] == 650.0 and "메일 2" in fields["grossWeightKg"]["origin"]  # 룰이 직접

    case_id = create(client, text_mail(base, subject="다른 요청"))["caseId"]
    reply(client, case_id, text_mail("총중량은 확인 후 다시 알려드리겠습니다.\n결제조건: T/T", subject=f"RE: [{case_id}]"))
    asked = client.get(f"/api/cases/{case_id}").json()["extraction"]["result"]["llm"]["asked"]
    assert "grossWeight" in asked  # 룰이 값을 못 읽은 언급 → LLM 확인


# ------------------------------------------------------------------ 월간 운임표 테스트 요율 (samples.md 3번)

FCL_SG = (
    "품목: 플라스틱 부품\n수량: 500박스\n총 부피: 48CBM\n총중량: 10,000kg\n선적항: 부산\n도착항: 싱가포르\n"
    "화물 준비일: 2026-11-05\n조건: CIF\n컨테이너: 40HQ 1대\n인보이스 금액: USD 15,000\n결제조건: T/T 30% Advance\n감사합니다."
)


def test_monthly_oft_cif_40hq_quote(client, connect):
    case = create(client, text_mail(FCL_SG))
    assert case["status"] == "계산완료", case["extraction"]["validation"]
    # 해상운임 USD 750 (월간 운임표 40'HC 싱가포르) + 부산 40HQ 부대비용 410,000+10,220+8,000+70,000
    assert case["extraction"]["validation"]["quote"]["totals"] == {"USD": 750.0, "KRW": 498220.0}
    ws = load_workbook(settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.xlsx").active
    assert ws["C17"].value == ": BUSAN - SINGAPORE" and ws["F23"].value == "PER CNTR · INCLUSIVE ISPS(DIRECT)"


def test_monthly_oft_fob_has_no_ocean_freight(client, connect):
    case = create(client, text_mail(FCL_SG.replace("조건: CIF", "조건: FOB")))
    assert case["extraction"]["validation"]["quote"]["totals"] == {"KRW": 498220.0}  # FOB는 해상운임 매수인 부담
    # FOB 예시 견적서 형식: POL만, REQUIRED CNTR, CONTAINER/Q/T/TOTAL, AT COST는 칸 합침
    ws = load_workbook(settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.xlsx").active
    assert (ws["B17"].value, ws["C17"].value) == ("** POL", ": BUSAN, KOREA")
    assert (ws["B19"].value, ws["C19"].value) == ("** REQUIRED CNTR", ": 40HQ' X 1")
    assert (ws["C22"].value, ws["D22"].value, ws["E22"].value) == ("CONTAINER", "Q/T", "TOTAL")
    at_cost = [r for r in range(23, 35) if ws[f"C{r}"].value == "AT COST"]
    assert at_cost and all(f"C{r}:E{r}" in {str(m) for m in ws.merged_cells.ranges} for r in at_cost)


def test_quote_pdf_is_made_with_libreoffice(client, connect, monkeypatch):
    from app import quotation

    if quotation.find_soffice() is None:
        pytest.skip("LibreOffice 없음 — 이 PC에서는 XLSX만 만든다")
    monkeypatch.setattr(settings, "quote_pdf", True)
    case = create(client, text_mail(FULL_LCL))
    quote = case["extraction"]["validation"]["quote"]
    pdf = settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.pdf"
    assert quote["pdf"] and pdf.read_bytes()[:4] == b"%PDF"
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/quote", json={}).json()
    assert sorted(a["filename"] for a in draft["attachments"]) == sorted([pdf.name, f"{case['caseId']}_견적서.xlsx"])


def test_quote_without_libreoffice_uses_builtin_pdf(client, connect, monkeypatch):
    """LibreOffice가 없어도 XLSX와 PDF(내장 엔진)가 함께 만들어진다 (MUST-SHIP ④)."""
    from app import quotation

    monkeypatch.setattr(settings, "quote_pdf", True)
    monkeypatch.setattr(quotation, "find_soffice", lambda: None)
    case = create(client, text_mail(FULL_LCL))
    quote = case["extraction"]["validation"]["quote"]
    assert case["status"] == "계산완료" and quote["file"].endswith(".xlsx")
    assert quote["pdf"] and quote["pdf"].endswith(".pdf")
    assert "LibreOffice" in quote["pdfNote"] and "내장 엔진" in quote["pdfNote"]
    assert (settings.storage_dir / quote["pdf"]).exists()


def test_quote_still_makes_xlsx_if_pdf_fails(client, connect, monkeypatch):
    """PDF를 어떤 방법으로도 못 만들어도 견적은 XLSX로 완료된다 (기존 동작 유지)."""
    from app import quotation, quote_pdf

    def broken(*args, **kwargs):
        raise RuntimeError("pdf engine down")

    monkeypatch.setattr(settings, "quote_pdf", True)
    monkeypatch.setattr(quotation, "find_soffice", lambda: None)
    monkeypatch.setattr(quote_pdf, "render_pdf", broken)
    case = create(client, text_mail(FULL_LCL))
    quote = case["extraction"]["validation"]["quote"]
    assert case["status"] == "계산완료" and quote["file"].endswith(".xlsx")
    assert quote["pdf"] is None and "XLSX만 생성" in quote["pdfNote"]


@pytest.mark.parametrize("container,reason", [
    ("20FT GP", "국내 부대비용"),  # 해상운임은 있지만 20ft 부대비용 자료 없음 → 불완전 견적 금지
])
def test_container_without_local_charges_is_held(client, connect, container, reason):
    case = create(client, text_mail(FCL_SG.replace("40HQ 1대", container).replace("48CBM", "25CBM")
                                    .replace("10,000kg", "8,000kg")))
    hold = case["extraction"]["validation"]["hold"]
    assert case["status"] == "보류" and reason in hold[0]["message"]


@pytest.mark.parametrize("container,port", [("40RF 1대", "싱가포르"), ("40' REEFER", "싱가포르"), ("40RF 1대", "제벨알리")])
def test_reefer_container_is_held_for_staff(client, connect, container, port):
    """냉동 컨테이너 요청은 요율이 있든 없든 자동 견적하지 않고 담당자 처리로 넘긴다 (정의서 3장 범위 제외, FR-308)."""
    case = create(client, text_mail(FCL_SG.replace("40HQ 1대", container).replace("싱가포르", port)))
    hold = case["extraction"]["validation"]["hold"]
    assert case["status"] == "보류" and hold[0]["code"] == "SPECIAL_CARGO" and "40RF" in hold[0]["message"]
    assert case["extraction"]["validation"]["quote"] is None


# ------------------------------------------------------------------ 컨테이너 대수 (금액이 대수만큼 늘어야 한다)


def test_stated_container_count_multiplies_per_container_charges(client, connect):
    case = create(client, text_mail(FCL_SG.replace("40HQ 1대", "40HQ 2대").replace("48CBM", "96CBM")))
    quote = case["extraction"]["validation"]["quote"]
    # 해상운임 750×2, THC 410,000×2 + WFG 10,220×2 + 씰 8,000×2 + 서류비 70,000×1(B/L당)
    assert quote["totals"] == {"USD": 1500.0, "KRW": 926440.0} and quote["containerCount"] == 2
    ws = load_workbook(settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.xlsx").active
    assert ws["C18"].value.endswith("(40HQ X 2)")


def test_missing_count_is_estimated_from_volume_and_explained(client, connect):
    body = FCL_SG.replace("컨테이너: 40HQ 1대", "컨테이너: 40HQ").replace("48CBM", "100CBM").replace("10,000kg", "20,000kg")
    case = create(client, text_mail(body))
    quote = case["extraction"]["validation"]["quote"]
    assert quote["containerCount"] == 2 and "100CBM" in quote["containerCountNote"]  # 76CBM×0.88 넘음 → 2대
    ws = load_workbook(settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.xlsx").active
    assert any("Container q'ty : 2 x 40HQ (estimated from cargo 100CBM / 20,000KG)" in (ws[f"B{r}"].value or "")
               for r in range(39, 48))


@pytest.mark.parametrize("stated,asked", [("1대", True), ("2대", False)])
def test_over_weight_uses_weight_per_container(client, connect, stated, asked):
    body = FCL_SG.replace("40HQ 1대", f"40HQ {stated}").replace("10,000kg", "28,000kg")
    case = create(client, text_mail(body))
    codes = {i["code"] for i in case["extraction"]["validation"]["issues"]}
    assert ("OVER_WEIGHT" in codes) is asked


def test_over_volume_with_stated_count(client, connect):
    """FR-302: 적힌 대수로 실을 수 없는 부피면 묻는다 (40HQ 1대 = 76CBM)."""
    case = create(client, text_mail(FCL_SG.replace("48CBM", "90CBM")))
    codes = {i["code"] for i in case["extraction"]["validation"]["issues"]}
    assert "OVER_VOLUME" in codes and case["status"] == "정보부족"
    fits = create(client, text_mail(FCL_SG.replace("48CBM", "90CBM").replace("40HQ 1대", "40HQ 2대")))
    assert "OVER_VOLUME" not in {i["code"] for i in fits["extraction"]["validation"]["issues"]}


def test_container_without_weight_limit_is_info_only(client, connect):
    """허용중량 자료가 없는 컨테이너(40RF)는 검사를 못 했다고 담당자에게만 알린다 (화주 문항 아님)."""
    case = create(client, text_mail(FCL_SG.replace("40HQ 1대", "40RF 1대")))
    note = [i for i in case["extraction"]["validation"]["issues"] if i["code"] == "NO_WEIGHT_LIMIT"]
    assert note and note[0]["severity"] == "Info" and note[0]["question"] is None


@pytest.mark.parametrize("commodity,held", [("냉동 만두", True), ("스프링노트", False)])
def test_special_cargo_is_held(client, connect, commodity, held):
    """FR-308: 냉동·특수화물 키워드가 있으면 자동 견적을 멈추고 담당자 처리로 넘긴다."""
    case = create(client, text_mail(FULL_LCL.replace("스프링노트", commodity)))
    codes = [h["code"] for h in case["extraction"]["validation"]["hold"]]
    assert ("SPECIAL_CARGO" in codes) is held and (case["status"] == "보류") is held


def test_validation_rules_are_data(client, connect, monkeypatch):
    """질문 문구·필수 항목을 데이터 파일로 바꾸면 코드 수정 없이 반영된다 (NFR-06)."""
    custom = dict(validation.rules())
    custom["questions"] = {**custom["questions"], "cargoReadyDate": "출고 예정일을 알려주세요 (테스트 문구)"}
    custom["required"] = [f for f in custom["required"] if f != "paymentTerm"]
    monkeypatch.setattr(validation, "rules", lambda: custom)
    questions = "\n".join(create(client, eml("01_lcl_request.eml"))["extraction"]["validation"]["questions"])
    assert "출고 예정일을 알려주세요 (테스트 문구)" in questions and "결제 조건" not in questions


def test_busan_singapore_lcl_test_rate(client, connect):
    case = create(client, text_mail(FULL_LCL.replace("인천", "부산").replace("시드니", "싱가포르")))
    assert case["status"] == "계산완료", case["extraction"]["validation"]
    assert case["extraction"]["validation"]["quote"]["totals"] == {"USD": 500.0, "KRW": 140203.0}


def test_invoice_questions_not_duplicated_and_split_reply_merges(client, connect):
    case = create(client, eml("01_lcl_request.eml"))
    questions = case["extraction"]["validation"]["questions"]
    invoice = [q for q in questions if "인보이스" in q]
    assert len(invoice) == 1 and "금액과 통화" in invoice[0]  # 통화 질문이 따로 또 나가지 않는다

    after = reply(client, case["caseId"], text_mail("인보이스 금액: 8,500\n통화: USD", subject=f"RE: [{case['caseId']}]"))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["invoiceValue"]["value"] == 8500 and fields["ccy"]["value"] == "USD"
    assert not [q for q in after["extraction"]["validation"]["questions"] if "인보이스" in q]


def test_currency_only_question_when_amount_known(client, connect):
    case = create(client, text_mail(FULL_LCL.replace("USD 3,000", "3,000")))
    invoice = [q for q in case["extraction"]["validation"]["questions"] if "인보이스" in q]
    assert len(invoice) == 1 and "통화" in invoice[0] and "금액과" not in invoice[0]


def test_reply_answering_by_question_numbers(client, connect):
    """보완 요청을 보낸 뒤 화주가 번호에 맞춰 값만 적어 회신해도 채워진다."""
    case = create(client, eml("01_lcl_request.eml"))
    draft = client.post(f"/api/cases/{case['caseId']}/drafts/supplement", json={}).json()
    questions = case["extraction"]["validation"]["questions"]
    assert draft["htmlBody"].count("<li>") == len(questions)
    order = {"선적항": "인천", "화물 준비일": "2026-10-20", "인보이스": "USD 8,500", "결제 조건": "T/T 30% Advance",
             "조건을": "CIF", "부피": "0.38CBM"}
    lines = []
    for n, q in enumerate(questions, start=1):
        answer = next((a for key, a in order.items() if key in q), "확인 중")
        lines.append(f"{n}. {answer}")
    after = reply(client, case["caseId"], text_mail("\n".join(lines), subject=f"RE: [{case['caseId']}] 보완"))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["pol"]["value"] == "KRINC" and fields["cargoReadyDate"]["value"] == "2026-10-20"
    assert fields["invoiceValue"]["value"] == 8500 and fields["paymentTerm"]["value"] == "T/T 30% Advance"
    assert after["status"] == "계산완료", after["extraction"]["validation"]


def test_unreadable_value_is_asked_with_original_text(client, connect):
    body = FULL_LCL.replace("예상 총중량: 500kg", "예상 총중량: 오백 킬로 정도")
    case = create(client, text_mail(body))
    q = [x for x in case["extraction"]["validation"]["questions"] if "총중량" in x]
    assert len(q) == 1 and "「예상 총중량: 오백 킬로 정도」" in q[0] and "500kg 형식" in q[0]
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["grossWeightKg"]["status"] == "review" and fields["grossWeightKg"]["candidate"] == "오백 킬로 정도"

    report = client.get("/api/extraction/review-report").json()
    assert report["unparsed"]["grossWeightKg"]["examples"][0]["case"] == case["caseId"]  # 놓친 표현 모아 보기


def test_unsupported_incoterms_gets_a_proper_question(client, connect):
    case = create(client, text_mail(FULL_LCL.replace("조건: CIF", "조건: FCA")))
    q = [x for x in case["extraction"]["validation"]["questions"] if "FCA" in x]
    assert len(q) == 1 and "FOB·CIF·DDP·EXW" in q[0]



# ------------------------------------------------------------------ LLM 중심 추출 (LLM이 모든 항목, 규칙은 교차 확인)


class LineLlm:
    """실제 LLM처럼 '정리된 값 + 줄 번호'로 답하는 가짜 LLM. answers = {항목: (값, 근거 줄에 든 글자)}"""

    def __init__(self, answers: dict, multi=False):
        self.answers, self.multi, self.prompts = answers, multi, []

    def complete_json(self, system, user, schema):
        self.prompts.append(user)
        lines = {int(m[1]): m[2] for m in __import__("re").finditer(r"^L(\d+): (.*)$", user, __import__("re").M)}
        fields = {}
        for key, (value, where) in self.answers.items():
            hit = [n for n, text in lines.items() if any(w in text for w in where.split("|"))]
            if hit:
                fields[key] = {"v": value, "l": hit[-1:] if "|" not in where else hit}
        return {"fields": fields, "multipleItems": self.multi}


PROSE = ("안녕하세요, 그린오피스 최서연입니다.\n인천에서 시드니로 LCL로 보낼 노트가 열 박스 있어요.\n"
         "한 박스가 가로 310 세로 450 높이 270mm이고 무게는 다 합쳐 오백 킬로 정도예요.\n"
         "10월 20일쯤 준비되고 CIF로 부탁드려요.\n금액은 3,000이고 통화는 USD입니다.\n결제는 티티 30% 선결제로 할게요.")


def test_llm_centric_reads_free_prose(client, connect):
    llm = LineLlm({
        "commodity": ("노트", "노트가"), "quantity": ("10 CTNS", "열 박스"),
        "boxDimensions": ("310 x 450 x 270 mm", "가로"), "grossWeight": ("500 kg", "오백"),
        "portOfLoading": ("인천", "인천에서"), "portOfDischarge": ("시드니", "시드니로"),
        "cargoReadyDate": ("2026-10-20", "10월 20일"), "incoterms": ("CIF", "CIF로"), "container": ("LCL", "LCL"),
        "invoiceValue": ("3000 USD", "금액은"), "paymentTerm": ("T/T 30% Advance", "티티"),
    })
    connect(llm, mode="hybrid")
    case = create(client, text_mail(PROSE))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert "직전 질문" not in llm.prompts[0] and "L2:" in llm.prompts[0]
    for name, value in [("commodity", "노트"), ("qty", 10), ("boxL", 310), ("grossWeightKg", 500.0), ("pol", "KRINC"),
                        ("pod", "AUSYD"), ("incoterms", "CIF"), ("containerType", "LCL"), ("invoiceValue", 3000),
                        ("ccy", "USD")]:
        assert fields[name]["value"] == value, name
    assert fields["paymentTerm"]["value"] == "T/T 30% Advance"  # '티티 30% 선결제'를 정리 (원문에 결제 방식 단어 있음)
    assert fields["cargoReadyDate"]["value"] == "2026-10-20" and fields["cargoReadyDate"]["score"] < 1.0


def test_llm_reads_numbered_reply_with_previous_questions(client, connect):
    llm = LineLlm({"portOfLoading": ("인천", "1. 인천"), "cargoReadyDate": ("2026-10-20", "2. 2026"),
                   "paymentTerm": ("T/T", "3. T/T")})
    connect(llm, mode="hybrid")
    case_id = create(client, text_mail(FULL_LCL.replace("선적항: 인천\n", "").replace("화물 준비일: 2026-10-20\n", "")
                                       .replace("결제조건: T/T 30% Advance\n", "")))["caseId"]
    client.post(f"/api/cases/{case_id}/drafts/supplement", json={})
    reply(client, case_id, text_mail("1. 인천\n2. 2026-10-20\n3. T/T", subject=f"RE: [{case_id}]"))
    assert "직전 질문" in llm.prompts[-1] and "1) 선적항" in llm.prompts[-1]
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case_id}/fields").json()}
    assert (fields["pol"]["value"], fields["cargoReadyDate"]["value"], fields["paymentTerm"]["value"]) == \
        ("KRINC", "2026-10-20", "T/T")


def test_small_model_mistakes_are_not_saved(client, connect):
    llm = LineLlm({"portOfLoading": ("인천", "픽업지"),  # 픽업지를 선적항으로 착각
                   "grossWeight": ("50 kg", "총중량")})  # 숫자를 잘못 옮김
    connect(llm, mode="hybrid")
    case = create(client, eml("01_lcl_request.eml"))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["pol"]["value"] is None  # 규칙이 '픽업지'로 읽은 줄 → 선적항으로 저장 안 함
    assert fields["grossWeightKg"]["value"] == 500.0  # 원문 숫자와 다른 LLM 값은 버리고 규칙 값 유지



def test_llm_payment_term_needs_payment_word_in_source(client, connect):
    llm = LineLlm({"paymentTerm": ("L/C 30 days", "30일")})  # 원문엔 결제 방식이 없는데 L/C로 지어냄
    connect(llm, mode="hybrid")
    case = create(client, text_mail(FULL_LCL.replace("결제조건: T/T 30% Advance", "대금은 30일 후에 드려요")))
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["paymentTerm"]["value"] is None and fields["paymentTerm"]["status"] == "review"


def test_llm_wrong_line_for_invoice_is_not_asked_to_customer(client, connect):
    """실제 사례: 작은 모델이 '예상 총중량: 약 360kg' 줄을 인보이스 금액으로 짚음 → 화주 질문에 그 줄이 나오면 안 된다."""
    body = ("품목: 문구류\n박스 수량: 8박스\n박스 1개 규격: 310 × 450 × 270mm\n예상 총중량: 약 360kg\n선적항: 인천\n"
            "도착항: 시드니\n조건: CIF\n선적방식: LCL")
    llm = LineLlm({"invoiceValue": ("360 kg", "총중량"), "grossWeight": ("360 kg", "총중량")})
    connect(llm, mode="hybrid")
    case = create(client, text_mail(body))
    questions = case["extraction"]["validation"]["questions"]
    invoice = [q for q in questions if "인보이스" in q]
    assert len(invoice) == 1 and "총중량" not in invoice[0] and "금액과 통화" in invoice[0]
    assert not any("「" in q for q in questions)  # LLM이 짚은 줄을 인용하는 질문 없음
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["invoiceValue"]["status"] == "missing" and fields["grossWeightKg"]["value"] == 360.0


def test_llm_unreadable_value_is_for_staff_not_quoted_to_customer(client, connect):
    llm = LineLlm({"grossWeight": ("반", "무게는")})
    connect(llm, mode="hybrid")
    case = create(client, text_mail(FULL_LCL.replace("예상 총중량: 500kg", "무게는 대략 반 정도")))
    q = [x for x in case["extraction"]["validation"]["questions"] if "총중량" in x]
    assert q and "「" not in q[0]  # 일반 질문
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["grossWeightKg"]["status"] == "review" and fields["grossWeightKg"]["candidate"] == "반"  # 담당자에겐 보임


def test_question_particles_follow_final_consonant():
    from app.validation import josa

    assert ("총중량" + josa("총중량", "을", "를"), "포장 형태" + josa("포장 형태", "을", "를")) == ("총중량을", "포장 형태를")
    assert "360kg" + josa("360kg", "이라고", "라고") == "360kg이라고"


def test_reply_merge_quirks_from_real_test(client, connect):
    """실제 테스트(②번 메일)에서 본 어색한 점: 규격 '(확인 필요)', 픽업지=회사명, 담당자=회신 계정 이름."""
    original = ("품목: 문구류(노트)\n박스 수량: 20박스\n박스 규격: 40 x 30 x 25 cm\n총중량: 180kg\n도착항: 싱가포르\n"
                "선적방식: LCL\n\n감사합니다.\n이지은 과장\n서울문구 주식회사")
    llm = LineLlm({"boxDimensions": ("40 x 30 x 25", "박스 규격")})  # 작은 모델이 단위를 빼먹음 → 확신 낮은 값
    connect(llm, mode="hybrid")
    first = client.post("/api/cases", json=text_mail(original, sender="jieun@seoul.example.com")).json()["case"]
    case_id = first["caseId"]
    client.post(f"/api/cases/{case_id}/drafts/supplement", json={})
    reply_mail = text_mail("1. 부산\n2. 2026-10-27\n3. CIF\n4. USD 2,400\n5. T/T 100% Advance", subject=f"RE: [{case_id}]")
    reply_mail["snapshot"]["from"] = {"name": "김민종", "email": "minjong@gmail.example.com"}
    case = reply(client, case_id, reply_mail)
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case_id}/fields").json()}
    assert fields["boxL"]["status"] == "filled" and fields["boxL"]["value"] == 400  # 확신 낮은 LLM 값이 흔들지 않음
    assert fields["contactName"]["value"] == "이지은 과장"  # 서명 이름이 회신 계정 이름보다 우선
    assert "pickupLocation" not in case["extraction"]["result"]["llm"]["asked"]
    assert case["status"] == "계산완료"


def test_attachment_values_flow_through_outlook_eml(client, connect):
    """패널이 보내는 .eml(첨부 포함) → 서버가 첨부 저장 → 추출이 첨부(PDF·XLSX)에서 값을 읽는다 (FR-204)."""
    import io
    from email.message import EmailMessage

    from openpyxl import Workbook

    from .test_extraction import tiny_pdf

    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = "부산-싱가포르 LCL 견적 (패킹리스트 첨부)", "buyer@x.co", "quote@leona.example.com"
    msg["Message-ID"] = "<attach-test@x.co>"
    msg.set_content("품목: 전자부품\n선적항: 부산\n도착항: 싱가포르\n조건: CIF\n선적방식: LCL\n화물 준비일: 2026-10-25\n"
                    "결제조건: T/T\n첨부 참고 부탁드립니다.")
    msg.add_attachment(tiny_pdf(["Package: 12 CTNS", "Carton Size: 60 x 40 x 35 cm", "Gross Weight: 240 KGS"]),
                       maintype="application", subtype="pdf", filename="packing_list.pdf")
    wb = Workbook()
    wb.active.append(["Invoice Value", "USD 6,000"])
    buf = io.BytesIO()
    wb.save(buf)
    msg.add_attachment(buf.getvalue(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="invoice.xlsx")

    body = {"emlBase64": base64.b64encode(msg.as_bytes()).decode(), "actor": "quote@leona.example.com"}
    case = create(client, body)
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    assert fields["grossWeightKg"]["value"] == 240.0 and "첨부 packing_list.pdf" in fields["grossWeightKg"]["origin"]
    assert fields["invoiceValue"]["value"] == 6000 and "첨부 invoice.xlsx" in fields["invoiceValue"]["origin"]
    assert case["status"] == "계산완료"  # 본문 + 첨부 두 개로 견적까지


# ------------------------------------------------------------------ 발주 측 샘플 요율표 (docs/양식/샘플_요율표.xlsx)

CLIENT_RATES = Path(__file__).resolve().parents[2] / "docs" / "양식" / "샘플_요율표.xlsx"
CLIENT_CASE = ("품목: 자동차 부품\n수량: 200박스\n총 부피: 25CBM\n총중량: 8,000kg\n선적항: 상하이\n도착항: 부산\n"
               "화물 준비일: 2026-10-05\n조건: CIF\n컨테이너: 20GP 1대\n인보이스 금액: USD 15,000\n결제조건: T/T\n{carrier}감사합니다.")


@pytest.fixture
def client_rates(client, connect):
    res = client.post("/api/rates/import", files={"file": ("샘플_요율표.xlsx", CLIENT_RATES.read_bytes(), "application/octet-stream")})
    assert res.status_code == 200 and res.json()["rates"] == 35, res.json()


# 인코텀즈별 비용 (발주 측 회신 2026-10-06): 수출자 기준 EXW 전부 제외 / FOB 출발지만 / CIF 출발지+해상운임+보험 / DDP 전부.
# 수입자는 그 반대. 이 요율표는 중국·베트남·태국 → 부산·인천이라 화주가 수입자다 (도착항이 국내).
# 단, 수입 CIF는 요율표 '조회예시' 시트대로 전부 넣는다.
# 금액: 해상운임 HMM 450·SITC 430, OTHC 120, DTHC 130, DOC 30, B/L 30, 취급 50, BAF(출발지) 50


def test_client_import_fob_has_freight_and_destination(client, client_rates):
    """수입 FOB: 해상운임 + 도착지 비용 (출발지 비용·보험은 수출자 몫). 선사·스케줄·근거 행은 '조회예시' 시트와 같다."""
    case = create(client, text_mail(CLIENT_CASE.format(carrier="선사: HMM\n").replace("조건: CIF", "조건: FOB")))
    quote = case["extraction"]["validation"]["quote"]
    assert case["status"] == "계산완료", case["extraction"]["validation"]
    assert quote["tradeRole"] == "IMPORT" and quote["carrier"] == "HMM"
    assert quote["totals"] == {"USD": 690.0}  # 450 + DTHC 130 + DOC 30 + B/L 30 + 취급 50
    assert (quote["etd"], quote["eta"]) == ("2026-10-06", "2026-10-08")  # 조회예시 스케줄
    lines = {line["charge"]: (line["amount"], line["source"]) for line in quote["lines"]}
    assert lines["OCEAN FREIGHT"] == (450.0, "OceanFreight#1")  # 근거 행까지 예시와 같다
    assert not {"ORIGIN THC", "BUNKER SURCHARGE", "CARGO INSURANCE"} & set(lines)
    fields = {f["field"]: f for f in client.get(f"/api/cases/{case['caseId']}/fields").json()}
    detail = client.get(f"/api/cases/{case['caseId']}").json()
    ws = load_workbook(settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.xlsx").active
    remarks = [ws[f"B{r}"].value or "" for r in range(39, 48)]
    assert any("ETD / ETA : 2026-10-06 / 2026-10-08 (HMM)" in r for r in remarks)
    assert any("FREE TIME : DET 7 days / DEM 7 days" in r for r in remarks)
    assert "* Destination charges are for consignee's account." in remarks  # FOB·CIF 견적서 고지
    assert fields["pol"]["value"] == "CNSHA" and detail["status"] == "계산완료"


def test_client_rates_lowest_carrier_when_not_specified(client, client_rates):
    case = create(client, text_mail(CLIENT_CASE.format(carrier="").replace("조건: CIF", "조건: FOB")))
    quote = case["extraction"]["validation"]["quote"]
    assert quote["carrier"] == "SITC" and quote["totals"] == {"USD": 670.0}  # 선사 미지정 → 최저가 SITC 430
    assert (quote["etd"], quote["eta"]) == ("2026-10-07", "2026-10-09")


def test_client_rates_follow_cargo_ready_date(client, client_rates):
    """GRI 인상: 선적일이 10/16 이후면 후반기 행(HMM 520)을 쓴다 — '오늘'이 아니라 선적일 기준."""
    body = CLIENT_CASE.format(carrier="선사: HMM\n").replace("조건: CIF", "조건: FOB").replace("2026-10-05", "2026-10-20")
    quote = create(client, text_mail(body))["extraction"]["validation"]["quote"]
    assert quote["rateDate"] == "2026-10-20" and quote["totals"] == {"USD": 760.0}  # 690 - 450 + 520


def test_client_import_cif_follows_lookup_example(client, client_rates):
    """수입 CIF는 발주 측 '조회예시' 시트대로 출발지·해상운임·도착지·보험을 전부 넣는다 (CNSHA→KRPUS 20GP, HMM).

    예시 합계는 890이지만 보험료를 계약금액 × 110% × 0.2% = 33으로 계산해(발주 측 회신 2026-10-06) 893이 된다.
    """
    case = create(client, text_mail(CLIENT_CASE.format(carrier="선사: HMM\n")))
    quote = case["extraction"]["validation"]["quote"]
    assert case["status"] == "계산완료", case["extraction"]["validation"]
    lines = {line["charge"]: line["amount"] for line in quote["lines"]}
    assert lines == {"OCEAN FREIGHT": 450.0, "ORIGIN THC": 120.0, "DESTINATION THC": 130.0, "DOCUMENTATION FEE": 30.0,
                     "B/L FEE": 30.0, "HANDLING CHARGE": 50.0, "CARGO INSURANCE": 33.0, "BUNKER SURCHARGE": 50.0}
    assert quote["totals"] == {"USD": 893.0} and quote["carrier"] == "HMM"
    assert (quote["etd"], quote["eta"]) == ("2026-10-06", "2026-10-08")  # 조회예시 스케줄


def test_client_import_exw_has_everything(client, client_rates):
    """수입 EXW: 수출자가 아무것도 내지 않는다 → 출발지·해상운임·도착지 전부 (보험은 요율표가 CIF에만 부과)."""
    body = CLIENT_CASE.format(carrier="선사: HMM\n").replace("조건: CIF", "조건: EXW")
    quote = create(client, text_mail(body))["extraction"]["validation"]["quote"]
    lines = {line["charge"] for line in quote["lines"]}
    assert {"OCEAN FREIGHT", "ORIGIN THC", "DESTINATION THC", "BUNKER SURCHARGE"} <= lines
    assert quote["totals"] == {"USD": 860.0}  # 450 + 120 + 130 + 30 + 30 + 50 + BAF 50


def test_nothing_to_quote_is_held(client, client_rates):
    """수입 DDP·수출 EXW: 상대방이 전부 부담해 견적에 넣을 비용이 없다 → 견적서를 만들지 않고 담당자에게 넘긴다."""
    ddp = create(client, text_mail(CLIENT_CASE.format(carrier="").replace("조건: CIF", "조건: DDP")))
    hold = ddp["extraction"]["validation"]["hold"]
    assert ddp["status"] == "보류" and "DDP 조건의 수입자" in hold[0]["message"]
    exw = create(client, text_mail(FULL_LCL.replace("조건: CIF", "조건: EXW")))  # 인천 → 시드니 = 수출
    hold = exw["extraction"]["validation"]["hold"]
    assert exw["status"] == "보류" and "EXW 조건의 수출자" in hold[0]["message"]


def test_incoterm_rule_table():
    """수출자 기준 규칙과 수입자(반대) — DB 없이 규칙 표만 확인."""
    from types import SimpleNamespace as NS

    from app import quotation as q

    assert q.trade_role(NS(pol="KRPUS", pod="SGSIN")) == "EXPORT" and q.trade_role(NS(pol="CNSHA", pod="KRPUS")) == "IMPORT"
    assert q.trade_role(NS(pol="CNSHA", pod="SGSIN")) is None
    assert q.included_parts("EXW", "EXPORT") == set() and q.included_parts("FOB", "EXPORT") == {"ORIGIN"}
    assert q.included_parts("CIF", "EXPORT") == {"ORIGIN", "FREIGHT", "INSURANCE"}
    assert q.included_parts("DDP", "EXPORT") == {"ORIGIN", "FREIGHT", "INSURANCE", "DEST"}
    assert q.included_parts("FOB", "IMPORT") == {"FREIGHT", "INSURANCE", "DEST"} and q.included_parts("DDP", "IMPORT") == set()
    assert q.included_parts("EXW", "IMPORT") == q.included_parts("CIF", "IMPORT") == {"ORIGIN", "FREIGHT", "INSURANCE", "DEST"}
    assert q.incoterm_remarks("DDP", "en") == ["* Customs duty and VAT are not included and will be charged separately."]
    assert q.incoterm_remarks("DDP", "ko") == ["* 관세·부가세 별도"] and q.incoterm_remarks("EXW", "en") == []
    assert q.incoterm_remarks("CIF", "ko") == ["* 도착지 비용 수하인 부담"]


def test_insurance_is_110_percent_of_invoice_value():
    """보험료 = 계약금액 × 110% × 보험요율, 최저 금액 적용 (발주 측 회신 2026-10-06)."""
    from types import SimpleNamespace as NS

    from app import quotation as q

    ins = NS(charge_code="INS", rate_pct=0.002, min_amount=30, currency="USD")
    assert q.percent_amount(ins, NS(invoiceValue=15000, ccy="USD")) == (33.0, "INV.V x 110% x 0.2% (MIN USD 30)")
    assert q.percent_amount(ins, NS(invoiceValue=10000, ccy="USD"))[0] == 30.0  # 22 → 최저 30
    other = NS(charge_code="HANDLING", rate_pct=0.01, min_amount=None, currency="USD")
    assert q.percent_amount(other, NS(invoiceValue=1000, ccy="USD")) == (10.0, "INV.V x 1%")  # 보험이 아니면 110% 없음


def test_validity_stops_at_rate_end_date(client, client_rates, monkeypatch):
    """유효기간 = 발행일 + 14일, 견적에 쓴 요율의 종료일이 더 빠르면 그날까지 (HMM 전반기 운임은 10/15까지)."""
    monkeypatch.setattr(validation, "today_kst", lambda: date(2026, 10, 10))
    body = CLIENT_CASE.format(carrier="선사: HMM\n").replace("조건: CIF", "조건: FOB").replace("2026-10-05", "2026-10-12")
    quote = create(client, text_mail(body))["extraction"]["validation"]["quote"]
    assert quote["validUntil"] == "2026-10-15"  # 10/10 + 14일 = 10/24 이지만 요율이 10/15에 끝난다


@pytest.mark.parametrize("days,valid_until", [(14, "2026-10-13"), (7, "2026-10-06"), (1, "2026-09-30"),
                                              (30, "2026-10-13"), (0, "2026-09-30")])
def test_quote_validity_days_is_selectable(client, client_rates, monkeypatch, days, valid_until):
    """FR-507: 견적 유효기간은 설정(QUOTE_VALIDITY_DAYS)으로 1~14일 중에서 고른다. 범위를 벗어나면 1 또는 14."""
    monkeypatch.setattr(settings, "quote_validity_days", days)
    case = create(client, text_mail(CLIENT_CASE.format(carrier="")))
    quote = case["extraction"]["validation"]["quote"]
    assert quote["validUntil"] == valid_until  # 테스트의 '오늘'은 2026-09-29
    ws = load_workbook(settings.storage_dir / case["caseId"] / "quotes" / f"{case['caseId']}_견적서.xlsx").active
    assert any(f"* VALIDITY : {valid_until}" == (ws[f"B{r}"].value or "") for r in range(39, 48))


def test_client_rates_bad_file_changes_nothing(client, client_rates):
    import io

    from openpyxl import load_workbook as load

    wb = load(CLIENT_RATES)
    wb["OceanFreight"]["D5"] = "SIDNEY"  # 첫 운임 행의 pod를 틀리게
    buf = io.BytesIO()
    wb.save(buf)
    res = client.post("/api/rates/import", files={"file": ("bad.xlsx", buf.getvalue(), "application/octet-stream")})
    assert res.status_code == 400 and any("OceanFreight 5행 pod" in e for e in res.json()["detail"]["errors"])
    assert client.get("/api/rates/summary").json()  # 기존 요율은 그대로
