"""첨부 텍스트 추출 (FR-204, S): PDF·XLSX(패킹리스트 등)의 글자를 읽어 추출 대상에 넣는다.

스캔 이미지 PDF(글자 없음)는 빈 문자열 — OCR은 선택 사항이라 넣지 않았다.
표는 '이름<TAB>값' 줄로 바꿔 rules.py가 읽게 한다.
  - 세로 표 (이름 | 값): 그대로 한 줄씩
  - 가로 표 (제목 줄 Description | Q'ty | CBM | G.W  … 아래 TOTAL 줄): 제목과 합계 줄(없으면 값 줄이 하나일 때 그 줄)을 짝지음
  - PDF는 화면 배치대로 읽는다(layout) — 기본 방식은 표를 '값 칸 전부 → 이름 칸 전부'로 흩어 놓아 짝이 깨진다
"""

import csv
import io
import logging
import re
from pathlib import Path

log = logging.getLogger(__name__)
MAX_CHARS = 8_000  # 60초 안에 끝나도록 첨부 1개당 길이 제한
SUPPORTED = {".pdf", ".xlsx", ".xlsm", ".csv", ".txt"}


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        return ""
    try:
        if suffix == ".pdf":
            text = _pdf(data)
        elif suffix in (".xlsx", ".xlsm"):
            text = _xlsx(data)
        elif suffix == ".csv":
            text = "\n".join("\t".join(row) for row in csv.reader(io.StringIO(_decode(data))))
        else:
            text = _decode(data)
    except Exception as error:  # 깨진 첨부 하나 때문에 추출 전체가 실패하지 않게
        log.warning("attachment text failed %s: %s", filename, error)
        return ""
    return text[:MAX_CHARS]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader  # BSD-3

    reader = PdfReader(io.BytesIO(data))
    rows: list[list[str]] = []
    for page in reader.pages[:10]:
        try:
            text = page.extract_text(extraction_mode="layout") or ""  # 같은 줄의 글자를 한 줄로 (표 짝 유지)
        except Exception:
            text = page.extract_text() or ""
        for line in text.split("\n"):
            cells = [c.strip() for c in re.split(r"\s{3,}|\t", line) if c.strip()]  # 넓은 공백 = 칸 구분
            if cells:
                rows.append(cells)
    return "\n".join(table_lines(rows))


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook  # MIT

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows: list[list[str]] = []
    for ws in wb.worksheets[:5]:
        for row in ws.iter_rows(values_only=True, max_row=300):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(cells)
    return "\n".join(table_lines(rows))


_TOTAL = re.compile(r"^(total|totals|합계|총계|계|sum|grand total)\b", re.I)


def _label_field(cell: str) -> str | None:
    """칸이 라벨이면 그 항목 이름 (숫자가 든 칸은 값)."""
    from .rules import label_scores

    if re.search(r"\d", cell):
        return None
    scores = label_scores(cell)
    return scores[0][0] if scores and scores[0][1] >= 0.9 else None


def _is_label(cell: str) -> bool:
    return _label_field(cell) is not None


def _is_header(cells: list[str]) -> bool:
    """가로 표 제목 줄: 칸 3개 이상, 대부분 라벨이고 서로 다른 항목이 2개 이상 ('품목 | (Commodity)'는 같은 항목이라 아님)."""
    fields = [_label_field(c) for c in cells]
    distinct = {f for f in fields if f}
    return len(cells) >= 3 and len(distinct) >= 2 and sum(f is not None for f in fields) >= len(cells) - 1


def table_lines(rows: list[list[str]]) -> list[str]:
    """표의 칸들 → 규칙이 읽는 '이름<TAB>값' 줄. 원래 줄도 남겨 LLM이 문맥을 보게 한다."""
    out: list[str] = []
    i = 0
    while i < len(rows):
        cells = rows[i]
        if _is_header(cells):  # 가로 표 제목 줄
            body = []
            j = i + 1
            while j < len(rows) and not _is_header(rows[j]):
                body.append(rows[j])
                j += 1
            total = next((r for r in body if _TOTAL.match(r[0])), None)
            pick = total or (body[0] if len(body) == 1 else None)  # 품목 줄이 여럿이고 합계가 없으면 짝짓지 않음
            out.append("\t".join(cells))
            out += ["\t".join(r) for r in body]
            if pick:
                values = pick[1:] if total else pick  # 합계 줄의 'TOTAL' 칸은 값이 아니다
                offset = len(cells) - len(values)  # 오른쪽 끝을 맞춘다 (합계 줄은 앞쪽 No·품목명 칸이 비어 있다)
                for k, value in enumerate(values):
                    label = cells[k + offset] if 0 <= k + offset < len(cells) else None
                    if label and _is_label(label):
                        out.append(f"{label}\t{value}")
            i = j
            continue
        if len(cells) >= 2:  # 세로 표: 앞 칸들 = 이름, 마지막 칸 = 값 ('품목 (Commodity) | 전자부품')
            out.append(" ".join(cells[:-1]) + "\t" + cells[-1])
        else:
            out.append(cells[0])
        i += 1
    return out
