"""견적서 PDF 내장 엔진 — LibreOffice가 없거나 변환에 실패할 때 쓴다 (MUST-SHIP ④: PDF·XLSX).

quotation.render_xlsx()가 채운 양식 XLSX를 그대로 읽어 같은 배치로 그린다.
→ 문구·비용 행·비고는 XLSX와 항상 같다 (PDF용 내용을 따로 만들지 않는다).

한글 폰트는 PDF에 내장한다 (내장하지 않으면 고객 PC·모바일 뷰어에서 한글이 깨질 수 있음).
  Windows: 맑은 고딕(malgun.ttf / malgunbd.ttf) 자동 사용
  그 밖: 환경 변수 LEONA_PDF_FONT=<TTF 경로> 또는 나눔고딕 설치
"""

from __future__ import annotations

import io
import logging
import os
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.pdfmetrics import registerFontFamily
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph

log = logging.getLogger(__name__)

# (본문 TTF, 굵은 TTF) 후보 — 앞에서부터 있는 것을 쓴다
_FONTS = [
    (os.environ.get("LEONA_PDF_FONT", ""), os.environ.get("LEONA_PDF_FONT_BOLD", "")),
    ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),
    ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
    ("/Library/Fonts/NanumGothic.ttf", "/Library/Fonts/NanumGothicBold.ttf"),
    ("/System/Library/Fonts/AppleSDGothicNeo.ttc", ""),
    ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", ""),
]
_font_name: str | None = None


def _font() -> str:
    """한글 폰트 등록 (처음 한 번). 굵은 글꼴이 없으면 본문 글꼴로 대신한다."""
    global _font_name
    if _font_name:
        return _font_name
    for regular, bold in _FONTS:
        if not regular or not Path(regular).exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("LeonaKR", regular, subfontIndex=0))
        except Exception:  # noqa: BLE001 — 읽을 수 없는 폰트면 다음 후보
            continue
        bold_name = "LeonaKR"
        if bold and Path(bold).exists():
            try:
                pdfmetrics.registerFont(TTFont("LeonaKR-Bold", bold, subfontIndex=0))
                bold_name = "LeonaKR-Bold"
            except Exception:  # noqa: BLE001
                pass
        registerFontFamily("LeonaKR", normal="LeonaKR", bold=bold_name, italic="LeonaKR", boldItalic=bold_name)
        _font_name = "LeonaKR"
        return _font_name
    log.warning("내장 가능한 한글 TTF 없음 → CID 폰트 사용 (일부 뷰어에서 한글이 깨질 수 있음). LEONA_PDF_FONT 설정 권장")
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    _font_name = "HYGothic-Medium"
    return _font_name


# ------------------------------------------------------------------ 셀 값 → 표시 문자열

def _fmt(cell) -> str:
    """양식의 숫자 서식("₩"#,##0 / "US$"#,##0.00)을 흉내 내 문자열로."""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        fmt = cell.number_format or ""
        prefix = "₩" if "₩" in fmt else "US$" if "US$" in fmt else "€" if "€" in fmt else ""
        if prefix or "#,##0" in fmt:
            decimals = 2 if "0.00" in fmt else 0
            return f"{prefix}{v:,.{decimals}f}"
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.3f}".rstrip("0").rstrip(".")
    if hasattr(v, "strftime"):
        return v.strftime("%Y. %m. %d")
    return str(v)


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ------------------------------------------------------------------ 본체

def render_pdf(xlsx: bytes, item_rows: tuple[int, int] = (23, 34), total_row: int = 35,
               remark_rows: tuple[int, int] = (38, 49)) -> bytes:
    """XLSX의 인쇄 영역·열 너비·행 높이·셀 스타일을 같은 비율로 PDF에 그린다."""
    from math import floor
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile
    from openpyxl.utils import get_column_letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    font = _font()
    bold = "LeonaKR-Bold" if "LeonaKR-Bold" in pdfmetrics.getRegisteredFontNames() else font
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb.active
    # 업체 양식의 인쇄 영역 A1:F49. 원본 열 너비/행 높이를 읽고 빈 비용 행도 그린다.
    # Excel 열 너비는 기본 글꼴 숫자 폭 기준이다(이 양식: 맑은 고딕 11pt).
    digit_width = round(max(pdfmetrics.stringWidth(str(d), font, wb._fonts[0].sz or 11) for d in range(10)) * 4 / 3)
    widths = [floor(ws.column_dimensions[get_column_letter(c)].width * digit_width + 5) * .75 for c in range(1, 7)]
    heights = [ws.row_dimensions[r].height or ws.sheet_format.defaultRowHeight for r in range(1, 50)]
    xs, ys = [0.0], [0.0]
    for width in widths:
        xs.append(xs[-1] + width)
    for height in heights:
        ys.append(ys[-1] + height)
    page_w, page_h = A4
    margins = ws.page_margins
    left, top = margins.left * 72, page_h - margins.top * 72
    scale = min((page_w - (margins.left + margins.right) * 72) / xs[-1],
                (page_h - (margins.top + margins.bottom) * 72) / ys[-1])
    if ws.print_options.horizontalCentered:
        left = (page_w - xs[-1] * scale) / 2
    buffer = io.BytesIO()
    canvas = Canvas(buffer, pagesize=A4)
    canvas.setTitle(str(ws["B12"].value or "Quotation"))
    canvas.translate(left, top)
    canvas.scale(scale, scale)

    merged = {}
    for area in ws.merged_cells.ranges:
        for row in range(area.min_row, area.max_row + 1):
            for col in range(area.min_col, area.max_col + 1):
                merged[row, col] = area

    theme_ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    theme = ET.fromstring(wb.loaded_theme) if wb.loaded_theme else None
    theme_names = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink")

    def color(value, default=colors.black):
        from openpyxl.styles.colors import COLOR_INDEX
        if value is not None and value.type == "rgb":
            return colors.HexColor("#" + value.rgb[-6:])
        if value is not None and value.type == "indexed" and value.indexed < len(COLOR_INDEX):
            return colors.HexColor("#" + COLOR_INDEX[value.indexed][-6:])
        if value is not None and value.type == "theme" and theme is not None and value.theme < len(theme_names):
            node = theme.find(f".//a:clrScheme/a:{theme_names[value.theme]}", theme_ns)
            if node is not None and len(node):
                rgb = node[0].get("lastClr") or node[0].get("val")
                channels = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
                tint = value.tint
                channels = [v * (1 + tint) if tint < 0 else v * (1 - tint) + tint for v in channels]
                return colors.Color(*channels)
        return default

    text_runs = []
    for row in range(1, 50):
        for col in range(1, 7):
            cell = ws.cell(row, col)
            area = merged.get((row, col))
            if area and (row, col) != (area.min_row, area.min_col):
                continue
            end_col = min(area.max_col if area else col, 6)
            end_row = min(area.max_row if area else row, 49)
            x, y = xs[col - 1], -ys[end_row]
            width, height = xs[end_col] - x, ys[end_row] - ys[row - 1]
            if cell.fill.patternType == "solid":
                canvas.setFillColor(color(cell.fill.fgColor, colors.white))
                canvas.rect(x, y, width, height, fill=1, stroke=0)
            for side, coords in (
                (cell.border.left, (x, y, x, y + height)),
                (cell.border.right, (x + width, y, x + width, y + height)),
                (cell.border.top, (x, y + height, x + width, y + height)),
                (cell.border.bottom, (x, y, x + width, y)),
            ):
                if side and side.style:
                    canvas.setStrokeColor(color(side.color))
                    canvas.setLineWidth(1.5 if side.style in ("medium", "thick") else .6)
                    canvas.line(*coords)
            text = _fmt(cell)
            if not text:
                continue
            # 셀 테두리가 없는 안내/화물 정보는 Excel처럼 오른쪽 빈 칸으로 이어 쓴다.
            bordered = any(s and s.style for s in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom))
            if not bordered and not area:
                if col in (2, 3):
                    end = col
                    while end < 6 and ws.cell(row, end + 1).value is None:
                        end += 1
                    width = xs[end] - x
                elif cell.alignment.horizontal == "right":
                    start = col
                    while start > 2 and ws.cell(row, start - 1).value is None:
                        start -= 1
                    x = xs[start - 1]
                    width = xs[end_col] - x
            if cell.coordinate == "B12":
                width = xs[4] - xs[1]  # 원본 Subject 도형의 실제 폭
            alignment = {"center": 1, "right": 2}.get(cell.alignment.horizontal, 0)
            size = float(cell.font.sz or 10)
            paragraph = None
            while size >= 6:
                style = ParagraphStyle("cell", fontName=bold if cell.font.b else font,
                                       fontSize=size, leading=size * 1.12, alignment=alignment,
                                       textColor=color(cell.font.color), wordWrap="CJK")
                paragraph = Paragraph(_esc(text).replace("\n", "<br/>"), style)
                pw, ph = paragraph.wrap(max(width - 5, 1), height)
                if ph <= height - 2:
                    break
                size -= .5
            else:
                raise ValueError(f"원본 양식 {cell.coordinate} 칸에 내용이 들어가지 않습니다.")
            vertical = cell.alignment.vertical
            py = y + (height - ph) / 2 if vertical == "center" else y + height - ph - 1 if vertical == "top" else y + 1
            text_runs.append((paragraph, x + 2.5, py))

    # 빈 셀의 배경색이 왼쪽에서 이어진 안내 문구를 덮지 않도록 글자는 마지막에 그린다.
    for paragraph, x, y in text_runs:
        paragraph.drawOn(canvas, x, y)

    # 원본 그림 위치를 그대로 사용한다. 로고를 별도의 머리글로 재배치하지 않는다.
    for logo in ws._images:
        anchor = logo.anchor
        origin = anchor._from
        x = xs[origin.col] + origin.colOff / 12700
        y = ys[origin.row] + origin.rowOff / 12700
        if hasattr(anchor, "ext"):
            width, height = anchor.ext.cx / 12700, anchor.ext.cy / 12700
        else:
            width = xs[anchor.to.col] + anchor.to.colOff / 12700 - x
            height = ys[anchor.to.row] + anchor.to.rowOff / 12700 - y
        canvas.drawImage(ImageReader(io.BytesIO(logo._data())), x, -y - height, width, height, mask="auto")

    # openpyxl이 읽지 않는 원본 Subject 사각형도 OOXML의 앵커 위치로 복원한다.
    ns = {"x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(io.BytesIO(xlsx)) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/drawings/drawing") or not name.endswith(".xml"):
                continue
            for anchor in ET.fromstring(archive.read(name)):
                if anchor.find("x:sp", ns) is None:
                    continue
                start, end = anchor.find("x:from", ns), anchor.find("x:to", ns)
                if start is None or end is None:
                    continue
                def point(node):
                    return (xs[int(node.findtext("x:col", namespaces=ns))] + int(node.findtext("x:colOff", namespaces=ns)) / 12700,
                            ys[int(node.findtext("x:row", namespaces=ns))] + int(node.findtext("x:rowOff", namespaces=ns)) / 12700)
                x1, y1 = point(start)
                x2, y2 = point(end)
                canvas.setStrokeColor(colors.black)
                canvas.setLineWidth(.75)
                canvas.rect(x1, -y2, x2 - x1, y2 - y1, fill=0, stroke=1)
    canvas.showPage()
    canvas.save()
    return buffer.getvalue()
