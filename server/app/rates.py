"""요율·구간 룰을 엑셀로 관리 (NFR-06: 요율은 코드 수정 없이 데이터로 변경).

    python -m app.rates template 요율_양식.xlsx     빈 양식 (작성 안내 시트 포함)
    python -m app.rates export   현재_요율.xlsx      DB에 있는 요율 내려받기 (고쳐서 다시 올리면 된다)
    python -m app.rates import   요율.xlsx           검증 후 올리기 — 틀린 줄이 하나라도 있으면 아무것도 바꾸지 않는다

API: GET /api/rates/template · GET /api/rates/export · POST /api/rates/import · GET /api/rates/summary

올리기 규칙: 엑셀의 '출처'가 같은 요율은 통째로 바꾼다 (예: '2026-10 D선사 월간운임'을 다시 올리면 그 출처만 교체).
발주 측 요율표 형식을 받으면 그 형식을 바로 읽는 변환기를 이 모듈에 추가한다.
"""

import argparse
import io
import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from .config import KST
from .extraction import parsers as P
from .models import ContainerType, LaneRule, Rate

RATE_SHEET, LANE_SHEET, GUIDE_SHEET = "요율", "구간룰", "작성 안내"
RATE_COLUMNS = ["구분코드", "견적서 표기", "선적항(POL)", "도착항(POD)", "컨테이너", "인코텀즈", "기준", "단가", "통화",
                "비고", "적용 시작", "적용 종료", "출처"]
LANE_COLUMNS = ["선적항(POL)", "도착항(POD)", "컨테이너", "견적 통화", "운송일수", "DET 프리타임", "DEM 프리타임", "견적 유효일수"]

BASIS_IN = {"R/T": "PER_RT", "RT": "PER_RT", "CBM": "PER_RT", "CNTR": "PER_CNTR", "컨테이너": "PER_CNTR",
            "B/L": "PER_BL", "BL": "PER_BL", "건": "PER_BL", "TRIP": "PER_TRIP", "운행": "PER_TRIP",
            "AT COST": "AT_COST", "ATCOST": "AT_COST", "실비": "AT_COST"}
BASIS_OUT = {"PER_RT": "R/T", "PER_CNTR": "CNTR", "PER_BL": "B/L", "PER_TRIP": "TRIP", "AT_COST": "AT COST"}
KNOWN_CODES = {
    "OCEAN_FREIGHT": "해상운임 (CIF·DDP·EXW 견적에 필수)", "THC": "출발지 터미널 비용 (없으면 부대비용 요율 없음으로 보류)",
    "CFS": "CFS 창고료", "WFG": "부두사용료(Wharfage)", "SHUTTLE": "셔틀 비용", "DOC_FEE": "서류 발급비",
    "SEAL": "씰 비용", "CUSTOMS": "수출 통관비", "INSURANCE": "적하보험료", "PSF": "PSF", "PFS": "PFS", "SHORING": "쇼링 비용",
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")


class RateImportError(ValueError):
    def __init__(self, errors: list[str]):
        super().__init__(f"요율 엑셀에 틀린 곳이 {len(errors)}개 있어 아무것도 바꾸지 않았습니다.")
        self.errors = errors


# ------------------------------------------------------------------ 양식·내보내기


def _header(ws, columns: list[str]) -> None:
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(col) * 2 + 2)
    ws.freeze_panes = "A2"


def _guide(wb: Workbook) -> None:
    ws = wb.create_sheet(GUIDE_SHEET)
    rows = [
        ["요율 엑셀 작성 안내"],
        [],
        ["칸", "쓰는 법"],
        ["구분코드", "아래 코드 중 하나 (새 비용은 영문 대문자로 새 코드를 만들어도 된다)"],
        ["견적서 표기", "견적서에 찍힐 이름 (예: OCEAN FREIGHT, THC, DOCUMENT FEE)"],
        ["선적항·도착항", "UN/LOCODE(KRPUS) 또는 항구 이름(부산, SINGAPORE). 비우면 '모든 항구'"],
        ["컨테이너", "LCL / 20FT GP / 40FT GP / 40HQ / 40RF. 비우면 '모든 컨테이너'"],
        ["인코텀즈", "FOB / CIF / DDP / EXW. 비우면 '모든 조건'. 해상운임은 보통 CIF로 둔다 (FOB는 매수인 부담)"],
        ["기준", "R/T(LCL, 부피·중량 톤) / CNTR(컨테이너당) / B/L(건당) / TRIP(운행당) / AT COST(실비, 단가 비움)"],
        ["단가·통화", "숫자만 (쉼표 가능). 통화는 USD·KRW 등 3자리, 비우면 KRW"],
        ["적용 시작·종료", "YYYY-MM-DD. 비우면 기간 제한 없음. 월간 운임은 그 달 1일~말일로 적는다"],
        ["출처", "예: 2026-10 D선사 월간운임. 같은 출처로 다시 올리면 그 출처 요율만 통째로 바뀐다"],
        [],
        ["구분코드", "설명"],
        *[[k, v] for k, v in KNOWN_CODES.items()],
        [],
        ["'구간룰' 시트", "구간별 운송일수·프리타임·견적 유효일수. 올리면 기존 구간룰은 모두 이 내용으로 바뀐다 (시트가 비어 있으면 그대로 둔다)"],
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = RATE_SHEET
    _header(ws, RATE_COLUMNS)
    ws.append(["OCEAN_FREIGHT", "OCEAN FREIGHT", "KRPUS", "SGSIN", "40HQ", "CIF", "CNTR", 750, "USD", "INCLUSIVE ISPS",
               "2026-10-01", "2026-10-31", "2026-10 D선사 월간운임"])
    ws.append(["THC", "THC", "KRPUS", "", "40HQ", "", "CNTR", 410000, "KRW", "PER CNTR", "", "", "부산 부대비용"])
    lanes = wb.create_sheet(LANE_SHEET)
    _header(lanes, LANE_COLUMNS)
    lanes.append(["KRPUS", "SGSIN", "40HQ", "USD", 5, 7, 7, 14])
    _guide(wb)
    return _save(wb)


def export(session: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = RATE_SHEET
    _header(ws, RATE_COLUMNS)
    for r in session.scalars(select(Rate).order_by(Rate.source, Rate.pol, Rate.pod, Rate.sort_no, Rate.id)):
        ws.append([r.charge_code, r.charge_label, r.pol or "", r.pod or "", r.container_type or "", r.incoterms or "",
                   BASIS_OUT.get(r.basis, r.basis), r.unit_price, r.currency, r.remark or "",
                   r.valid_from.isoformat() if r.valid_from else "", r.valid_until.isoformat() if r.valid_until else "",
                   r.source or ""])
    lanes = wb.create_sheet(LANE_SHEET)
    _header(lanes, LANE_COLUMNS)
    for rule in session.scalars(select(LaneRule).order_by(LaneRule.pol, LaneRule.pod)):
        lanes.append([rule.pol or "", rule.pod or "", rule.container_type or "", rule.quote_currency,
                      rule.transit_days, rule.free_time_det, rule.free_time_dem, rule.validity_days])
    _guide(wb)
    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ 가져오기 (검증 → 전부 성공일 때만 반영)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value, where: str, errors: list[str], required: bool) -> float | None:
    text = _text(value).replace(",", "")
    if not text:
        if required:
            errors.append(f"{where}: 값이 비어 있습니다")
        return None
    try:
        return float(text)
    except ValueError:
        errors.append(f"{where}: 숫자가 아닙니다 ({value})")
        return None


def _date(value, where: str, errors: list[str]) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = P.parse_date(str(value))
    if parsed is None:
        errors.append(f"{where}: 날짜 형식이 아닙니다 (YYYY-MM-DD) — {value}")
        return None
    return date.fromisoformat(parsed.value)


def _port(value, where: str, errors: list[str]) -> str | None:
    text = _text(value)
    if not text:
        return None
    code = P.port_code(text)
    if code is None and P.LOCODE.match(text.upper()):
        code = text.upper()  # 목록에 없는 코드라도 형식이 맞으면 받는다
    if code is None:
        hint = P.port_candidates(text)
        errors.append(f"{where}: 항구를 찾을 수 없습니다 ({text}){' — 혹시 ' + ', '.join(hint) + '?' if hint else ''}")
    return code


def _container(value, where: str, errors: list[str], known: set[str]) -> str | None:
    text = _text(value)
    if not text:
        return None
    if text in known:
        return text
    parsed = P.parse_container(text)
    if parsed and parsed.value in known:
        return parsed.value
    errors.append(f"{where}: 컨테이너 마스터에 없는 값입니다 ({text}) — {', '.join(sorted(known))} 중 하나")
    return None


def _rows(ws, columns: list[str], errors: list[str]):
    header = [_text(c.value) for c in ws[1]]
    missing = [c for c in columns if c not in header]
    if missing:
        errors.append(f"'{ws.title}' 시트 첫 줄에 칸 이름이 없습니다: {', '.join(missing)} (양식을 내려받아 쓰세요)")
        return
    index = {name: header.index(name) for name in columns}
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v in (None, "") for v in row):
            continue
        yield n, {name: (row[i] if i < len(row) else None) for name, i in index.items()}


def parse_workbook(session: Session, data: bytes) -> tuple[list[Rate], list[LaneRule] | None, list[str]]:
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as error:  # 엑셀이 아닌 파일
        raise RateImportError([f"엑셀(.xlsx) 파일을 열 수 없습니다: {error}"]) from error
    if RATE_SHEET not in wb.sheetnames:
        raise RateImportError([f"'{RATE_SHEET}' 시트가 없습니다 (양식을 내려받아 쓰세요)"])
    known = set(session.scalars(select(ContainerType.code)))
    default_source = f"엑셀 업로드 {datetime.now(KST):%Y-%m-%d}"

    rates: list[Rate] = []
    for n, row in _rows(wb[RATE_SHEET], RATE_COLUMNS, errors):
        where = f"{RATE_SHEET} {n}행"
        code = _text(row["구분코드"]).upper().replace(" ", "_")
        if not code:
            errors.append(f"{where}: 구분코드가 비어 있습니다")
        key = _text(row["기준"]).upper()
        basis = BASIS_IN.get(key) or BASIS_IN.get(key.replace(" ", ""))
        if basis is None:
            errors.append(f"{where}: 기준은 R/T·CNTR·B/L·TRIP·AT COST 중 하나여야 합니다 ({row['기준']})")
        terms = _text(row["인코텀즈"]).upper() or None
        if terms and terms not in P.INCOTERMS:
            errors.append(f"{where}: 인코텀즈는 FOB·CIF·DDP·EXW 중 하나이거나 비워야 합니다 ({terms})")
        currency = (_text(row["통화"]) or "KRW").upper()
        if len(currency) != 3 or not currency.isalpha():
            errors.append(f"{where}: 통화는 USD·KRW 같은 3자리여야 합니다 ({currency})")
        price = _number(row["단가"], f"{where} 단가", errors, required=basis != "AT_COST")
        start, end = _date(row["적용 시작"], f"{where} 적용 시작", errors), _date(row["적용 종료"], f"{where} 적용 종료", errors)
        if start and end and start > end:
            errors.append(f"{where}: 적용 시작이 종료보다 늦습니다")
        rates.append(Rate(
            charge_code=code, charge_label=_text(row["견적서 표기"]) or code.replace("_", " "),
            pol=_port(row["선적항(POL)"], f"{where} 선적항", errors), pod=_port(row["도착항(POD)"], f"{where} 도착항", errors),
            container_type=_container(row["컨테이너"], f"{where} 컨테이너", errors, known), incoterms=terms,
            basis=basis or "PER_BL", unit_price=price if basis != "AT_COST" else None, currency=currency,
            remark=_text(row["비고"]) or None, sort_no=10 if code == "OCEAN_FREIGHT" else 50,
            valid_from=start, valid_until=end, source=_text(row["출처"]) or default_source,
        ))

    lanes: list[LaneRule] | None = None
    if LANE_SHEET in wb.sheetnames:
        parsed = list(_rows(wb[LANE_SHEET], LANE_COLUMNS, errors))
        if parsed:
            lanes = []
            for n, row in parsed:
                where = f"{LANE_SHEET} {n}행"
                days = [_number(row[c], f"{where} {c}", errors, required=False) for c in LANE_COLUMNS[4:]]
                lanes.append(LaneRule(
                    pol=_port(row["선적항(POL)"], f"{where} 선적항", errors),
                    pod=_port(row["도착항(POD)"], f"{where} 도착항", errors),
                    container_type=_container(row["컨테이너"], f"{where} 컨테이너", errors, known),
                    quote_currency=(_text(row["견적 통화"]) or "USD").upper(),
                    transit_days=int(days[0]) if days[0] is not None else None,
                    free_time_det=int(days[1]) if days[1] is not None else None,
                    free_time_dem=int(days[2]) if days[2] is not None else None,
                    validity_days=int(days[3]) if days[3] is not None else 14,
                ))
    if not rates and lanes is None and not errors:
        errors.append("올릴 요율이 한 줄도 없습니다")
    return rates, lanes, errors


def import_workbook(session: Session, data: bytes) -> dict:
    """검증을 모두 통과하면 반영하고 요약을 돌려준다. 틀린 곳이 있으면 RateImportError (DB는 그대로)."""
    rates, lanes, errors = parse_workbook(session, data)
    if errors:
        raise RateImportError(errors)
    sources = sorted({r.source for r in rates})
    if sources:
        session.execute(delete(Rate).where(Rate.source.in_(sources)))
    session.add_all(rates)
    if lanes is not None:
        session.execute(delete(LaneRule))
        session.add_all(lanes)
    session.commit()
    warnings = []
    unknown = sorted({r.charge_code for r in rates} - set(KNOWN_CODES))
    if unknown:
        warnings.append(f"처음 보는 구분코드: {', '.join(unknown)} (견적서에 그대로 찍힌다)")
    return {"rates": len(rates), "laneRules": None if lanes is None else len(lanes), "replacedSources": sources,
            "warnings": warnings}


def summary(session: Session) -> list[dict]:
    """출처별 요율 개수·적용 기간 — 오래된 운임이 남아 있는지 확인용."""
    today = datetime.now(KST).date()
    rows = session.execute(
        select(Rate.source, func.count(), func.min(Rate.valid_from), func.max(Rate.valid_until)).group_by(Rate.source)
    ).all()
    return [{"source": s or "(출처 없음)", "count": c, "validFrom": f.isoformat() if f else None,
             "validUntil": u.isoformat() if u else None, "expired": bool(u and u < today)} for s, c, f, u in rows]


# ------------------------------------------------------------------ 명령줄


def main() -> None:
    from .db import SessionLocal, init_db

    parser = argparse.ArgumentParser(description="요율 엑셀 관리")
    parser.add_argument("action", choices=["template", "export", "import"])
    parser.add_argument("file")
    args = parser.parse_args()
    init_db()
    with SessionLocal() as session:
        if args.action == "template":
            open(args.file, "wb").write(template())
            print(f"빈 양식을 만들었습니다: {args.file}")
        elif args.action == "export":
            open(args.file, "wb").write(export(session))
            print(f"현재 요율을 내려받았습니다: {args.file}")
        else:
            try:
                result = import_workbook(session, open(args.file, "rb").read())
            except RateImportError as error:
                print(error)
                for line in error.errors:
                    print("  -", line)
                sys.exit(1)
            print(f"요율 {result['rates']}줄을 올렸습니다. 바뀐 출처: {', '.join(result['replacedSources']) or '-'}")
            for w in result["warnings"]:
                print("  참고:", w)


if __name__ == "__main__":
    main()
