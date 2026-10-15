"""견적서 생성: 요율 마스터로 금액을 계산하고 발주 측 양식(견적서_양식.xlsx)에 채워 저장한다.

- 계산은 룰 베이스만 (LLM 없음). 단가 × 수량 = 금액, 통화별 합계 (양식 'TOTAL (USD+KRW)').
- 요율이 없는 구간은 견적을 만들지 않는다 (FR-501) → None.
- 파생 필드(quoteCurrency·transitTime·freeTimeDet/Dem·validityDays·exchangeRateAsOf)를 lane_rules에서 채운다 (FR-507·508).
- 파일은 storage/<케이스ID>/quotes/<케이스ID>_견적서.xlsx → A트랙 [견적서 송부 초안]이 그대로 첨부한다.
  다시 계산하면 이전 파일은 quote-history/로 옮긴다 (송부 초안에 옛 버전이 같이 붙지 않도록).
- 케이스 ID는 하단 Remark에 표기하고 원본 Subject 줄은 그대로 둔다 (FR-102).
- 원본 배치를 유지하며 고정 문구는 기본 영문으로 출력한다 (FR-510). QUOTE_LANGUAGE=ko이면 원본 한글 문구를 쓴다.

PDF는 LibreOffice로 같은 XLSX를 변환하고, 없거나 실패하면 내장 엔진(quote_pdf.py)으로 만들어 옆에 둔다 → 송부 초안에 PDF·XLSX가 함께 첨부된다 (MUST-SHIP ④).
"""

import logging
import math
from datetime import date
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
from copy import copy
from io import BytesIO
from pathlib import Path
from datetime import timedelta, timezone

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from . import drafts
from .config import KST, settings
from .extraction import ports
from .models import Case, ContainerType, LaneRule, Quote, QuoteItem, QuoteTotal, Rate
from .quote_template import save_filled_template

# 해상운임이 견적에 반드시 들어가야 하는 조건. FOB는 해상운임을 매수인이 부담해 국내 비용만 견적한다.
FREIGHT_REQUIRED = ("CIF", "DDP", "EXW")

ITEM_FIRST_ROW, ITEM_LAST_ROW, TOTAL_ROW = 23, 34, 35  # 양식의 비용 12칸 + 합계 행
REMARK_FIRST_ROW, REMARK_LAST_ROW = 39, 47
MONEY_FORMAT = {"KRW": '"₩"#,##0', "USD": '"US$"#,##0.00'}


def _applies(column, value):
    """요율 행의 조건 칸이 비어 있으면 모든 값에 적용."""
    return or_(column.is_(None), column == value)


def find_rates(session: Session, qi, on_date) -> list[Rate]:
    rows, _reason = rates_with_reason(session, qi, on_date)
    return rows


def rate_date(qi, today):
    """요율을 고르는 기준일 = 선적일(화물 준비일). 없거나 지난 날이면 오늘 (발주 측 요율표 조회 규칙)."""
    try:
        ready = date.fromisoformat(qi.cargoReadyDate) if qi.cargoReadyDate else None
    except ValueError:
        ready = None
    return ready if ready and ready >= today else today


def _carrier_key(value: str | None) -> str:
    return (value or "").strip().upper()


class RateSelectionError(ValueError):
    """자동으로 비교할 수 없는 요율은 담당자 확인으로 넘긴다."""


def choose_carrier(rows: list[Rate], wanted: str | None) -> tuple[list[Rate], str | None]:
    """해상운임 행이 선사별로 여럿이면 하나만 남긴다: 지정 선사 → 없으면 최저가 (발주 측 규칙).

    선사 없는(ANY) 행은 선사 무관 운임으로 함께 경쟁한다. 부대비용은 선사와 무관하게 그대로 둔다.
    """
    freight = [r for r in rows if r.charge_code == "OCEAN_FREIGHT"]
    if len(freight) <= 1:
        return rows, (freight[0].carrier if freight and _carrier_key(freight[0].carrier) not in ("", "ANY") else None)
    pool = freight
    if wanted:
        named = [r for r in freight if _carrier_key(r.carrier) == _carrier_key(wanted)]
        pool = named or freight  # 지정 선사의 운임이 없으면 최저가로 (견적서 비고에 표시)
    currencies = {(r.currency or "").strip().upper() for r in pool}
    if len(currencies) > 1:
        raise RateSelectionError(
            f"해상운임 후보의 통화가 서로 다릅니다 ({', '.join(sorted(currencies))}) — "
            "환산 기준 없이 최저가를 자동 선택할 수 없습니다. 담당자가 선사·요율을 확인해 주세요 (FR-501)"
        )
    best = min(pool, key=lambda r: (r.unit_price if r.unit_price is not None else float("inf"), r.id))
    keep = [r for r in rows if r.charge_code != "OCEAN_FREIGHT" or r is best]
    carrier = best.carrier if _carrier_key(best.carrier) not in ("", "ANY") else None
    return keep, carrier


# ---------------------------------------------------------------- 인코텀즈별 비용 분기 (FR-503, 발주 측 회신 2026-10-06)
# 수출자 기준으로 견적에 넣는 비용 묶음. 수입자는 그 반대(수출자가 내지 않는 쪽)를 넣는다.
#   EXW: 전부 제외 / FOB: 출발지 비용만 / CIF: 출발지 + 해상운임 + 보험 (도착지 제외) / DDP: 전부 (관세·부가세는 '별도' 표기)
# 단, 수입 CIF는 발주 측 요율표 '조회예시' 시트(상하이→부산 CIF, 출발지·해상운임·도착지·보험 전부 합산)를 따른다.
ORIGIN, FREIGHT, INSURANCE, DEST = "ORIGIN", "FREIGHT", "INSURANCE", "DEST"
ALL_PARTS = frozenset({ORIGIN, FREIGHT, INSURANCE, DEST})
EXPORTER_PAYS = {
    "EXW": frozenset(),
    "FOB": frozenset({ORIGIN}),
    "CIF": frozenset({ORIGIN, FREIGHT, INSURANCE}),
    "DDP": ALL_PARTS,
}
IMPORTER_PAYS = {**{terms: ALL_PARTS - pays for terms, pays in EXPORTER_PAYS.items()},
                 "CIF": ALL_PARTS}  # 조회예시 기준 (수출자의 반대라면 도착지 비용만이지만 예시는 전부 넣는다)
EXPORT, IMPORT = "EXPORT", "IMPORT"
INSURANCE_CODES = ("INS", "INSURANCE")
HOME_COUNTRY = "KR"


def trade_role(qi) -> str | None:
    """화주가 수출자인지 수입자인지 — 6장에 칸이 없어 항구 국가로 본다: 선적항이 국내면 수출, 도착항이 국내면 수입."""
    if (qi.pol or "").upper().startswith(HOME_COUNTRY):
        return EXPORT
    if (qi.pod or "").upper().startswith(HOME_COUNTRY):
        return IMPORT
    return None


def included_parts(incoterms: str | None, role: str | None) -> frozenset | None:
    """이 조건·역할의 견적에 넣는 비용 묶음. 모르는 조건이면 None."""
    if role is None:
        return None
    return (EXPORTER_PAYS if role == EXPORT else IMPORTER_PAYS).get((incoterms or "").upper())


def charge_part(rate: Rate) -> str | None:
    """요율 행이 어느 비용 묶음인지. None = 구분 없는 공통 비용(서류비·B/L 등) — 견적에 다른 비용이 있으면 함께 넣는다."""
    code = (rate.charge_code or "").upper()
    if code == "OCEAN_FREIGHT":
        return FREIGHT
    if code in INSURANCE_CODES:
        return INSURANCE
    direction = (getattr(rate, "direction", None) or "").upper()
    if direction in (ORIGIN, DEST):
        return direction
    # 방향 칸이 없는 요율(우리 양식): 선적항만 적힌 행은 출발지 비용, 도착항만 적힌 행은 도착지 비용
    if getattr(rate, "pol", None) and not getattr(rate, "pod", None):
        return ORIGIN
    if getattr(rate, "pod", None) and not getattr(rate, "pol", None):
        return DEST
    return None


def same_table_rows(rows: list[Rate]) -> list[Rate]:
    """요율표가 여러 벌 올라와 있을 때 비용이 겹치지 않게 한다.

    항구를 정하지 않은 공통 비용(ALL — 서류비·B/L·보험 등)은 '그 요율표가 이 구간의 요율을 갖고 있을 때'만 쓴다.
    예) 인천→시드니는 A 요율표에만 있는데 B 요율표의 공통 서류비·보험료까지 붙어 서류비·보험료가 두 번 청구되는 것을 막는다.
    """
    lane_sources = {r.source for r in rows if r.pol or r.pod}
    return [r for r in rows if r.pol or r.pod or r.source in lane_sources]


def one_table_per_charge(rows: list[Rate]) -> list[Rate]:
    """같은 비용 항목(THC·서류비 등)이 두 요율표에 모두 있으면 한 요율표의 것만 쓴다 — 두 번 청구하지 않는다.

    우선순위: 실제 요율표(이름이 TEST로 시작하지 않는 것) → 나중에 올린 요율표. 해상운임은 choose_carrier가 하나만 고른다.
    서로 다른 비용을 나눠 가진 요율표(예: 해상운임 표 + 국내 부대비용 표)는 그대로 함께 쓴다.
    """
    newest: dict = {}
    for r in rows:
        newest[r.source] = max(newest.get(r.source, 0), r.id or 0)

    def rank(source):
        return ((source or "").startswith("TEST"), -newest[source])

    def charge_key(row):
        # 같은 THC라도 출발지와 도착지 비용은 서로 다른 청구다.
        return row.charge_code, charge_part(row)

    sources: dict = {}
    for r in rows:
        sources.setdefault(charge_key(r), set()).add(r.source)
    keep = {code: min(found, key=rank) for code, found in sources.items()}
    return [r for r in rows if r.charge_code == "OCEAN_FREIGHT" or r.source == keep[charge_key(r)]]


def rates_with_reason(session: Session, qi, on_date, carrier: str | None = None) -> tuple[list[Rate], str | None]:
    """견적에 쓸 요율 행과, 못 쓰면 그 이유 (FR-501: 요율이 모자라면 견적을 만들지 않는다)."""
    lane = f"{qi.pol}→{qi.pod} {qi.containerType}"
    rows = session.scalars(
        select(Rate)
        .where(
            _applies(Rate.pol, qi.pol),
            _applies(Rate.pod, qi.pod),
            _applies(Rate.container_type, qi.containerType),
            _applies(Rate.incoterms, qi.incoterms),
            or_(Rate.valid_from.is_(None), Rate.valid_from <= on_date),
            or_(Rate.valid_until.is_(None), Rate.valid_until >= on_date),
        )
        .order_by(Rate.sort_no, Rate.id)
    ).all()
    rows = list(rows)
    # 구간(pol 또는 pod)이 지정된 요율이 하나도 없으면 '요율 없는 구간'으로 본다 — 공통 부대비용만으로 견적 금지
    # (조건별로 비용을 빼기 전에 본다: 수입 CIF처럼 공통 도착지 비용만 남는 견적도 구간 요율이 있어야 한다)
    if not any(r.pol or r.pod for r in rows):
        # 날짜만 안 맞는 경우(화물 준비일이 요율 유효기간 밖)는 '미등록'이 아니다 → 사유를 구분해 알린다
        other_dates = [r for r in session.scalars(select(Rate).where(
            _applies(Rate.pol, qi.pol), _applies(Rate.pod, qi.pod),
            _applies(Rate.container_type, qi.containerType), _applies(Rate.incoterms, qi.incoterms),
        )).all() if r.pol or r.pod]
        if other_dates:
            starts = [r.valid_from for r in other_dates if r.valid_from]
            ends = [r.valid_until for r in other_dates if r.valid_until]
            period = f"{min(starts) if starts else ''} ~ {max(ends) if ends else ''}"
            return [], (f"{on_date} 기준으로 유효한 요율이 없습니다 {lane} "
                        f"(등록된 요율 기간 {period}) — 요율표 갱신 필요 (FR-501)")
        return [], f"요율 미등록 구간 {lane} (FR-501)"
    rows = same_table_rows(rows)
    role = trade_role(qi)
    if role is None:
        return [], f"수출·수입을 구분할 수 없는 구간 {qi.pol}→{qi.pod} (선적항·도착항 모두 국내 아님) — 담당자 처리"
    parts = included_parts(qi.incoterms, role)
    if parts is not None:
        if not parts:
            who = "수출자" if role == EXPORT else "수입자"
            return [], f"{qi.incoterms} 조건의 {who}는 견적에 넣을 비용이 없습니다 (상대방 부담) — 담당자 확인 (FR-503)"
        rows = [r for r in rows if charge_part(r) is None or charge_part(r) in parts]
    rows = one_table_per_charge(rows)
    try:
        rows, _ = choose_carrier(rows, carrier)
    except RateSelectionError as error:
        return [], str(error)
    # 해상운임을 넣어야 하는 조건인데 이 구간의 해상운임 요율이 없으면 견적 불가 (부대비용만으로 만들지 않는다)
    needs_freight = FREIGHT in parts if parts is not None else qi.incoterms in FREIGHT_REQUIRED
    if needs_freight and not any(r.charge_code == "OCEAN_FREIGHT" for r in rows):
        return [], f"해상운임 요율 없음 {lane} {qi.incoterms} (FR-501)"
    # 터미널 비용(THC·OTHC·DTHC)이 없으면 부대비용 요율이 없는 것 — 해상운임만 있는 불완전한 견적을 만들지 않는다
    if not any("THC" in (r.charge_code or "") for r in rows):
        return [], f"국내 부대비용(THC 등) 요율 없음 {qi.pol} {qi.containerType} (FR-501)"
    # CIF는 보험료가 조건에 포함된다 — 보험 요율이 없으면 보험료가 조용히 빠진 견적을 만들지 않는다
    if ((qi.incoterms or "").upper() == "CIF" and (parts is None or INSURANCE in parts)
            and not any(charge_part(r) == INSURANCE for r in rows)):
        return [], f"보험료 요율 없음 {lane} CIF — 보험료가 빠진 견적은 만들지 않습니다 (FR-501)"
    return list(rows), None


def lane_rule(session: Session, qi) -> LaneRule | None:
    rows = session.scalars(
        select(LaneRule).where(
            _applies(LaneRule.pol, qi.pol), _applies(LaneRule.pod, qi.pod),
            _applies(LaneRule.container_type, qi.containerType),
        )
    ).all()
    # 조건이 많이 맞는(구체적인) 룰 우선
    return max(rows, key=lambda r: sum(x is not None for x in (r.pol, r.pod, r.container_type)), default=None)


def revenue_ton(qi) -> float:
    """LCL 청구 단위 R/T = max(CBM, 중량톤), 최소 1 R/T (예시 견적서 '1CBM 기준')."""
    cbm = qi.totalCbm
    if cbm is None and None not in (qi.boxL, qi.boxW, qi.boxH, qi.qty):
        cbm = qi.boxL * qi.boxW * qi.boxH / 1e9 * qi.qty
    ton = (qi.grossWeightKg or 0) / 1000
    return round(max(1.0, cbm or 0, ton), 3)


USABLE_RATIO = 0.88  # 컨테이너 내부 부피 중 실제 적재 가능한 비율 (산정용 보수 기준)


def container_count(session: Session, qi, stated: int | None) -> tuple[int | None, str | None]:
    """FCL 컨테이너 대수 → (대수, 산정 근거). 메일에 적힌 대수가 우선, 없으면 부피·중량으로 필요한 대수를 계산.

    6장에 대수 필드가 없어 추출 결과의 보조값(containerCount)으로 받는다. LCL이면 (None, None).
    """
    if qi.containerType in (None, "LCL"):
        return None, None
    if stated:
        return stated, None
    ct = session.get(ContainerType, qi.containerType)
    cbm = qi.totalCbm
    if cbm is None and None not in (qi.boxL, qi.boxW, qi.boxH, qi.qty):
        cbm = qi.boxL * qi.boxW * qi.boxH / 1e9 * qi.qty
    needs = [1]
    if ct and ct.max_cbm and cbm:
        needs.append(math.ceil(cbm / (ct.max_cbm * USABLE_RATIO)))
    if ct and ct.max_gross_kg and qi.grossWeightKg:
        needs.append(math.ceil(qi.grossWeightKg / ct.max_gross_kg))
    count = max(needs)
    basis = []
    if cbm:
        basis.append(f"{_fmt_num(round(cbm, 2))}CBM")
    if qi.grossWeightKg:
        basis.append(f"{_fmt_num(qi.grossWeightKg)}KG")
    note = f"컨테이너 대수: 화물 {'·'.join(basis) or '정보'} 기준 {qi.containerType} {count}대로 산정" if basis else None
    return count, note


def quantity(basis: str, qi, count: int | None = None) -> float | None:
    if basis == "PER_RT":
        return revenue_ton(qi)
    if basis in ("AT_COST", "PERCENT"):
        return None
    if basis == "PER_CNTR":
        return float(count or 1)
    return 1.0  # PER_BL·PER_SHIPMENT·PER_TRIP


def money(value: float, currency: str) -> float:
    return float(round(value)) if currency == "KRW" else round(value + 1e-9, 2)


class QuoteGenerationError(ValueError):
    """필수 산출물 두 개를 완성하지 못했을 때 C 검증 결과로 안내한다."""


class InsuranceCurrencyError(ValueError):
    """환산 기준 없이 서로 다른 통화로 비율 보험료를 산정하지 않는다."""


def require_insurance_currency(rate: Rate, qi) -> None:
    if (rate.charge_code or "").upper() not in INSURANCE_CODES:
        return
    if qi.ccy and qi.ccy != rate.currency:
        raise InsuranceCurrencyError(
            f"보험료 산정 통화 불일치 — 인보이스 {qi.ccy} / 보험요율 {rate.currency}. "
            "환산 기준이 없어 자동 견적을 보류합니다. 담당자가 통화와 요율을 확인해 주세요."
        )


class QuoteInformationError(ValueError):
    """FR-507 필수 정보가 부족하거나 유효하지 않아 견적을 보류한다."""

    def __init__(self, issues: dict[str, str]):
        self.fields = list(issues)
        super().__init__(
            "FR-507 필수 견적 정보 확인 필요 — " + "; ".join(issues.values())
            + ". 요율·룰 정보를 보완한 뒤 다시 실행해 주세요."
        )


def export_fob_schedule(session: Session, qi, on_date: date,
                        wanted_carrier: str | None) -> tuple[Rate | None, str | None]:
    """수출 FOB의 일정만 조회한다. 선택한 해상운임은 청구 목록에 추가하지 않는다.

    incoterms는 비용 적용 조건이므로 일정 조회에는 쓰지 않는다. 구간·컨테이너·
    유효기간과 기존 선사 선택 규칙은 그대로 적용한다 (FR-503·507).
    """
    rows = list(session.scalars(select(Rate).where(
        Rate.charge_code == "OCEAN_FREIGHT",
        _applies(Rate.pol, qi.pol), _applies(Rate.pod, qi.pod),
        _applies(Rate.container_type, qi.containerType),
        or_(Rate.valid_from.is_(None), Rate.valid_from <= on_date),
        or_(Rate.valid_until.is_(None), Rate.valid_until >= on_date),
    ).order_by(Rate.sort_no, Rate.id)).all())
    try:
        selected, carrier = choose_carrier(rows, wanted_carrier)
    except RateSelectionError as error:
        raise QuoteInformationError({"carrier": str(error)}) from error
    return (selected[0] if selected else None), carrier


def required_schedule(qi, freight: Rate | None, on_date: date, today: date,
                      rates: list[Rate]) -> tuple[str, str, str]:
    """파일·견적 버전을 만들기 전에 일정, Free Time, 유효기간을 확인한다."""
    issues: dict[str, str] = {}
    for field, label, minimum in (
        ("transitTime", "Transit Time(운송일수)", 0),
        ("freeTimeDet", "Free Time DET", 0),
        ("freeTimeDem", "Free Time DEM", 0),
        ("validityDays", "견적 유효일수", 1),
    ):
        value = getattr(qi, field)
        if value is None:
            issues[field] = f"{label} 누락"
        elif type(value) is not int or value < minimum:
            issues[field] = f"{label}은 {minimum} 이상의 정수여야 합니다"

    def read_date(field: str, value: str | None) -> date | None:
        if not value:
            issues[field] = f"{field.upper()} 누락"
            return None
        try:
            return date.fromisoformat(value)
        except (TypeError, ValueError):
            issues[field] = f"{field.upper()} 날짜 형식 오류"
            return None

    etd = read_date("etd", freight.etd if freight else None)
    raw_eta = freight.eta if freight else None
    if not raw_eta and etd is not None and "transitTime" not in issues:
        try:
            eta = etd + timedelta(days=qi.transitTime)  # ETA = ETD + 운송일수 (0일도 유효)
        except OverflowError:
            eta = None
            issues["eta"] = "ETA 계산 가능한 날짜 범위 초과"
    else:
        eta = read_date("eta", raw_eta)
    if etd is not None and etd < on_date:
        issues["etd"] = f"ETD {etd}가 선적 기준일 {on_date}보다 이전입니다"
    if etd is not None and eta is not None and eta < etd:
        issues["eta"] = f"ETA {eta}가 ETD {etd}보다 이전입니다"

    valid_until = None
    if "validityDays" not in issues:
        try:
            valid_until = valid_until_date(today, qi.validityDays, rates)
        except OverflowError:
            issues["validUntil"] = "견적 유효기간 계산 가능한 날짜 범위 초과"
        if valid_until is not None and valid_until < today:
            issues["validUntil"] = f"견적 유효기간 {valid_until}이 발행일 {today}보다 이전입니다"
    if issues:
        raise QuoteInformationError(issues)
    return etd.isoformat(), eta.isoformat(), valid_until.isoformat()


def build(session: Session, case: Case, extraction: dict | None = None) -> dict | None:
    started = time.perf_counter()  # 견적서 생성 시간 측정 (NFR-02: 10초 이내)
    qi = case.quote_input
    stated = ((extraction or {}).get("fields") or {}).get("containerCount") or {}
    count, count_note = container_count(session, qi, stated.get("value") if stated.get("status") == "filled" else None)
    from .validation import today_kst  # 순환 import 방지

    today = today_kst()
    wanted = ((extraction or {}).get("fields") or {}).get("carrier") or {}
    wanted_carrier = wanted.get("value") if wanted.get("status") == "filled" else None
    on_date = rate_date(qi, today)
    rates, _reason = rates_with_reason(session, qi, on_date, wanted_carrier)
    if not rates:
        return None
    # 적용 대상 보험료만 확인한다. 보류 시 견적 버전·파일을 만들지 않는다.
    for rate in rates:
        if rate.basis == "PERCENT":
            require_insurance_currency(rate, qi)
    _, carrier = choose_carrier(rates, wanted_carrier)
    freight = next((r for r in rates if r.charge_code == "OCEAN_FREIGHT"), None)
    if (qi.incoterms or "").upper() == "FOB" and trade_role(qi) == EXPORT:
        freight, carrier = export_fob_schedule(session, qi, on_date, wanted_carrier)

    rule = lane_rule(session, qi)
    qi.quoteCurrency = rule.quote_currency if rule else "USD"
    # 운송일수·프리타임: 고른 해상운임 행(선사별)이 우선, 없으면 구간 룰 (FR-507)
    qi.transitTime = (freight.transit_days if freight and freight.transit_days is not None
                      else (rule.transit_days if rule else None))
    qi.freeTimeDet = (freight.free_time_det if freight and freight.free_time_det is not None
                      else (rule.free_time_det if rule else None))
    qi.freeTimeDem = (freight.free_time_dem if freight and freight.free_time_dem is not None
                      else (rule.free_time_dem if rule else None))
    qi.validityDays = rule.validity_days if rule else default_validity_days()
    qi.exchangeRateAsOf = today.isoformat()  # 실제 환율은 출항일 기준 적용 — 견적서 비고에 고지 (FR-508)
    etd, eta, valid_until = required_schedule(qi, freight, on_date, today, rates)

    version = (session.scalar(select(func.max(Quote.version_no)).where(Quote.case_pk == case.id)) or 0) + 1
    quote = Quote(case=case, version_no=version, valid_until=valid_until,
                  carrier=carrier, etd=etd, eta=eta)
    totals: dict[str, float] = {}
    for rate in rates:
        qty = quantity(rate.basis, qi, count)
        amount, unit_price, remark = None, rate.unit_price, rate.remark
        if rate.basis == "PERCENT":
            amount, remark = percent_amount(rate, qi)
            unit_price = None
        elif qty is not None and rate.unit_price is not None:
            amount = money(rate.unit_price * qty, rate.currency)
        if amount is not None:
            totals[rate.currency] = money(totals.get(rate.currency, 0) + amount, rate.currency)
        quote.items.append(QuoteItem(
            charge_code=rate.charge_code, charge_label=rate.charge_label, basis=rate.basis,
            rate=unit_price, qty=qty, amount=amount, currency=rate.currency, remark=remark, source_ref=rate.source_ref,
        ))
    for currency in sorted(totals, key=lambda c: (c != "USD", c)):
        quote.totals.append(QuoteTotal(currency=currency, total=totals[currency]))
    session.add(quote)
    session.flush()

    try:
        xlsx = render_xlsx(case, quote, count, count_note)
        if not settings.quote_pdf:
            raise QuoteGenerationError("PDF 생성 설정이 꺼져 있습니다. PDF·XLSX를 함께 생성해야 합니다.")
        pdf, pdf_note = to_pdf(xlsx)
        if not pdf or not pdf.startswith(b"%PDF"):
            raise QuoteGenerationError(pdf_note or "PDF 생성 결과가 없습니다.")
        # 두 파일을 먼저 완성한 후 저장한다. 한쪽 저장에 실패해도 부분 파일은 첨부되지 않게 치운다.
        retire_files(case, version - 1)
        try:
            quote.xlsx_path = _save(case, quote, "xlsx", xlsx)
            pdf_path = _save(case, quote, "pdf", pdf)
        except Exception:
            retire_files(case, version)
            raise
    except Exception as error:
        for item in [*quote.items, *quote.totals]:
            session.delete(item)
        session.delete(quote)  # 실패한 산출물을 정상 견적 버전으로 남기지 않는다.
        session.flush()
        retire_files(case, version - 1)
        raise QuoteGenerationError(f"견적서 생성 실패 — PDF·XLSX 송부를 보류합니다: {error}") from error
    session.flush()
    elapsed_ms = round((time.perf_counter() - started) * 1000)
    return {
        "elapsedMs": elapsed_ms,  # 요율 조회·금액 계산·XLSX·PDF 생성까지 걸린 시간
        "withinLimit": elapsed_ms <= QUOTE_TIME_LIMIT_MS,
        "version": quote.version_no,
        "file": quote.xlsx_path,
        "pdf": pdf_path,
        "pdfNote": pdf_note,
        "validUntil": quote.valid_until,
        "totals": {t.currency: t.total for t in quote.totals},
        "items": len(quote.items),
        "containerCount": count,
        "containerCountNote": count_note,
        "rateSources": sorted({r.source for r in rates if r.source}),
        "rateDate": on_date.isoformat(),
        "tradeRole": trade_role(qi),
        "carrier": carrier,
        "carrierRequested": wanted_carrier,
        "etd": etd, "eta": eta,
        "scheduleUnconfirmed": False,  # 기존 반환 형식 유지. 일정 미확보 시에는 견적 자체를 보류한다.
        "lines": [{"charge": i.charge_label, "amount": i.amount, "currency": i.currency, "source": i.source_ref}
                  for i in quote.items],  # 금액별 근거 (어느 요율표 행) FR-509
    }


VALIDITY_MIN, VALIDITY_MAX = 1, 14
QUOTE_TIME_LIMIT_MS = 10_000  # NFR-02: 견적서 생성 10초 이내


def valid_until_date(today: date, days: int, rates: list[Rate]) -> date:
    """견적 유효기간 = 발행일 + 유효일수(달력일). 견적에 쓴 요율의 종료일이 더 빠르면 그날까지 (발주 측 회신 2026-10-06)."""
    until = today + timedelta(days=days)
    ends = [r.valid_until for r in rates if r.valid_until is not None]
    return min([until, *ends])


def default_validity_days() -> int:
    """견적 유효기간 기본값(일). .env의 QUOTE_VALIDITY_DAYS로 1~14 중에서 고른다 (범위를 벗어나면 1 또는 14)."""
    try:
        days = int(getattr(settings, "quote_validity_days", VALIDITY_MAX))
    except (TypeError, ValueError):
        days = VALIDITY_MAX
    return min(VALIDITY_MAX, max(VALIDITY_MIN, days))


INSURED_VALUE_RATIO = 1.10  # 보험료 = 계약금액 × 110% × 보험요율 (발주 측 회신 2026-10-06)


def percent_amount(rate: Rate, qi) -> tuple[float | None, str | None]:
    """화물가액 × 비율, 최저 금액. 보험료는 계약금액의 110%에 보험요율을 곱한다 (예: 15,000 × 110% × 0.2% = 33, 최저 30).

    통화가 다른 비율 보험료는 담당자 확인 전까지 견적을 보류한다(환율 미반영).
    """
    require_insurance_currency(rate, qi)
    pct = rate.rate_pct or 0
    insured = (rate.charge_code or "").upper() in INSURANCE_CODES
    ratio = INSURED_VALUE_RATIO if insured else 1.0
    label = ("INV.V x 110% x " if insured else "INV.V x ") + f"{pct * 100:g}%" + (
        f" (MIN {rate.currency} {rate.min_amount:,.0f})" if rate.min_amount else "")
    if qi.invoiceValue is None or (qi.ccy and qi.ccy != rate.currency):
        return None, f"{label} — 인보이스 통화가 달라 실비 청구"
    amount = max(qi.invoiceValue * ratio * pct, rate.min_amount or 0)
    return money(amount, rate.currency), label


# ------------------------------------------------------------------ XLSX 양식 채우기


def port_label(code: str | None) -> str:
    names = ports().get(code or "", ())
    english = [n for n in names if n.isascii() and len(n) > 2]
    return english[0] if english else (code or "")


def _fmt_num(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:,.0f}" if float(value).is_integer() else f"{value:,.2f}"


def total_text(quote: Quote) -> str:
    parts = [f"{t.currency}{t.total:,.2f}" if t.currency == "USD" else f"{t.currency}{t.total:,.0f}" for t in quote.totals]
    return " + ".join(parts)


# 원본 파일은 보존하고 생성본의 고정 문구만 선택한 언어로 채운다 (기본 영문, FR-510).
# 화주가 적은 회사명·담당자명은 임의로 번역하지 않는다.
TEXT = {
    "en": {
        "to": "To : {customer} / Attn : {contact}", "to_one": "To : {name}",
        "from": "From : LEONA SEA & AIR CO., LTD. / Quotation Team",
        "subject": "Subject : Ocean Freight Quotation",
        "greeting1": "* We wish your company continued success.",
        "greeting2": "* Please find our quotation below for your reference.",
        "thanks": "* Thank you.", "seal": "(Company seal omitted)",
        "rt_basis": " (based on {rt} R/T)",
        # FR-508: 견적 통화·환율 기준일 명시 + '유효기간 내 환율 변동에 따라 변동될 수 있음' 고지
        "fx": "* Quote currency: {currency} / Exchange rate as of {as_of} (rate on the actual sailing date applies).",
        "subject_to_change": "* Subject to change with exchange rate fluctuation within validity and with cargo details.",
        "carrier": "* CARRIER : {carrier}",
        "schedule_tbc": "* ETD / ETA : to be confirmed{carrier} (no sailing on or after the cargo ready date in the current schedule)",
        "transit": "* T/T : approx. {days} days",
        "free_time": "* FREE TIME : DET {det} days / DEM {dem} days",
        "validity": "* VALIDITY : {date}",
        "duty_separate": "* Customs duty and VAT are not included and will be charged separately.",
        "quote_no": "* QUOTE NO. : {case_id} (Rev. {version})",
    },
    "ko": {
        "to": "수 신 : {customer} / {contact}", "to_one": "수 신 : {name}",
        "from": "발 신 : 레오나 해운항공㈜ / 견적 담당 드림",
        "rt_basis": " ({rt} R/T 기준)",
        "fx": "* 견적 통화 {currency} · 환율 기준일 {as_of} (실제 출항일 환율 적용)",
        "subject_to_change": "* 유효기간 내 환율 변동 및 화물 DETAIL의 변경에 따라 상기 견적이 달라질 수 있습니다.",
        "carrier": "* 선사 : {carrier}",
        "schedule_tbc": "* ETD / ETA : 확인 후 안내{carrier} (화물 준비일 이후 출항 스케줄 미등록)",
        "transit": "* T/T : 약 {days}일",
        "free_time": "* FREE TIME : DET {det}일 / DEM {dem}일",
        "validity": "* VALIDITY : {date}",
        "duty_separate": "* 관세·부가세 별도",
        "quote_no": "* 견적번호 : {case_id} (v{version})",
    },
}
BASIS_LABEL = {"PER_RT": "PER R/T", "PER_CNTR": "PER CNTR", "PER_BL": "PER B/L", "PER_SHIPMENT": "PER SHIPMENT",
               "PER_TRIP": "PER TRIP"}
# 요율 비고에 섞인 한글 → 영문 (영문 견적서). 여기 없는 한글 비고는 부과 기준(PER CNTR 등)으로 대신한다
REMARK_EN = (("(MIN 기준)", "(MIN)"), ("보험요율", "insurance rate"),
             ("— 인보이스 통화가 달라 실비 청구", "— charged at cost (invoice currency differs)"))
_HANGUL = re.compile("[가-힣]")
_COUNT_NOTE = re.compile(r"컨테이너 대수: 화물 (.+) 기준 (.+) (\d+)대로 산정")


def language() -> str:
    """FR-510 기본 영문. 명시적으로 ko를 선택했을 때만 원본 한글 문구를 쓴다."""
    lang = (getattr(settings, "quote_language", None) or "en").strip().lower()
    return lang if lang in TEXT else "en"


# 발주 측 요율표(OceanFreight·Surcharges 시트)에서 온 행의 근거 표기 (rates.py가 source_ref에 넣는다)
CLIENT_TABLE_REFS = ("OceanFreight#", "Surcharges#")


def item_remark(item: QuoteItem, lang: str) -> str:
    """비용 줄의 REMARK 칸. 영문 견적서에서는 한글 비고를 영문으로 바꾸고, 바꿀 수 없으면 부과 기준을 적는다.

    발주 측 요율표의 remark 칸은 요율 관리용 내부 메모다('GRI 인상 → 새 행으로 처리', '최저가 선택 규칙 예시' 등).
    고객에게 나가는 견적서에는 언어와 관계없이 찍지 않고 부과 기준(PER CNTR 등)만 적는다.
    화물가액 % 줄(보험)의 비고는 요율표 글자가 아니라 시스템이 만든 산정식이라 그대로 둔다.
    """
    fallback = BASIS_LABEL.get(item.basis, "")
    from_client_table = (getattr(item, "source_ref", None) or "").startswith(CLIENT_TABLE_REFS)
    if from_client_table and item.basis != "PERCENT":
        return fallback
    text = item.remark or fallback
    if lang != "en" or not _HANGUL.search(text):
        return text
    for korean, english in REMARK_EN:
        text = text.replace(korean, english)
    return fallback if _HANGUL.search(text) else text


def kst(moment):
    """DB의 시각(UTC, 시간대 표시가 없을 수 있음) → 한국 시각. 견적서 DATE가 자정~오전 9시에 하루 전으로 찍히지 않게 한다."""
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(KST)


def recipient_line(case: Case, lang: str) -> str:
    """견적서 수신 줄 (FR-208 To/Attn). 회사명·담당자명을 못 뽑았으면 회신 주소(없으면 요청 메일의 보낸 사람)를 적는다."""
    qi, t = case.quote_input, TEXT[lang]
    customer, contact = (qi.customerName or "").strip(), (qi.contactName or "").strip()
    if customer and contact:
        return t["to"].format(customer=customer, contact=contact)
    if customer or contact:
        return t["to_one"].format(name=customer or contact)
    fallback = (qi.contactEmail or "").strip()
    if not fallback:
        try:
            original = drafts._original(case)
            fallback = (original.from_name or original.from_email or "").strip()
        except Exception:  # noqa: BLE001 — 수신 줄 때문에 견적서가 실패하면 안 된다
            fallback = ""
    return t["to_one"].format(name=fallback)


def incoterm_remarks(incoterms: str | None, lang: str) -> list[str]:
    """조건별 고지 문구. DDP의 관세·부가세 별도 안내를 유지한다."""
    t, terms = TEXT[lang], (incoterms or "").upper()
    if terms == "DDP":
        return [t["duty_separate"]]
    return []


def count_note_text(note: str, lang: str) -> str:
    """컨테이너 대수 산정 근거 (container_count의 한글 문장 → 영문 견적서용)."""
    m = _COUNT_NOTE.fullmatch(note) if lang == "en" else None
    if not m:
        return note
    return f"Container q'ty : {m[3]} x {m[2]} (estimated from cargo {m[1].replace('·', ' / ')})"


def _korean_date(iso: str | None) -> str:
    """'2026-09-30' → '2026년 9월 30일' (예시 견적서 VALIDITY 표기)"""
    if not iso:
        return ""
    y, m, d = iso.split("-")
    return f"{int(y)}년 {int(m)}월 {int(d)}일"


# 비용 적용 여부·금액은 rates_with_reason/build가 정한다. 아래는 출력 위치만 정한다.
COST_ROWS = {
    "OCEAN_FREIGHT": 23, "BAF": 23,
    "THC": 24, "OTHC": 24, "DTHC": 24,
    "DOC_FEE": 25, "DOC": 25, "BL": 25,
    "SEAL": 26, "WFG": 27, "WHARFAGE": 27, "PSF": 28, "PFS": 29,
    "CFS": 30, "SHUTTLE": 31, "SHORING": 32, "CUSTOMS": 33, "INSURANCE": 34, "INS": 34,
}
BREAKDOWN_LABELS = {"OCEAN_FREIGHT": "Base", "BAF": "BAF", "OTHC": "Origin THC", "DTHC": "Destination THC",
                    "DOC": "DOC", "DOC_FEE": "DOC", "BL": "B/L"}


def display_groups(items):
    """동일 통화·부과 기준·수량의 비용만 합친다. 원래 QuoteItem은 변경하지 않는다."""
    groups, extras = {}, []
    has_freight = any(i.charge_code == "OCEAN_FREIGHT" for i in items)
    for item in sorted(items, key=lambda i: i.charge_code == "BAF"):
        row = COST_ROWS.get(item.charge_code)
        if row is None or (item.charge_code == "BAF" and not has_freight):
            extras.append(item)
            continue
        group = groups.setdefault(row, [])
        first = group[0] if group else None
        compatible = first is None or (
            (first.currency, first.basis, first.qty) == (item.currency, item.basis, item.qty)
            and item.qty is not None and item.qty > 0
            and all(i.rate is not None and i.amount is not None for i in (first, item))
            and item.basis not in ("PERCENT", "AT_COST")
        )
        (group if compatible else extras).append(item)
    return groups, extras


def extra_charge_text(item, lang):
    name = "Handling Charge" if item.charge_code == "HANDLING" else item.charge_label
    amount = f"{item.currency}{item.amount:,.2f}" if item.amount is not None else "AT COST"
    calc = (f"{item.currency}{item.rate:,.2f} {BASIS_LABEL.get(item.basis, item.basis)} x {_fmt_num(item.qty)} = "
            if item.rate is not None and item.qty is not None else "")
    included = ("TOTAL 포함" if lang == "ko" else "included in TOTAL") if item.amount is not None else (
        "실비 별도 / TOTAL 미포함" if lang == "ko" else "at cost / excluded from TOTAL")
    return f"* {name}: {calc}{amount} ({included})"


def render_xlsx(case: Case, quote: Quote, count: int | None = None, count_note: str | None = None,
                schedule_unconfirmed: bool = False) -> bytes:
    """발주 측 XLSX의 고정 비용 12행·도형·행 높이·병합을 그대로 두고 값만 채운다."""
    qi = case.quote_input
    items = quote.items
    wb = load_workbook(settings.quote_template)
    ws = wb.active
    is_lcl = qi.containerType == "LCL"

    # 인쇄: 양식의 인쇄 범위(A:F)를 한 페이지 폭에 맞춘다 — 없으면 회사 주소·DATE·REMARK(F열)가 2쪽으로 밀린다
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    lang = language()
    t = TEXT[lang]
    ws["B9"] = recipient_line(case, lang)
    ws["B10"] = t["from"]
    # 영문 생성본의 문구만 바꾸며 글꼴·배치·원본 파일은 그대로 둔다. ko는 원본 셀을 유지한다.
    if lang == "en":
        for ref, key in (("B12", "subject"), ("B14", "greeting1"), ("B15", "greeting2"),
                         ("B48", "thanks"), ("F49", "seal")):
            ws[ref] = t[key]
    # 케이스 번호는 하단 Remark에만 넣는다.

    size = f" / {qi.boxL} X {qi.boxW} X {qi.boxH}MM" if None not in (qi.boxL, qi.boxW, qi.boxH) else ""
    cbm = f" / {_fmt_num(qi.totalCbm)}CBM" if qi.totalCbm is not None else ""
    cargo = f": {qi.qty} {qi.qtyUnit}{size} / {_fmt_num(qi.grossWeightKg)}KGS{cbm}"
    ws["C17"] = f": {port_label(qi.pol)} - {port_label(qi.pod)}"
    if is_lcl:
        ws["C18"] = ": BY CONTAINER SHIPMENT (LCL)"
        rt = next((i.qty for i in items if i.basis == "PER_RT" and i.qty is not None), None)
        ws["C19"] = cargo + (t["rt_basis"].format(rt=_fmt_num(rt)) if rt is not None else "")
        ws["D22"] = "R/T"
    else:
        ws["C18"] = f": BY CONTAINER SHIPMENT ({qi.containerType} X {count or 1})"
        ws["C19"] = cargo
        ws["D22"] = "20FT" if (qi.containerType or "").startswith("20") else qi.containerType
    ws["C20"] = f": {qi.incoterms}"  # 발주 측 양식은 이 줄(PAYMENT TERM)에 인코텀즈를 적는다 (예시 견적서 기준)
    ws["F21"] = f"{kst(quote.created_at):%Y-%m-%d}" if quote.created_at else ""

    for row in range(ITEM_FIRST_ROW, ITEM_LAST_ROW + 1):
        for col in "CDEF":
            ws[f"{col}{row}"] = None
    groups, extras = display_groups(items)
    for row, group in groups.items():
        item = group[0]
        if item.basis == "PERCENT" and item.amount is not None:  # 화물가액 % (산정식은 비고 칸)
            ws[f"E{row}"] = item.amount
            ws[f"E{row}"].number_format = MONEY_FORMAT.get(item.currency, "#,##0.00")
        elif item.basis == "AT_COST" or item.rate is None:
            ws[f"E{row}"] = "AT COST"
        else:
            ws[f"C{row}"] = sum(i.rate for i in group)
            ws[f"C{row}"].number_format = MONEY_FORMAT.get(item.currency, "#,##0.00")
            ws[f"D{row}"] = item.qty
            ws[f"E{row}"] = money(sum(i.amount for i in group), item.currency)
            ws[f"E{row}"].number_format = MONEY_FORMAT.get(item.currency, "#,##0.00")
        if len(group) == 1:
            ws[f"F{row}"] = item_remark(item, lang)
        else:
            ws[f"F{row}"] = "\n".join(
                f"{BREAKDOWN_LABELS.get(i.charge_code, i.charge_code)} {i.currency}{i.rate:,.2f} {BASIS_LABEL.get(i.basis, i.basis)}"
                for i in group)
            font = copy(ws[f"F{row}"].font)
            font.sz = 8.5
            ws[f"F{row}"].font = font
    ws[f"C{TOTAL_ROW}"] = total_text(quote)

    as_of = qi.exchangeRateAsOf or (f"{kst(quote.created_at):%Y-%m-%d}" if quote.created_at else "-")
    # 통화코드 단일 필드와 계산은 유지하고, 고지에는 실제 합계의 통화를 모두 표시한다 (FR-508).
    currencies = sorted({total.currency for total in quote.totals}, key=lambda c: (c != "USD", c))
    currency_notice = " / ".join(currencies) or qi.quoteCurrency or "USD"
    remarks = [t["fx"].format(currency=currency_notice, as_of=as_of), t["subject_to_change"]]
    if schedule_unconfirmed:
        remarks.append(t["schedule_tbc"].format(carrier=f" ({quote.carrier})" if quote.carrier else ""))
    else:
        carrier = f" ({quote.carrier})" if quote.carrier else ""
        tbc = "확인 필요" if lang == "ko" else "to be confirmed"
        remarks.append(f"* ETD / ETA : {quote.etd or tbc} / {quote.eta or tbc}{carrier}")
    tbc = "확인 필요" if lang == "ko" else "to be confirmed"
    remarks.append(t["transit"].format(days=qi.transitTime) if qi.transitTime is not None else f"* T/T : {tbc}")
    det = f"{qi.freeTimeDet}{'일' if lang == 'ko' else ' days'}" if qi.freeTimeDet is not None else tbc
    dem = f"{qi.freeTimeDem}{'일' if lang == 'ko' else ' days'}" if qi.freeTimeDem is not None else tbc
    remarks.append(f"* FREE TIME : DET {det} / DEM {dem}")
    if count_note:
        remarks.append(f"* {count_note_text(count_note, lang)}")
    remarks.extend(incoterm_remarks(qi.incoterms, lang))
    remarks.append(t["validity"].format(date=_korean_date(quote.valid_until) if lang == "ko" else quote.valid_until or ""))
    remarks.append(t["quote_no"].format(case_id=case.case_id, version=quote.version_no))
    remarks.extend(extra_charge_text(item, lang) for item in extras)
    if len(remarks) > REMARK_LAST_ROW - REMARK_FIRST_ROW + 1:
        raise QuoteGenerationError("원본 양식의 Remark 공간이 부족합니다. 비용이나 필수 안내를 생략할 수 없어 담당자 확인이 필요합니다.")
    for offset, row in enumerate(range(REMARK_FIRST_ROW, REMARK_LAST_ROW + 1)):
        ws[f"B{row}"] = remarks[offset] if offset < len(remarks) else None

    return save_filled_template(wb, settings.quote_template)


def _save(case: Case, quote: Quote, ext: str, data: bytes) -> str:
    """storage/<케이스ID>/quotes/<케이스ID>_견적서.<ext> — 이전 판은 build()가 retire_files()로 먼저 옮긴다."""
    return drafts.save_quote_file(case, f"{case.case_id}_견적서.{ext}", data)


def retire_files(case: Case, version: int) -> list[str]:
    """자동으로 만든 견적서(PDF·XLSX)를 quotes/에서 quote-history/로 함께 옮긴다. 옮긴 파일 이름을 돌려준다.

    quotes/ 폴더의 파일은 [견적서 송부 초안]에 그대로 첨부된다. 그래서
    - 새 판을 만들 때 이전 판의 PDF·XLSX를 한 번에 옮기고 (PDF 생성이 실패해도 옛 PDF가 남지 않게),
    - 새 견적서를 만들지 못한 경우(정보부족·보류)에도 옛 견적서를 치운다.
    담당자가 직접 올린 다른 이름의 파일은 건드리지 않는다.
    """
    moved: list[str] = []
    for ext in ("xlsx", "pdf"):
        current = settings.storage_dir / case.case_id / "quotes" / f"{case.case_id}_견적서.{ext}"
        if not current.exists():
            continue
        history = settings.storage_dir / case.case_id / "quote-history"
        history.mkdir(parents=True, exist_ok=True)
        target = history / f"{case.case_id}_견적서_v{version}.{ext}"
        if target.exists():
            target = history / f"{case.case_id}_견적서_v{version}_{uuid.uuid4().hex[:8]}.{ext}"
        os.replace(current, target)
        moved.append(target.name)
    return moved


def retire_current(session: Session, case: Case) -> list[str]:
    """지금 조건으로는 견적서를 낼 수 없을 때 부른다 — 마지막 판 번호로 보관 폴더에 옮긴다."""
    version = session.scalar(select(func.max(Quote.version_no)).where(Quote.case_pk == case.id)) or 0
    return retire_files(case, version)


# ------------------------------------------------------------------ PDF (LibreOffice)

log = logging.getLogger(__name__)
_WINDOWS_SOFFICE = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
)
_MACOS_SOFFICE = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "~/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


def find_soffice() -> str | None:
    """명시한 경로 → PATH → 운영체제별 기본 설치 경로. 잘못된 명시 경로는 자동 대체하지 않는다."""
    if settings.soffice_path:
        configured = Path(settings.soffice_path).expanduser()
        return str(configured) if configured.is_file() else None
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    defaults = _MACOS_SOFFICE if sys.platform == "darwin" else _WINDOWS_SOFFICE
    for candidate in defaults:
        path = Path(candidate).expanduser()
        if path.is_file():
            return str(path)
    return None


def _validate_pdf(data: bytes) -> None:
    """파일 확장자나 헤더만 믿지 않고 페이지·내용 스트림까지 읽을 수 있는지 확인한다."""
    if not data.startswith(b"%PDF-"):
        raise ValueError("PDF header missing")
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted or not reader.pages:
        raise ValueError("PDF is encrypted or has no pages")
    for page in reader.pages:
        contents = page.get_contents()
        if contents is not None:
            contents.get_data()


def to_pdf(xlsx: bytes) -> tuple[bytes | None, str | None]:
    """양식 XLSX → PDF. LibreOffice가 있으면 그것으로, 없거나 실패하면 내장 엔진(quote_pdf)으로 만든다.

    둘 다 실패하면 (None, 사유) — build가 두 파일의 송부를 함께 보류한다.
    """
    pdf, reason = _soffice_pdf(xlsx)
    if pdf:
        return pdf, None
    try:
        from . import quote_pdf

        pdf = quote_pdf.render_pdf(xlsx)
        _validate_pdf(pdf)
        return pdf, f"{reason} — 내장 엔진으로 PDF를 만들었습니다"
    except Exception as error:  # noqa: BLE001 — 검증 결과에 실패 원인을 전달한다
        log.warning("builtin pdf failed: %s", error)
        return None, f"{reason}, 내장 엔진도 실패 ({error.__class__.__name__}) — PDF·XLSX 송부 보류"


def _soffice_pdf(xlsx: bytes) -> tuple[bytes | None, str | None]:
    """LibreOffice 변환. 못 만들면 (None, 사유)."""
    soffice = find_soffice()
    if soffice is None:
        return None, "LibreOffice 실행 파일 없음 — 설치 경로 또는 SOFFICE_PATH 설정을 확인해 주세요"
    try:
        with tempfile.TemporaryDirectory(prefix="leona-pdf-") as tmp:
            src = Path(tmp) / "quote.xlsx"
            src.write_bytes(xlsx)
            profile = (Path(tmp) / "profile").as_uri()  # 담당자가 LibreOffice를 켜 둬도 충돌하지 않게 별도 프로필
            result = subprocess.run(
                [soffice, "--headless", "--norestore", f"-env:UserInstallation={profile}",
                 "--convert-to", "pdf", "--outdir", tmp, str(src)],
                capture_output=True, timeout=settings.pdf_timeout_s, check=False,
            )
            if result.returncode != 0:
                return None, (f"LibreOffice 변환 실패 (종료 코드 {result.returncode}) — "
                              "LibreOffice 설치 상태와 실행 권한을 확인해 주세요")
            out = Path(tmp) / "quote.pdf"
            if not out.is_file():
                return None, "LibreOffice 변환 결과 없음 — PDF 파일이 생성되지 않았습니다"
            pdf = out.read_bytes()
            try:
                _validate_pdf(pdf)
            except Exception as error:  # noqa: BLE001 — 손상된 변환 결과는 송부하지 않는다
                log.warning("invalid LibreOffice PDF: %s", error.__class__.__name__)
                return None, "LibreOffice PDF 오류 — 파일이 손상되었거나 읽을 수 있는 페이지가 없습니다"
            return pdf, None
    except subprocess.TimeoutExpired:
        return None, (f"LibreOffice 변환 시간 초과 ({settings.pdf_timeout_s:g}초) — "
                      "LibreOffice 실행 상태를 확인해 주세요")
    except OSError as error:
        log.warning("pdf conversion or file access failed: %s", error.__class__.__name__)
        return None, (f"LibreOffice 실행·파일 처리 실패 ({error.__class__.__name__}) — "
                      "설치 경로와 실행·임시 폴더 접근 권한을 확인해 주세요")
